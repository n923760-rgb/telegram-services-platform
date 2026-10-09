"""Source-preserving text tables and static row-level review; no customer formulas."""

from copy import copy
from io import BytesIO

from openpyxl import load_workbook

from app.builders import excel
from app.builders.quality import validate
from app.builders.schema import Table
from app.core.i18n import tr


def _sheet(title, columns, rows, *, literal=False):
    book = load_workbook(BytesIO(excel.build(Table(title=title, columns=columns, rows=rows))))
    sheet = book.active
    if literal:
        # Explicit text cell types preserve the exact string, without apostrophe prefixes.
        # Empty CSV fields display as blank cells. Never ask Excel to execute source text.
        for index, values in enumerate([columns, *rows], 3):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(index, column, value)
                cell.data_type = "s"
                cell.number_format = "@"
        for index, column in enumerate(columns):
            sheet.tables["Records"].tableColumns[index].name = column
    for dimension in sheet.column_dimensions.values():
        dimension.width = max(14, dimension.width)
    return book


def build(headers, rows, *, trim, language):
    # Called with validated parsed CSV. Original record order and count never change.
    columns = [header.strip() for header in headers]
    clean = [[value.strip(" \t") if trim else value for value in row] for row in rows]
    audit = []
    seen = {}
    counts = {"rows": len(rows), "blanks": 0, "duplicates": 0, "changes": 0}
    for index, (source, current) in enumerate(zip(rows, clean, strict=True), 1):
        blanks = [str(i) for i, value in enumerate(current, 1) if value == ""]
        changes = [
            str(i)
            for i, (old, new) in enumerate(zip(source, current, strict=True), 1)
            if old != new
        ]
        key = tuple(current)
        duplicate = seen.get(key)
        seen.setdefault(key, index)
        counts["blanks"] += len(blanks)
        counts["duplicates"] += duplicate is not None
        counts["changes"] += len(changes)
        audit.append([index, ", ".join(blanks), duplicate, ", ".join(changes)])
    record_heading = tr("csv_review_source_record", language)
    candidate = record_heading
    suffix = 2
    while candidate.casefold() in {column.casefold() for column in columns}:
        candidate = f"{record_heading} ({suffix})"
        suffix += 1
    # A stable generated reference stays attached when the customer sorts a table.
    data_columns = [*columns, candidate]
    data_rows = [[*row, str(index)] for index, row in enumerate(clean, 1)]
    source_rows = [[*row, str(index)] for index, row in enumerate(rows, 1)]
    book = _sheet(tr("csv_review_data", language), data_columns, data_rows, literal=True)
    book.active.title = "Data"
    source = _sheet(
        tr("csv_review_source", language), [*headers, candidate], source_rows, literal=True
    )
    report = _sheet(
        tr("csv_review_audit", language),
        [
            tr(key, language)
            for key in (
                "csv_review_record",
                "csv_review_blank_columns",
                "csv_review_duplicate_of",
                "csv_review_changed_columns",
            )
        ],
        audit,
    )
    for donor, name in [(source, "Source"), (report, "Review")]:
        sheet = book.create_sheet(name)
        original = donor.active
        for row in original:
            for cell in row:
                target = sheet.cell(cell.row, cell.column, cell.value)
                target.data_type = cell.data_type
                for attribute in ("font", "fill", "border", "alignment", "protection"):
                    setattr(target, attribute, copy(getattr(cell, attribute)))
                target.number_format = cell.number_format
        for key, dimension in original.row_dimensions.items():
            sheet.row_dimensions[key] = copy(dimension)
        for key, dimension in original.column_dimensions.items():
            sheet.column_dimensions[key] = copy(dimension)
        for attribute in (
            "sheet_properties",
            "sheet_format",
            "page_setup",
            "page_margins",
            "oddFooter",
        ):
            setattr(sheet, attribute, copy(getattr(original, attribute)))
        sheet.sheet_view.rightToLeft = original.sheet_view.rightToLeft
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A4"
        sheet.print_area = original.print_area
        sheet.print_title_rows = original.print_title_rows
        for merged in original.merged_cells:
            sheet.merge_cells(str(merged))
        table = original.tables["Records"]
        table.name = table.displayName = name + "Records"
        sheet.add_table(table)
        donor.close()
    review = book["Review"]
    rules = tr(
        "csv_review_rules",
        language,
        mode=tr("csv_review_trim" if trim else "csv_review_preserve", language),
    )
    review["A2"] = tr("csv_review_summary", language, **counts) + "\n" + rules
    review["A2"].alignment = copy(review["A1"].alignment)
    review.merge_cells("A2:D2")
    review.row_dimensions[2].height = 120
    # Header normalization is explicit and independently visible in Source/Data.
    book["Data"]["A2"] = tr(
        "csv_review_mode_trim" if trim else "csv_review_mode_preserve", language
    )
    book["Data"].merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(data_columns))
    book["Data"]["A2"].alignment = copy(book["Data"]["A1"].alignment)
    book["Data"].row_dimensions[2].height = 75
    content = BytesIO()
    book.save(content)
    book.close()
    return validate(content.getvalue(), "xlsx"), counts
