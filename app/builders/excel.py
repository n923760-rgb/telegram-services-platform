import re
from datetime import date
from functools import lru_cache
from io import BytesIO
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from PIL import ImageFont

from app.builders.direction import is_rtl
from app.builders.quality import validate
from app.builders.schema import Table
from app.builders.word import INLINE_LTR, safe_text

HEADER_FILL = PatternFill("solid", fgColor="176B55")
HEADER_FONT = Font(name="DejaVu Sans", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="DejaVu Sans", size=11)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _reading_rtl(text):
    # Dates and ASCII contacts are identifiers, not evidence of English prose.
    # Strip them only for detection; preserve every cell value verbatim.
    return is_rtl(INLINE_LTR.sub("", text))


def cell_safe(value):
    if isinstance(value, str):
        value = safe_text(value)
        if value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


@lru_cache(maxsize=3)
def _measurement_font(size, bold):
    filename = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    for directory in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"):
        path = Path(directory) / filename
        if path.is_file():
            # Measure at 3x 96-DPI resolution: 11 points = 44 pixels here.
            return ImageFont.truetype(str(path), size=round(size * 4))
    # Retain non-Docker builder compatibility; native visual qualification uses
    # the declared DejaVu fonts installed in every production image.
    return None


def _text_pixels(text, *, size=11, bold=False):
    font = _measurement_font(size, bold)
    if font is None:
        return max(
            (
                sum(2 if ord(c) > 0x2E7F else 1 for c in line) * 7 * size / 11
                for line in text.split("\n")
            ),
            default=0,
        )
    return max((font.getlength(line) / 3 for line in text.split("\n")), default=0)


def _display_width(value, *, bold=False):
    text = "" if value is None else str(value)
    # Excel column units are based on its default font's digit width (~7px).
    # Leave a margin for native application/font substitution differences.
    return _text_pixels(text, bold=bold) * 1.12 / 7


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _wrapped_lines(text, width, *, size=11, bold=False):
    available = max(1, width * 7 - 12)
    total = 0
    measured = {}
    for chunk in ("" if text is None else str(text)).split("\n"):
        lines, used, gap = 1, 0, 0
        for token in re.findall(r"\s+|\S+", chunk):
            if token not in measured:
                measured[token] = _text_pixels(token, size=size, bold=bold) * 1.12
            pixels = measured[token]
            if token.isspace():
                gap += pixels
                continue
            # Native Excel wraps whole words before splitting long unbroken text.
            # Dividing total text width alone misses the unused space on each line.
            if used and used + gap + pixels > available:
                lines += 1
                used, gap = 0, 0
            used += gap + pixels
            if used > available:
                extra = ceil(used / available) - 1
                lines += extra
                used -= extra * available
            gap = 0
        total += lines
    return total


def build(data: Table) -> bytes:
    data = Table.model_validate(data)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    workbook.properties.title = safe_text(data.title)
    workbook.properties.creator = "Telegram Services"

    columns = [safe_text(column) for column in data.columns]
    rows = [[cell_safe(value) for value in row] for row in data.rows]

    probe = " ".join(columns) + " " + " ".join(str(value) for value in (rows[0] if rows else []))
    rtl = _reading_rtl(probe)
    sheet.sheet_view.rightToLeft = rtl
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = not rtl

    # The title is separate from the native sortable data table.
    sheet.append([safe_text(data.title)])
    sheet["A1"].data_type = "s"  # A title beginning with '=' is still customer text.
    sheet["A1"].font = Font(name="DejaVu Sans", bold=True, size=18, color="176B55")
    sheet["A1"].alignment = Alignment(
        horizontal="right" if rtl else "left", vertical="center", wrap_text=True
    )
    if len(columns) > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet.row_dimensions[1].height = 48
    sheet.append([])

    sheet.append(columns)
    for row in rows:
        sheet.append(row)

    header_align = Alignment(
        horizontal="right" if rtl else "left", vertical="center", wrap_text=True
    )
    # XML sanitation must not silently create duplicate/blank native table headers.
    if any(not c.strip() for c in columns) or len({c.casefold() for c in columns}) != len(columns):
        raise ValueError("invalid Excel table headers")
    for cell in sheet[3]:
        cell.data_type = "s"
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = header_align
    sheet.row_dimensions[3].height = 36

    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            value = cell.value
            kind = data.column_formats[cell.column - 1].kind if data.column_formats else None
            if kind == "date" and value is not None:
                cell.value = date.fromisoformat(value)
                cell.number_format = "yyyy-mm-dd"
            elif kind == "percent" and value is not None:
                cell.number_format = "0.###############%"
            elif isinstance(value, str):
                cell.number_format = "@"
            elif _is_number(value):
                cell.number_format = (
                    "#,##0"
                    if isinstance(value, int)
                    else "0.##############E+00"
                    if value and abs(value) < 1e-10
                    else "#,##0.###############"
                )
            cell.font = BODY_FONT
            cell.border = BORDER
            if _is_number(value):
                horizontal = "right"
            elif isinstance(value, str) and _reading_rtl(value):
                horizontal = "right"
            else:
                horizontal = "left"
            cell.alignment = Alignment(
                horizontal=horizontal,
                vertical="top",
                wrap_text=True,
                readingOrder=2 if isinstance(value, str) and _reading_rtl(value) else 1,
            )

    widths = [_display_width(column, bold=True) for column in columns]
    for row in rows:
        for index, value in enumerate(row):
            widths[index] = max(widths[index], _display_width(value))
    for index, width in enumerate(widths, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(60, max(10, width + 2))
    title_width = sum(
        sheet.column_dimensions[get_column_letter(i)].width for i in range(1, len(columns) + 1)
    )
    sheet.row_dimensions[1].height = min(
        409, max(48, 25 * _wrapped_lines(data.title, title_width, size=18, bold=True))
    )
    sheet.row_dimensions[3].height = min(
        409,
        max(
            36,
            16
            * max(
                _wrapped_lines(
                    value, sheet.column_dimensions[get_column_letter(i)].width, bold=True
                )
                for i, value in enumerate(columns, 1)
            )
            + 8,
        ),
    )

    for row in sheet.iter_rows(min_row=4):
        lines = 1
        for cell in row:
            text = "" if cell.value is None else str(cell.value)
            col_width = max(8, sheet.column_dimensions[get_column_letter(cell.column)].width or 8)
            lines = max(lines, _wrapped_lines(text, col_width))
        sheet.row_dimensions[row[0].row].height = min(409, max(22, 16 * lines + 6))

    last_column = get_column_letter(len(columns))
    ref = f"A3:{last_column}{sheet.max_row}"
    native_table = ExcelTable(displayName="Records", ref=ref)
    native_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(native_table)
    sheet.freeze_panes = "A4"
    # The native table owns its filter; a worksheet filter must not overlap it.
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:{last_column}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape" if len(columns) > 5 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1 if len(columns) <= 8 else 0
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "&P / &N"
    buffer = BytesIO()
    workbook.save(buffer)
    return validate(buffer.getvalue(), "xlsx")
