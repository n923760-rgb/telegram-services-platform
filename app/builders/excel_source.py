"""Native source references stay inside each sortable record; raw source stays separate."""

from copy import copy
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo

from app.builders import excel
from app.builders.direction import is_rtl
from app.builders.quality import validate
from app.builders.schema import Table
from app.core.i18n import tr


def build(table, lines, row_ids, quotes):
    if not 1 <= len(lines) <= 200 or any(len(line) > 4000 for line in lines):
        raise ValueError("source line limits")
    if len(row_ids) != len(table.rows) or len(quotes) != len(table.rows):
        raise ValueError("source row limits")
    if row_ids != sorted(set(row_ids)) or any(not 1 <= i <= len(lines) for i in row_ids):
        raise ValueError("source row IDs")
    language = "ar" if is_rtl(" ".join([table.title, *table.columns])) else "en"
    header = tr("excel_source_reference", language)
    columns = [*table.columns]
    while header.casefold() in {name.casefold() for name in columns}:
        header += "_"
    formats = [fmt.model_dump() for fmt in table.column_formats]
    augmented = Table.model_validate(
        {
            "title": table.title,
            "columns": [*columns, header],
            "rows": [row + [f"L{i:03d}"] for row, i in zip(table.rows, row_ids, strict=True)],
            "column_formats": formats + [{"kind": "text"}] if formats else [],
        }
    )
    book = load_workbook(BytesIO(excel.build(augmented)))
    data = book.active
    for row in data.iter_rows(min_row=4):
        for cell in row:
            alignment = copy(cell.alignment)
            alignment.indent = 1
            alignment.vertical = "center"
            if cell.column == len(augmented.columns):
                alignment.horizontal = "center"
                alignment.indent = 0
            cell.alignment = alignment
    for letter in data.column_dimensions:
        data.column_dimensions[letter].width = min(60, data.column_dimensions[letter].width + 2)
    for r, source_quotes in enumerate(quotes, 4):
        if len(source_quotes) != len(table.columns):
            raise ValueError("source quote columns")
        for c, quote in enumerate(source_quotes, 1):
            cell = data.cell(r, c)
            kind = table.column_formats[c - 1].kind if table.column_formats else "number"
            original = table.rows[r - 4][c - 1]
            if isinstance(original, str) and kind != "date":
                # Explicit literal cell type replaces the legacy visible apostrophe escape.
                cell.value = original
                cell.data_type = "s"
            if quote is None or type(cell.value) not in (int, float):
                continue
            literal = quote[:-1] if kind == "percent" else quote
            precision = max(0, -Decimal(literal).as_tuple().exponent)
            cell.number_format = (
                ("0" if kind == "percent" else "#,##0")
                + ("." + "0" * precision if precision else "")
                + ("%" if kind == "percent" else "")
            )
            letter = get_column_letter(c)
            data.column_dimensions[letter].width = min(
                60, max(data.column_dimensions[letter].width, excel._display_width(quote) + 2)
            )
    excel.configure_print(data)
    source = book.create_sheet("Source")
    source.sheet_view.showGridLines = False
    source.sheet_view.rightToLeft = excel._reading_rtl(" ".join(lines))
    source.append([tr("excel_source_title", language), "", ""])
    source["A1"].font = Font(name="DejaVu Sans", size=16, bold=True)
    source["A1"].alignment = Alignment(
        horizontal="right" if source.sheet_view.rightToLeft else "left", vertical="center"
    )
    source.row_dimensions[1].height = 36
    source.merge_cells("A1:C1")
    source.append([])
    source.append(
        [
            tr("excel_source_reference", language),
            tr("excel_source_part", language),
            tr("excel_source_text", language),
        ]
    )
    for i, line in enumerate(lines, 1):
        parts = [line[j : j + 240] for j in range(0, len(line), 240)]
        for j, part in enumerate(parts, 1):
            source.append([f"L{i:03d}", f"{j}/{len(parts)}", part])
    source.column_dimensions["A"].width = 20
    source.column_dimensions["B"].width = 10
    source.column_dimensions["C"].width = 76
    for row in source.iter_rows(min_row=3):
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.font = excel.HEADER_FONT if cell.row == 3 else excel.BODY_FONT
            cell.alignment = Alignment(
                horizontal="center"
                if cell.column < 3
                else "right"
                if excel._reading_rtl(str(cell.value))
                else "left",
                vertical="center",
                indent=0 if cell.column < 3 else 1,
                wrap_text=True,
                readingOrder=2 if excel._reading_rtl(str(cell.value)) else 1,
            )
            if cell.row == 3:
                cell.fill = excel.HEADER_FILL
        source.row_dimensions[row[0].row].height = max(
            36 if row[0].row == 3 else 22,
            16 * max(excel._wrapped_lines(row[0].value, 20), excel._wrapped_lines(row[2].value, 76))
            + 6,
        )
    source.freeze_panes = "B4"
    native = ExcelTable(displayName="SourceLines", ref=f"A3:C{source.max_row}")
    native.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    source.add_table(native)
    source.print_title_rows = "1:3"
    source.print_area = f"A1:C{source.max_row}"
    excel.configure_print(source)
    result = BytesIO()
    book.save(result)
    return validate(result.getvalue(), "xlsx")
