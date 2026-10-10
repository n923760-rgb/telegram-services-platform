"""Source-preserving extracted PDF cells and separately proposed headings."""

from copy import copy
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.builders import excel
from app.builders.quality import validate
from app.builders.schema import Table
from app.core.i18n import tr


def append_sheet(book, title, columns, rows, name, note):
    donor = load_workbook(BytesIO(excel.build(Table(title=title, columns=columns, rows=rows))))
    original = donor.active
    sheet = book.create_sheet(name)
    for row in original:
        for cell in row:
            target = sheet.cell(cell.row, cell.column, cell.value)
            target.data_type = cell.data_type
            for attribute in ("font", "fill", "border", "alignment", "protection"):
                setattr(target, attribute, copy(getattr(cell, attribute)))
            target.number_format = cell.number_format
    # Keep every input cell as text, including formulas, leading zeros and date strings.
    for r, values in enumerate([columns, *rows], 3):
        for c, value in enumerate(values, 1):
            cell = sheet.cell(r, c, value)
            cell.data_type, cell.number_format = "s", "@"
    for key, dimension in original.row_dimensions.items():
        sheet.row_dimensions[key] = copy(dimension)
    for key, dimension in original.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dimension)
    for attribute in (
        "sheet_properties",
        "sheet_format",
        "page_setup",
        "page_margins",
        "oddFooter",
    ):
        setattr(sheet, attribute, copy(getattr(original, attribute)))
    sheet.sheet_view.rightToLeft = original.sheet_view.rightToLeft
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    sheet.print_area, sheet.print_title_rows = original.print_area, original.print_title_rows
    for merged in original.merged_cells:
        sheet.merge_cells(str(merged))
    table = original.tables["Records"]
    table.name = table.displayName = name + "Records"
    for column, label in zip(table.tableColumns, columns, strict=True):
        column.name = label
    sheet.add_table(table)
    sheet["A2"] = note
    sheet["A2"].data_type = "s"
    sheet["A2"].alignment = copy(sheet["A1"].alignment)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    sheet.row_dimensions[2].height = 100
    donor.close()
    return sheet


def build(extraction, labels, *, language, enhanced):
    book = Workbook()
    book.remove(book.active)
    label_map = {table.table_id: table for table in labels.tables}
    audit = []
    for number, table in enumerate(extraction.tables, 1):
        proposal = label_map[number]
        reference = tr("pdf_tables_reference", language)
        while reference.casefold() in {label.casefold() for label in proposal.columns}:
            reference += "_"
        source_columns = [
            tr("pdf_tables_column", language, number=c) for c in range(1, len(table.rows[0]) + 1)
        ]
        rows = [[*row, f"T{number}:R{r}"] for r, row in enumerate(table.rows, 1)]
        note = tr(
            "pdf_tables_sheet_note",
            language,
            page=table.page,
            mode=tr("pdf_tables_labels" if enhanced else "pdf_tables_direct", language),
        )
        append_sheet(
            book, proposal.title, [*proposal.columns, reference], rows, f"Data{number}", note
        )
        append_sheet(
            book,
            tr("pdf_tables_source", language, number=number),
            [*source_columns, reference],
            rows,
            f"Source{number}",
            tr("pdf_tables_source_note", language, page=table.page),
        )
        audit.append(
            [
                str(number),
                str(table.page),
                str(len(table.rows)),
                str(len(table.rows[0])),
                str(table.accuracy),
                str(table.whitespace),
            ]
        )
    append_sheet(
        book,
        tr("pdf_tables_audit", language),
        [
            tr("pdf_tables_audit_" + key, language)
            for key in ("table", "page", "rows", "columns", "accuracy", "whitespace")
        ],
        audit,
        "Review",
        tr("pdf_tables_audit_note", language),
    )
    content = BytesIO()
    book.save(content)
    book.close()
    return validate(content.getvalue(), "xlsx")
