import re
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.builders.direction import has_arabic, is_rtl
from app.builders.pptx_layout import FONT, plan_slide
from app.builders.quality import validate
from app.builders.schema import Deck
from app.builders.word import DATE_PATTERN, DATE_TOKEN, safe_text

ACCENT = RGBColor(0x17, 0x6B, 0x55)
DARK = RGBColor(0x22, 0x36, 0x3D)
MUTED = RGBColor(0x58, 0x68, 0x6D)
PAPER = RGBColor(0xF7, 0xFA, 0xF8)
COVER = RGBColor(0x14, 0x2D, 0x36)
COVER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
COVER_MUTED = RGBColor(0xD6, 0xE5, 0xE0)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
DISPLAY_LTR = re.compile(
    r"https?://[^\s<>]+|www\.[^\s<>]+|[a-z0-9.+_-]+@[a-z0-9.-]+\.[a-z0-9-]+|" + DATE_PATTERN,
    re.IGNORECASE,
)


def _reading_rtl(text):
    # Contacts/dates are identifiers, not English prose. Use the same contact-first
    # matcher as display formatting, including addresses starting with a date.
    return is_rtl(DISPLAY_LTR.sub("", text))


def _style_run(run, size, bold, color):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    properties = run._r.get_or_add_rPr()
    properties.set("lang", "ar-SA" if has_arabic(run.text) else "en-US")
    for tag in ("a:ea", "a:cs"):
        element = OxmlElement(tag)
        element.set("typeface", FONT)
        properties.append(element)


def _paragraph(paragraph, text, rtl, align, size, bold, color, *, bullet=False):
    visible = safe_text(text)
    # Match complete contacts before dates: a date inside a URL/email is part of
    # that identifier and must not acquire invisible characters when copied.
    # Standalone dates retain their source notation and digit script.
    if rtl:
        visible = DISPLAY_LTR.sub(
            lambda match: (
                "\u200e" + match[0] + "\u200e" if DATE_TOKEN.fullmatch(match[0]) else match[0]
            ),
            visible,
        )
    paragraph.text = visible
    paragraph.alignment = align
    properties = paragraph._p.get_or_add_pPr()
    properties.set("rtl", "1" if rtl else "0")
    if bullet:
        properties.set("marR" if rtl else "marL", str(int(Pt(20))))
        properties.set("indent", str(-int(Pt(12))))
    marker = OxmlElement("a:buChar" if bullet else "a:buNone")
    if bullet:
        marker.set("char", "•")
    properties.append(marker)
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(10 if bullet else 0)
    paragraph.line_spacing = Pt(size * 1.2)
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    for run in paragraph.runs:
        _style_run(run, size, bold, color)


def _frame(frame):
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(0.1)
    frame.margin_top = frame.margin_bottom = Inches(0.1)


def _textbox(slide, name, x, y, width, height):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    _frame(shape.text_frame)
    return shape.text_frame


def _table(slide, item, layout):
    data = item.table
    native = slide.shapes.add_table(
        len(data.rows) + 1,
        len(data.columns),
        Inches(0.85),
        Inches(1.8),
        Inches(11.6),
        Pt(sum(layout.row_heights)),
    )
    native.name = "Source records"
    table = native.table
    # Match the exact widths used by the layout gate. Keep the native table's
    # total width exact after converting point measurements to integer EMUs.
    remaining_width = native.width
    for index, width in enumerate(layout.column_widths):
        native_width = remaining_width if index == len(layout.column_widths) - 1 else Pt(width)
        table.columns[index].width = native_width
        remaining_width -= native_width
    table.first_row = True
    table.horz_banding = False  # Explicit colors, consistent across Office themes.
    table._tbl.tblPr.set("rtl", "1" if _reading_rtl(" ".join(data.columns)) else "0")
    for row_index, values in enumerate([data.columns, *data.rows]):
        table.rows[row_index].height = Pt(layout.row_heights[row_index])
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            _frame(cell.text_frame)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                ACCENT
                if row_index == 0
                else RGBColor(0xF0, 0xF6, 0xF3)
                if row_index % 2
                else RGBColor(0xFF, 0xFF, 0xFF)
            )
            for edge in ("L", "R", "T", "B"):
                border = OxmlElement("a:ln" + edge)
                border.set("w", "9525")  # 0.75pt native cell borders.
                fill = OxmlElement("a:solidFill")
                color = OxmlElement("a:srgbClr")
                color.set("val", "FFFFFF" if row_index == 0 else "DFE8E4")
                fill.append(color)
                border.append(fill)
                cell._tc.get_or_add_tcPr().insert_element_before(border, "a:solidFill")
            rtl = _reading_rtl(value)
            _paragraph(
                cell.text_frame.paragraphs[0],
                value,
                rtl,
                PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT,
                21 if row_index == 0 else layout.body_size,
                row_index == 0,
                RGBColor(0xFF, 0xFF, 0xFF) if row_index == 0 else DARK,
            )


def build(data: Deck) -> bytes:
    data = Deck.model_validate(data)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    prs.core_properties.title = safe_text(data.slides[0].title)
    prs.core_properties.author = "Telegram Services"
    for index, item in enumerate(data.slides):
        slide = prs.slides.add_slide(blank)
        layout = plan_slide(item, cover_candidate=index == 0 and len(data.slides) > 1)
        cover = layout.kind == "cover"
        # Native fills keep every foreground item editable and require no image assets.
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = COVER if cover else PAPER
        title_frame = _textbox(
            slide, "Slide title", 0.85, 1.55 if cover else 0.45, 11.6, 1.6 if cover else 1.3
        )
        rtl = _reading_rtl(item.title)
        _paragraph(
            title_frame.paragraphs[0],
            item.title,
            rtl,
            PP_ALIGN.CENTER if cover else PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT,
            layout.title_size,
            True,
            COVER_TEXT if cover else ACCENT,
        )
        if item.table is not None:
            _table(slide, item, layout)
        else:
            body = _textbox(
                slide, "Slide content", 0.85, 3.55 if cover else 1.8, 11.6, 1.9 if cover else 4.85
            )
            for bullet_index, text in enumerate(item.bullets):
                paragraph = body.paragraphs[0] if bullet_index == 0 else body.add_paragraph()
                rtl = _reading_rtl(text)
                _paragraph(
                    paragraph,
                    text,
                    rtl,
                    PP_ALIGN.CENTER if cover else PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT,
                    layout.body_size,
                    False,
                    COVER_MUTED if cover else DARK,
                    bullet=not cover,
                )
                if cover:
                    paragraph.space_after = Pt(10 if bullet_index < len(item.bullets) - 1 else 0)
                elif bullet_index == len(item.bullets) - 1:
                    paragraph.space_after = Pt(0)
        footer = _textbox(slide, "Slide number", 11.15, 7.0, 1.3, 0.35)
        footer.margin_top = footer.margin_bottom = Pt(0)
        _paragraph(
            footer.paragraphs[0],
            f"{index + 1} / {len(data.slides)}",
            False,
            PP_ALIGN.RIGHT,
            12,
            False,
            COVER_MUTED if cover else MUTED,
        )

    buffer = BytesIO()
    prs.save(buffer)
    return validate(buffer.getvalue(), "pptx")
