"""Source-owned professional template; customer strings are escaped values only.

Cache immutable template bytes, never a rendered customer document. Native tables
are inserted by the existing Word builder at request-local, opaque placeholders.
"""

from functools import lru_cache
from io import BytesIO
from uuid import uuid4

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docxtpl import DocxTemplate

from app.builders.word_schema import WordDocument


@lru_cache(maxsize=1)
def _template_bytes() -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name, size in (("Normal", 11), ("Title", 22), ("Heading 1", 14), ("List Bullet", 11)):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.widow_control = True
        style.paragraph_format.space_after = Pt(6)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.15
    for name in ("Title", "Heading 1"):
        properties = doc.styles[name].element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
        style = doc.styles[name]
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0 if name == "Title" else 14)
        style.paragraph_format.space_after = Pt(14 if name == "Title" else 6)
    # Paragraph tags remove their own control paragraphs, avoiding blank pages.
    for text, style in (
        ("{%p if include_title %}", None),
        ("{{ title }}", "Title"),
        ("{%p endif %}", None),
        ("{%p for section in sections %}", None),
        ("{%p if section.heading %}", None),
        ("{{ section.heading }}", "Heading 1"),
        ("{%p endif %}", None),
        ("{%p for paragraph in section.paragraphs %}", None),
        ("{{ paragraph }}", None),
        ("{%p endfor %}", None),
        ("{%p for bullet in section.bullets %}", None),
        ("{{ bullet }}", "List Bullet"),
        ("{%p endfor %}", None),
        ("{%p for marker in section.markers %}", None),
        ("{{ marker }}", None),
        ("{%p endfor %}", None),
        ("{%p endfor %}", None),
    ):
        doc.add_paragraph(text, style=style)
    doc.core_properties.author = "Telegram Services"
    doc.core_properties.comments = ""
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render(data: WordDocument, *, include_title: bool):
    sections, slots = [], {}
    for section in data.sections:
        markers = []
        for table in section.tables:
            marker = uuid4().hex
            markers.append(marker)
            slots[marker] = table
        sections.append({**section.model_dump(), "markers": markers})
    template = DocxTemplate(BytesIO(_template_bytes()))
    template.render(
        {"title": data.title, "include_title": include_title, "sections": sections},
        autoescape=True,
    )
    return template.docx, slots
