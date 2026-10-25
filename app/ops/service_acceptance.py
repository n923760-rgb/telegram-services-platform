"""Live synthetic acceptance in a fresh, explicitly isolated test database only."""

import argparse
import asyncio
import hashlib
import json
import re
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.engine import make_url

from app.core.db import sessions
from app.core.models import (
    CostHold,
    CostUsage,
    Job,
    JobAttempt,
    Ledger,
    Order,
    Service,
    Setting,
    User,
)
from app.core.settings import config
from app.ops.acceptance_cases import OWNER, cases, picture, source_pdf
from app.ops.costs import uncertain_call
from app.orders.engine import confirm_structure, fail_locked, register, submit
from app.orders.locking import lock_order
from app.providers.ai.gateway import AI
from app.providers.ai.openai_compatible import OpenAICompatible
from app.providers.ai.review import QualityReview
from app.providers.ai.validation import validation_categories
from app.providers.runtime import runtime_for, storage
from app.providers.storage import OwnedStorage
from app.services.base import ServiceError
from app.services.catalog import CATEGORY_KEYS
from app.services.registry import registry
from app.wallet.ledger import apply, balance
from app.workers.runner import execute_job

LIMIT = Decimal("5")
CALLS = 96
ACTIVE = ("awaiting_payment", "queued", "processing", "waiting_confirmation", "delivering")
ASSETS = {
    name: f"{OWNER}/" + str(i) * 32 + "/source." + ("pdf" if name.startswith("pdf") else "png")
    for i, name in enumerate(("pdf", "pdf2", "image", "shaded", "mixed"), 1)
}


def isolated():
    cfg = config()
    database = make_url(cfg.database_url.get_secret_value()).database or ""
    root = cfg.storage_root
    if (
        cfg.app_env != "test"
        or not re.fullmatch(r"svc_[a-f0-9]{12}_acceptance_test", database)
        or not re.fullmatch(r"acceptance-[a-f0-9]{12}", root.name)
        or root.is_symlink()
        or cfg.admin_ids != [OWNER]
        or cfg.stars_enabled
        or not cfg.ai_enabled
    ):
        raise ServiceError("acceptance_isolation_required")
    for value in (cfg.max_daily_cost_sar, cfg.max_cost_per_user_per_day):
        if not value.is_finite() or not 0 < value <= LIMIT:
            raise ServiceError("acceptance_budget_required")
    if not 0 < cfg.max_job_cost_sar <= Decimal(".75"):
        raise ServiceError("acceptance_budget_required")
    return root


class AcceptanceProvider(OpenAICompatible):
    """Additional whole-run conservative bound; ordinary DB job budgets still apply."""

    def __init__(self):
        super().__init__()
        self.calls = 0
        self.bounds = Decimal("0")

    def upper_bound(self, *args):
        bound = super().upper_bound(*args)
        if self.calls >= CALLS or self.bounds + bound > LIMIT:
            raise ServiceError("acceptance_cost_cap")
        self.bounds += bound  # Never refund this independent run allowance.
        return bound

    async def call(self, *args):
        if self.calls >= CALLS:
            raise ServiceError("acceptance_cost_cap")
        self.calls += 1
        return await super().call(*args)


class NativeAcceptanceProvider(AcceptanceProvider):
    async def call(self, *args):
        raise ServiceError("acceptance_native_provider_forbidden")


def pptx_token_diagnostics(content, source):
    """Observe only counts/fixed field kinds in isolated acceptance, never literal values.

    Locations describe occurrences beyond the source budget in traversal order;
    they do not determine which copy is semantically appropriate to remove.
    """
    from app.services.protected_values import ProtectedValues
    from app.services.text_editing.schema import CONTACT_TOKEN
    from app.services.text_to_pptx.schema import DeckPlan, source_token_violations

    if (
        not isinstance(content, str)
        or len(content) > 100000
        or not isinstance(source, str)
        or len(source) > 12000
    ):
        return None
    try:
        values = ProtectedValues(source, CONTACT_TOKEN)
        plan = DeckPlan.model_validate(values.resolve(json.loads(content)))
    except ValueError:
        return None
    if plan.deck is None or plan.missing_information:
        return None
    totals = dict.fromkeys(
        ("numeric_added", "numeric_repeated", "contact_added", "contact_repeated"), 0
    )
    locations = list(source_token_violations(source, plan.deck))
    for location in locations:
        totals[location["kind"]] += location["occurrences"]
    return {"counts": totals, "locations": locations[:32], "truncated": len(locations) > 32}


class DiagnosticAI(AI):
    """Observe validation in isolated acceptance only; never record model/source text."""

    def __init__(self, *args, diagnostics, pptx_source=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.diagnostics = diagnostics
        self.pptx_source = pptx_source

    async def _extract(self, schema, *args, **kwargs):
        stage = "review" if schema is QualityReview else "generation"
        diagnostics, provider = self.diagnostics, self.provider
        pptx_source = self.pptx_source

        class ObservedSchema:
            diagnostic_categories = getattr(schema, "diagnostic_categories", {})

            @staticmethod
            def model_json_schema():
                return schema.model_json_schema()

            @staticmethod
            def model_validate_json(content):
                event = {"call": provider.calls, "stage": stage}
                try:
                    candidate = schema.model_validate_json(content)
                except ValueError as error:
                    # Do not retain messages, locations, input, context or custom error types.
                    event.update(outcome="invalid", categories=validation_categories(schema, error))
                    if (
                        stage == "generation"
                        and pptx_source is not None
                        and "pptx_source_tokens" in event["categories"]
                    ):
                        details = pptx_token_diagnostics(content, pptx_source)
                        if details is not None:
                            event["token_diagnostics"] = details
                    diagnostics.append(event)
                    raise
                diagnostics.append(event | {"outcome": "valid"})
                if stage == "review":
                    diagnostics[-1].update(accepted=candidate.accepted, issues=candidate.issues)
                else:
                    # Fixed booleans only: never copy clarification text or model content.
                    for flag in ("missing_information", "ambiguous"):
                        value = getattr(candidate, flag, None)
                        if type(value) is bool:
                            diagnostics[-1][flag] = value
                return candidate

        return await super()._extract(ObservedSchema, *args, **kwargs)


def selected_cases(templates, names=None):
    if names is None:
        return templates
    if not names or len(set(names)) != len(names) or set(names) - {c.name for c in templates}:
        raise ServiceError("acceptance_case_invalid")
    return [case for case in templates if case.name in names]


def native_cases(templates, names):
    if not names:
        raise ServiceError("acceptance_case_invalid")
    chosen = selected_cases(templates, names)
    if any(registry.types[case.slug].needs_ai(case.inputs) for case in chosen):
        raise ServiceError("acceptance_native_requires_no_ai")
    return chosen


def pixel_signature(image):
    """Decoded RGB dimensions/pixels of synthetic assets, independent of PNG encoding."""
    with image.convert("RGB") as rgb:
        return hashlib.sha256(f"{rgb.width}x{rgb.height}:".encode() + rgb.tobytes()).hexdigest()


def inspect_output(name, data, case):
    suffix = Path(name).suffix.lower()
    text = ""
    details = {}
    if suffix == ".pdf":
        pdf = PdfReader(BytesIO(data))
        if pdf.is_encrypted or len(pdf.pages) < case.pdf_pages:
            raise ServiceError("acceptance_output_invalid")
        if case.exact_pdf_pages is not None and len(pdf.pages) != case.exact_pdf_pages:
            raise ServiceError("acceptance_output_invalid")
        if case.pdf_rotations:
            rotations = tuple(page.rotation for page in pdf.pages)
            if rotations != case.pdf_rotations:
                raise ServiceError("acceptance_output_invalid")
            details["rotations"] = list(rotations)
        if case.pdf_page_tokens:
            if len(pdf.pages) != len(case.pdf_page_tokens):
                raise ServiceError("acceptance_output_invalid")
            for page, tokens in zip(pdf.pages, case.pdf_page_tokens, strict=True):
                page_text = page.extract_text() or ""
                if any(token not in page_text for token in tokens):
                    raise ServiceError("acceptance_source_token_missing")
        if case.pdf_image_sources and not case.pdf_image_hashes:
            raise ServiceError("acceptance_output_invalid")
        if case.pdf_image_hashes:
            if len(pdf.pages) != len(case.pdf_image_hashes):
                raise ServiceError("acceptance_output_invalid")
            for page, expected in zip(pdf.pages, case.pdf_image_hashes, strict=True):
                images = list(page.images)
                if len(images) != 1 or pixel_signature(images[0].image) != expected:
                    raise ServiceError("acceptance_output_invalid")
        details["pages"] = len(pdf.pages)
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif suffix == ".docx":
        doc = Document(BytesIO(data))
        if case.docx_tables:
            tables = tuple(
                tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
                for table in doc.tables
            )
            if tables != case.docx_tables:
                raise ServiceError("acceptance_output_invalid")
            details["editable_tables"] = len(tables)
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for r in t.rows for c in r.cells]
        )
    elif suffix == ".xlsx":
        book = load_workbook(BytesIO(data), data_only=False)
        if not set(case.required_sheets).issubset(book.sheetnames):
            raise ServiceError("acceptance_output_invalid")
        cells = [cell for sheet in book for row in sheet for cell in row]
        if any(cell.data_type == "f" or cell.hyperlink for cell in cells):
            raise ServiceError("acceptance_output_invalid")
        text = "\n".join(str(cell.value) for cell in cells if cell.value is not None)
        details["sheets"] = book.sheetnames
    elif suffix == ".pptx":
        slides = Presentation(BytesIO(data)).slides
        if not slides:
            raise ServiceError("acceptance_output_invalid")
        if case.pptx_slides is not None and len(slides) != case.pptx_slides:
            raise ServiceError("acceptance_output_invalid")
        texts = []
        for slide in slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        text = "\n".join(texts)
        details["slides"] = len(slides)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        with Image.open(BytesIO(data)) as image:
            details["pixels"] = list(image.size)
            image.verify()
    elif suffix == ".txt":
        text = data.decode("utf-8")
    else:
        raise ServiceError("acceptance_output_invalid")
    if suffix not in {".png", ".jpg", ".jpeg", ".pdf"} and not text.strip():
        raise ServiceError("acceptance_output_invalid")
    if any(token not in text for token in case.tokens):
        raise ServiceError("acceptance_source_token_missing")
    if name in case.changed_text_files:
        # Only fixtures with known errors or an actual language change opt in.
        # Whitespace-only changes do not demonstrate editing/translation value.
        if " ".join(text.split()) == " ".join(case.inputs["text"].split()):
            raise ServiceError("acceptance_text_unchanged")
        details["text_changed"] = True
    return {"name": name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), **details}


class Delivery:
    def __init__(self, root, store, case):
        self.root, self.store, self.case = root, store, case
        self.files = []
        self.sent = 0

    async def send(self, user_id, order_id, result):
        if user_id != OWNER:
            raise ServiceError("not_allowed")
        entries = [(a.filename, self.store.read(a.key, user_id)) for a in result.artifacts]
        if result.text:
            entries.append(("response.txt", result.text.encode("utf-8")))
        if not entries or len({name for name, _ in entries}) != len(entries):
            raise ServiceError("acceptance_output_invalid")
        if self.case.expected_files is not None and set(name for name, _ in entries) != set(
            self.case.expected_files
        ):
            raise ServiceError("acceptance_output_invalid")
        checked = []
        for name, data in entries:
            if name != Path(name).name or not re.fullmatch(r"[A-Za-z0-9_.-]+", name):
                raise ServiceError("acceptance_output_invalid")
            checked.append(inspect_output(name, data, self.case))
        self.root.mkdir(parents=True, exist_ok=True)
        for name, data in entries:
            (self.root / name).write_bytes(data)
        self.files = checked
        self.sent += 1


def materialize(case, owned, data):
    saved = {}

    def walk(value):
        if isinstance(value, dict):
            return {key: walk(item) for key, item in value.items()}
        if isinstance(value, list):
            return [walk(item) for item in value]
        if isinstance(value, str) and value in ASSETS.values():
            if value not in saved:
                name = next(name for name, key in ASSETS.items() if key == value)
                pdf = name.startswith("pdf")
                artifact = owned.save(
                    "source.pdf" if pdf else "source.png",
                    data[name],
                    "application/pdf" if pdf else "image/png",
                )
                saved[value] = artifact.key
            return saved[value]
        return value

    hashes = []
    for name in case.pdf_image_sources:
        with Image.open(BytesIO(data[name])) as image:
            hashes.append(pixel_signature(image))
    return replace(case, inputs=walk(case.inputs), pdf_image_hashes=tuple(hashes))


def save_report(path, report):
    # Recompute from the complete portfolio, even for a selected or interrupted run.
    report["service_readiness"] = service_readiness(cases(ASSETS), report["cases"])
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def service_readiness(templates, results):
    """Describe this run's coverage; never infer visual or launch approval from PASS."""
    expected = {case.name: case for case in templates}
    observed = {}
    for row in results:
        name = row.get("case")
        if (
            name not in expected
            or name in observed
            or row.get("service") != expected[name].slug
            or row.get("status") not in {"PASS", "FAIL"}
        ):
            raise ServiceError("acceptance_report_invalid")
        observed[name] = row
    services = []
    for category in CATEGORY_KEYS:
        for slug in sorted({case.slug for case in templates}):
            plugin = registry.types[slug]
            if plugin.category != category:
                continue
            values = [case for case in templates if case.slug == slug]
            passed = [c.name for c in values if observed.get(c.name, {}).get("status") == "PASS"]
            failed = [c.name for c in values if observed.get(c.name, {}).get("status") == "FAIL"]
            missing = [c.name for c in values if c.name not in observed]
            status = (
                "FAIL"
                if failed
                else "NOT_RUN"
                if len(missing) == len(values)
                else "INCOMPLETE"
                if missing
                else "PASS"
            )
            services.append(
                {
                    "service": slug,
                    "name_ar": plugin.name_ar,
                    "category": category,
                    "version": plugin.version,
                    "automated_status": status,
                    "passed_cases": passed,
                    "failed_cases": failed,
                    "not_run_cases": missing,
                    "manual_review_status": "NOT_RECORDED",
                    "manual_review_checks": sorted(
                        {c for case in values for c in case.review_checks}
                    ),
                    "next_action": (
                        "fix_failed_cases"
                        if failed
                        else "run_missing_cases"
                        if missing
                        else "review_source_fidelity_and_target_viewer"
                    ),
                }
            )
    return {
        "scope": "current_run_only; configuration_and_previous_runs_not_combined",
        "launch_status": "NOT_ASSESSED",
        "services": services,
    }


async def finish(order_id):
    async with sessions.begin() as db:
        order = await lock_order(db, order_id)
        if order.status in ACTIVE:
            previous_error = await db.scalar(
                select(JobAttempt.error_key)
                .join(Job, Job.id == JobAttempt.job_id)
                .where(Job.order_id == order_id)
                .order_by(JobAttempt.number.desc())
                .limit(1)
            )
            await fail_locked(db, order, previous_error or "acceptance_incomplete")
    async with sessions() as db:
        ids = list(
            (
                await db.scalars(
                    select(CostHold.id)
                    .join(Job, Job.id == CostHold.job_id)
                    .where(Job.order_id == order_id, CostHold.status == "open")
                )
            ).all()
        )
    for hold in ids:
        await uncertain_call(hold)


async def run_case(case, adapter, backend, output):
    oid = None
    error = None
    approvals = 0
    diagnostics = []
    delivery = Delivery(output / case.name, backend, case)
    before = adapter.calls
    try:
        async with sessions() as db:
            row = await db.get(Service, case.slug)
            price, version = row.price_halala, row.version
        oid = await submit(OWNER, case.slug, case.inputs, price, "acceptance:" + case.name, version)
        async with sessions() as db:
            jid = await db.scalar(select(Job.id).where(Job.order_id == oid))

        def runtime(job_id, user_id):
            if job_id != jid or user_id != OWNER:
                raise ServiceError("not_allowed")
            value = runtime_for(
                job_id, user_id, needs_ai=registry.types[case.slug].needs_ai(case.inputs)
            )
            if value.ai is not None:
                value.ai = DiagnosticAI(
                    adapter,
                    job_id,
                    user_id,
                    quality_review=True,
                    diagnostics=diagnostics,
                    pptx_source=case.inputs["text"] if case.slug == "text_to_pptx" else None,
                )
            return value

        ctx = {"delivery": delivery, "runtime_factory": runtime}
        for _ in range(3):
            await execute_job(ctx, str(jid))
            async with sessions() as db:
                order = await db.get(Order, oid)
                status, prepared = order.status, bool((order.result or {}).get("prepared_delivery"))
            if status != "waiting_confirmation":
                break
            approvals += 1
            await confirm_structure(oid, OWNER, True, expected_prepared=prepared)
        if status == "completed":
            count, sends = adapter.calls, delivery.sent
            await execute_job(ctx, str(jid))
            if adapter.calls != count or delivery.sent != sends:
                raise ServiceError("acceptance_replay_invalid")
    except Exception as exc:
        error = exc.key if isinstance(exc, ServiceError) else "acceptance_runtime_failed"
    finally:
        if oid is None:
            async with sessions() as db:
                oid = await db.scalar(
                    select(Order.id).where(
                        Order.user_id == OWNER, Order.idempotency_key == "acceptance:" + case.name
                    )
                )
        if oid:
            await finish(oid)
    async with sessions() as db:
        order = await db.get(Order, oid) if oid else None
        holds = (
            list(
                (
                    await db.scalars(
                        select(CostHold)
                        .join(Job, Job.id == CostHold.job_id)
                        .where(Job.order_id == oid)
                    )
                ).all()
            )
            if oid
            else []
        )
        usages = (
            list(
                (
                    await db.scalars(
                        select(CostUsage)
                        .join(Job, Job.id == CostUsage.job_id)
                        .where(Job.order_id == oid)
                    )
                ).all()
            )
            if oid
            else []
        )
        captures = (
            await db.scalar(
                select(func.count(Ledger.id)).where(
                    Ledger.order_id == oid, Ledger.kind == "capture"
                )
            )
            if oid
            else 0
        )
        success = bool(
            order
            and order.status == "completed"
            and delivery.sent == 1
            and captures == 1
            and not error
        )
        return {
            "case": case.name,
            "service": case.slug,
            "version": registry.types[case.slug].version,
            "ai_needed": registry.types[case.slug].needs_ai(case.inputs),
            "status": "PASS" if success else "FAIL",
            "order_status": order.status if order else None,
            "error": error or (order.error_key if order else None),
            "request": str(oid) if oid else None,
            "provider_calls": adapter.calls - before,
            "diagnostics": diagnostics,
            "approvals": approvals,
            "artifacts": delivery.files,
            "manual_review_checks": list(case.review_checks),
            "usage": [
                {
                    "calculated_sar": str(u.cost_sar),
                    "input_tokens": u.input_tokens,
                    "output_tokens": u.output_tokens,
                    "audit": u.rate_snapshot,
                }
                for u in usages
            ],
            "holds": [{"status": h.status, "reserved_sar": str(h.amount)} for h in holds],
        }


async def execute(case_names=None, *, native_only=False):
    root = isolated()
    templates = (native_cases if native_only else selected_cases)(cases(ASSETS), case_names)
    output = root / "results"
    if output.exists():
        raise ServiceError("acceptance_fresh_run_required")
    adapter = NativeAcceptanceProvider() if native_only else AcceptanceProvider()
    adapter.rates()
    async with sessions() as db:
        if any([await db.scalar(select(func.count(cls.id))) for cls in (User, Order, CostHold)]):
            raise ServiceError("acceptance_fresh_database_required")
    output.mkdir(parents=True)
    async with sessions.begin() as db:
        await registry.sync(db)
        for slug in {case.slug for case in templates}:
            (await db.get(Service, slug)).enabled = True
    await register(OWNER)
    async with sessions.begin() as db:
        prices = {row.slug: row.price_halala for row in (await db.scalars(select(Service))).all()}
        await apply(
            db, OWNER, "credit", sum(prices[case.slug] for case in templates), "acceptance-seed"
        )
    backend = storage()
    owned = OwnedStorage(backend, OWNER)
    pdf = source_pdf()
    data = {
        "pdf": pdf,
        "pdf2": source_pdf(second=True),
        "image": picture(),
        "shaded": picture(shaded=True),
        "mixed": picture(mixed=True),
    }
    report = {
        "status": "RUNNING",
        "native_only": native_only,
        "cases": [],
        "selected_cases": [case.name for case in templates],
        "services": len({case.slug for case in templates}),
        "host": urlsplit(config().ai_base_url).hostname,
        "model": config().ai_model,
        "max_call_bounds_sar": str(LIMIT),
        "max_calls": CALLS,
        "qualification": "synthetic_native_and_live_provider_acceptance; not_visual_mobile_semantic_or_invoice_certification",
    }
    path = output / "report.json"
    save_report(path, report)
    for template in templates:
        try:
            case = materialize(template, owned, data)
            row = await run_case(case, adapter, backend, output)
        except Exception as exc:
            row = {
                "case": template.name,
                "service": template.slug,
                "status": "FAIL",
                "error": exc.key if isinstance(exc, ServiceError) else "acceptance_runtime_failed",
            }
        report["cases"].append(row)
        save_report(path, report)
        console = {key: row.get(key) for key in ("case", "status", "error", "provider_calls")}
        if row["status"] == "FAIL" and row.get("diagnostics"):
            console["diagnostics"] = row["diagnostics"]
        print(json.dumps(console), flush=True)
    async with sessions() as db:
        funds = await balance(db, OWNER)
        active = await db.scalar(select(func.count(Order.id)).where(Order.status.in_(ACTIVE)))
        pause = await db.get(Setting, "PROVIDER_PAUSED")
        report.update(
            provider_calls=adapter.calls,
            conservative_call_bounds_sar=str(adapter.bounds),
            wallet_reserved_halala=funds.reserved,
            active_orders=active,
            test_provider_paused=bool(pause and pause.value.get("value")),
        )
    report["passed"] = sum(row["status"] == "PASS" for row in report["cases"])
    report["calculated_usage_sar"] = str(
        sum(
            (Decimal(u["calculated_sar"]) for row in report["cases"] for u in row.get("usage", [])),
            Decimal("0"),
        )
    )
    report["failed"] = len(templates) - report["passed"]
    report["status"] = (
        "PASS" if not report["failed"] and not active and funds.reserved == 0 else "REVIEW_REQUIRED"
    )
    save_report(path, report)
    print(
        json.dumps(
            {
                key: report[key]
                for key in (
                    "status",
                    "services",
                    "passed",
                    "failed",
                    "provider_calls",
                    "conservative_call_bounds_sar",
                    "calculated_usage_sar",
                    "active_orders",
                    "wallet_reserved_halala",
                    "service_readiness",
                )
            }
        )
    )
    return 0 if report["status"] == "PASS" else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--native-only", action="store_true")
    parser.add_argument("--case", action="append", dest="case_names")
    args = parser.parse_args()
    try:
        values = (native_cases if args.native_only else selected_cases)(
            cases(ASSETS), args.case_names
        )
    except ServiceError as error:
        print(json.dumps({"status": "BLOCKED", "error": error.key}))
        return 1
    if not args.execute:
        print(
            json.dumps(
                {
                    "status": "plan",
                    "services": len({c.slug for c in values}),
                    "cases": [c.name for c in values],
                    "max_call_bounds_sar": str(LIMIT),
                    "max_calls": CALLS,
                    "live_execution": "not_performed",
                    "service_readiness": service_readiness(cases(ASSETS), []),
                }
            )
        )
        return 0
    try:
        return asyncio.run(execute(args.case_names, native_only=args.native_only))
    except ServiceError as error:
        print(json.dumps({"status": "BLOCKED", "error": error.key}))
        return 1
    except Exception:
        print(json.dumps({"status": "BLOCKED", "error": "acceptance_runtime_failed"}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
