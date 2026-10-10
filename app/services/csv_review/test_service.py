from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.builders.csv_review import build
from app.core.i18n import CATALOGS
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.csv_review.schema import Inputs
from app.services.csv_review.service import SERVICE
from tests.csv_review_fixtures import inputs


def exported(values):
    parsed = Inputs.parse(values)
    headers, rows = parsed.records()
    data, counts = build(headers, rows, trim=parsed.whitespace == "trim", language=parsed.language)
    return load_workbook(BytesIO(data)), counts, data


@pytest.mark.parametrize("language", ["ar", "en"])
def test_source_values_counts_native_tables_and_row_references(language):
    book, counts, _ = exported(inputs(language))
    assert counts == {"rows": 4, "blanks": 2, "duplicates": 1, "changes": 1}
    assert book.sheetnames == ["Data", "Source", "Review"]
    assert book["Data"]["A4"].value == "00123"
    assert book["Data"]["C4"].value == "125.50"
    assert book["Data"]["C5"].value == "0" and book["Data"]["B5"].value is None
    assert book["Source"]["B4"].value == " أحمد " and book["Data"]["B4"].value == "أحمد"
    assert book["Data"]["D7"].value == "08/10/2026"
    assert [c.value for c in book["Review"][4]] == [1, None, None, "2"]
    assert [c.value for c in book["Review"][5]] == [2, "2, 4", None, None]
    assert [c.value for c in book["Review"][6]] == [3, None, 1, None]
    for name in book.sheetnames:
        sheet = book[name]
        assert sheet.max_row == 7 and sheet.freeze_panes == "A4"
        assert len(sheet.tables) == 1
        assert next(iter(sheet.tables.values())).ref == ("A3:D7" if name == "Review" else "A3:E7")
        assert sheet.sheet_view.rightToLeft == (language == "ar")
        assert sheet.print_title_rows == "$1:$3"
        assert all(cell.alignment.wrap_text for row in sheet.iter_rows(min_row=4) for cell in row)
    assert book["Data"]["C4"].number_format == "@"


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", " ="])
def test_formula_like_values_and_headers_survive_as_exact_text(prefix):
    value = prefix + 'HYPERLINK("https://example.invalid","x")'
    # Quote the comma-containing cell using actual CSV escaping.
    import csv
    from io import StringIO

    text = StringIO()
    writer = csv.writer(text)
    writer.writerow([prefix + "Header"])
    writer.writerow([value])
    book, _, content = exported({**inputs(), "text": text.getvalue()})
    for name in ("Data", "Source"):
        assert (
            book[name]["A4"].value == value.strip(" \t")
            if name == "Data"
            else book[name]["A4"].value == value
        )
        assert book[name]["A4"].data_type == "s"
    assert not any(
        c.data_type == "f" or c.hyperlink for sheet in book for row in sheet for c in row
    )
    with ZipFile(BytesIO(content)) as package:
        assert not any(
            "externallinks" in name.lower() or "vba" in name.lower() for name in package.namelist()
        )
        assert all(
            b"<f>" not in package.read(name)
            for name in package.namelist()
            if name.startswith("xl/worksheets/")
        )
    saved = BytesIO()
    book.save(saved)
    assert load_workbook(BytesIO(saved.getvalue()))["Source"]["A4"].value == value


def test_preserve_mode_retains_spaces_duplicates_and_all_empty_records():
    book, counts, _ = exported(
        {**inputs(whitespace="preserve"), "text": " ID ,Note\n00123, x \n00123,x\n,\n\n"}
    )
    assert counts == {"rows": 4, "blanks": 4, "duplicates": 1, "changes": 0}
    assert book["Source"]["A3"].value == " ID " and book["Data"]["A3"].value == "ID"
    assert book["Source"]["A4"].value == "00123" and book["Data"]["B4"].value == " x "
    assert book["Review"]["C7"].value == 3
    assert book["Data"].max_row == 7


@pytest.mark.parametrize(
    "delimiter, text",
    [
        ("semicolon", 'A;B\n00123;"x;y"'),
        ("tab", "A\tB\n00123\t0"),
        ("comma", '\ufeffA,B\r\n00123,"first\r\nsecond"'),
    ],
)
def test_explicit_delimiters_quoted_cells_bom_and_multiline(delimiter, text):
    book, _, _ = exported({**inputs(), "text": text, "delimiter": delimiter})
    assert book["Source"]["A4"].value == "00123"
    assert book["Source"]["B4"].value in {"x;y", "0", "first\nsecond"}


@pytest.mark.parametrize(
    "text",
    [
        "",
        "A,B",
        "A,B\n1",
        "A,B\n1,2,3",
        'A\n"unfinished',
        "A,a\n1,2",
        " ,B\n1,2",
        "A\nvalue\x00",
        "A\n" + "x" * 501,
        "A\n" + "x\n" * 501,
        ",".join(str(i) for i in range(21)) + "\n" + "," * 20,
        'A\n"' + "x\n" * 11 + '"',
    ],
)
def test_invalid_or_oversized_source_rejected_during_admission(text):
    with pytest.raises(ServiceError, match="input_invalid"):
        SERVICE.needs_ai({**inputs(), "text": text})


def test_max_records_and_exact_utf8_text_limits():
    book, counts, _ = exported({**inputs(), "text": "ID\n" + "00017\n" * 500})
    assert counts["rows"] == 500 and counts["duplicates"] == 499
    assert book["Data"].max_row == 503 and book["Source"]["A503"].value == "00017"
    with pytest.raises(ServiceError):
        Inputs.parse({**inputs(), "text": "A\n" + "x" * 11999})


async def test_service_owns_artifact_without_ai_and_catalogs_cover_intake(tmp_path):
    assert not SERVICE.enabled_by_default and not SERVICE.requires_ai
    assert SERVICE.needs_ai(inputs()) is False
    service = SERVICE()
    storage = LocalStorage(tmp_path)
    service.runtime = SimpleNamespace(ai=None, storage=OwnedStorage(storage, 1))
    result = await service.run(inputs())
    assert len(result.artifacts) == 1 and result.artifacts[0].filename == "review.xlsx"
    assert "4" in result.preview and not result.needs_confirmation
    assert "4" in result.preview_localizations["ar"] and "4" in result.preview_localizations["en"]
    assert (
        load_workbook(BytesIO(storage.read(result.artifacts[0].key, 1)))["Data"]["A4"].value
        == "00123"
    )
    with pytest.raises(ServiceError, match="not_allowed"):
        storage.read(result.artifacts[0].key, 2)
    for field in SERVICE.input_schema.fields:
        for key in [field.prompt_key, *field.choice_keys]:
            assert all(key in CATALOGS[lang] for lang in ("ar", "en"))
    with pytest.raises(ServiceError):
        SERVICE.needs_ai({**inputs(), "__continuation": {}})


def test_stable_record_reference_survives_reordering_and_heading_collision():
    book, _, _ = exported({**inputs("en"), "text": "Source record,Name\n00123, A \n00123,A"})
    sheet = book["Data"]
    assert sheet["C3"].value == "Source record (2)"
    assert sheet["C4"].value == "1" and sheet["C5"].value == "2"
    # The added reference moves with its record instead of depending on row position.
    reversed_rows = list(reversed(list(sheet.iter_rows(min_row=4, values_only=True))))
    assert reversed_rows[0][-1] == "2" and reversed_rows[1][-1] == "1"
    assert book["Review"]["C5"].value == 1
    assert book["Source"]["C4"].value == "1"


def test_trim_changes_only_ascii_edge_spaces_and_tabs_and_keeps_unicode_and_newlines():
    text = 'A,B\n" \t00123\t ","first\nsecond"\n\u00a0,0'
    book, counts, _ = exported({**inputs(), "text": text})
    assert counts == {"rows": 2, "blanks": 0, "duplicates": 0, "changes": 1}
    assert book["Data"]["A4"].value == "00123"
    assert book["Source"]["A4"].value == " \t00123\t "
    assert book["Data"]["B4"].value == "first\nsecond"
    assert book["Data"]["A5"].value == "\u00a0"


@pytest.mark.parametrize("header", [" " * 81 + "A", "\u00a0", "A,\u00a0A", '"A\u2028B"'])
def test_headers_reject_oversized_blank_normalized_collisions_and_line_breaks(header):
    with pytest.raises(ServiceError, match="input_invalid"):
        Inputs.parse({**inputs(), "text": header + "\n" + ("1,2" if "," in header else "1")})
