from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.builders.csv_review import build
from app.builders.csv_similarity import review
from app.services.base import ServiceError
from app.services.csv_review.schema import Inputs
from app.services.csv_review.service import SERVICE
from tests.csv_review_fixtures import inputs as legacy_inputs
from tests.csv_similarity_fixtures import inputs


def exported(values):
    parsed = Inputs.parse(values)
    headers, rows = parsed.records()
    data, counts = build(
        headers,
        rows,
        trim=parsed.whitespace == "trim",
        language=parsed.language,
        similarity=parsed.similarity == "conservative",
    )
    return load_workbook(BytesIO(data)), counts


@pytest.mark.parametrize("language", ["ar", "en"])
def test_report_candidates_have_source_references_and_preserve_all_cells(language):
    values = inputs(language)
    book, counts = exported(values)
    assert counts == {
        "rows": 5,
        "blanks": 1,
        "duplicates": 1,
        "changes": 0,
        "similar_pairs": 2,
        "shown_pairs": 2,
        "omitted_pairs": 0,
    }
    assert book.sheetnames == ["Data", "Source", "Review", "Similarities"]
    sheet = book["Similarities"]
    candidates = list(sheet.iter_rows(min_row=4, values_only=True))
    assert [(row[0], row[1], row[3]) for row in candidates] == [(1, 2, "2"), (2, 4, "2")]
    assert all(90 <= row[2] < 100 for row in candidates)
    assert sheet.tables["SimilaritiesRecords"].ref == "A3:D5"
    assert sheet.sheet_view.rightToLeft == (language == "ar")
    assert book["Review"]["C7"].value == 1
    _, original = Inputs.parse(values).records()
    for name in ("Data", "Source"):
        actual = list(book[name].iter_rows(min_row=4, values_only=True))
        assert [
            ["" if value is None else value for value in row[:-1]] for row in actual
        ] == original
        assert [row[-1] for row in actual] == [str(i) for i in range(1, 6)]
        assert book[name]["C4"].value == "=1+1" and book[name]["C4"].data_type == "s"
    assert not any(
        cell.data_type == "f" or cell.hyperlink for sheet in book for row in sheet for cell in row
    )


@pytest.mark.parametrize(
    "a,b",
    [
        ("Invoice 00123", "Invoice 00124"),
        ("فاتورة ٠٠١٢٣", "فاتورة ٠٠١٢٤"),
        ("125.50", "125.51"),
        ("Date 2026-10-09", "Date 2026-10-08"),
        ("Account １２３", "Account １２４"),
        ("Record ²", "Record ³"),
        ("", "Southern Office Services"),
        (" \t", "Southern Office Services"),
    ],
)
def test_changed_numbers_and_blank_values_never_become_candidates(a, b):
    result = review([["Shared long company name", a], ["Shared long company name", b]])
    assert result.total == 0 and not result.pairs


def test_unchanged_common_cells_cannot_mask_a_disagreeing_column():
    common = ["Southern Office Services"] * 19
    assert review([common + ["accepted"], common + ["rejected"]]).total == 0


def test_column_boundaries_are_not_lost_by_concatenation():
    assert review([["abcdefghijk", "x"], ["abcdefghij", "kx"]]).total == 0


def test_exact_duplicates_are_only_in_existing_equality_review():
    assert review([["Same", "00123"], ["Same", "00123"]]).total == 0


def test_ascii_edges_only_follow_existing_selected_rule():
    text = "Name\n Southern Office Services \nSouthern Office Services"
    preserved, preserved_counts = exported({**inputs(), "text": text})
    trimmed, trimmed_counts = exported({**inputs(), "text": text, "whitespace": "trim"})
    assert preserved_counts["similar_pairs"] == 1 and preserved_counts["duplicates"] == 0
    assert trimmed_counts["similar_pairs"] == 0 and trimmed_counts["duplicates"] == 1
    assert (
        preserved["Source"]["A4"].value
        == trimmed["Source"]["A4"].value
        == " Southern Office Services "
    )
    assert preserved["Data"]["A4"].value == " Southern Office Services "
    assert trimmed["Data"]["A4"].value == "Southern Office Services"


@pytest.mark.parametrize("value", ["ai", "merge", "off", True])
def test_unknown_or_non_string_modes_rejected_before_reservation(value):
    with pytest.raises(ServiceError, match="input_invalid"):
        SERVICE.needs_ai({**inputs(), "similarity": value})


def test_legacy_inputs_default_off_and_allow_500_records():
    values = {**legacy_inputs(), "text": "Name\n" + "Name\n" * 500}
    assert Inputs.parse(values).similarity is None
    SERVICE.input_schema.validate_inputs(values)
    book, counts = exported(values)
    assert book.sheetnames == ["Data", "Source", "Review"]
    assert counts["rows"] == 500 and "similar_pairs" not in counts
    with pytest.raises(ServiceError, match="input_invalid"):
        SERVICE.needs_ai({**values, "similarity": "conservative"})


def test_pair_cap_reports_omissions_in_stable_source_order():
    rows = [
        ["Southern Office Services" if i % 2 == 0 else "Southern Office Service"]
        for i in range(200)
    ]
    result = review(rows)
    assert result.total == 10000 and len(result.pairs) == 1000
    assert result.pairs[0].first == 1 and result.pairs[0].second == 2
    keys = [(pair.second, pair.first) for pair in result.pairs]
    assert keys == sorted(keys)
    values = {**inputs(), "text": "Name\n" + "\n".join(row[0] for row in rows)}
    book, counts = exported(values)
    assert counts["similar_pairs"] == 10000 and counts["omitted_pairs"] == 9000
    assert book["Data"].max_row == 203 and book["Similarities"].max_row == 1003
    assert "9000" in book["Similarities"]["A2"].value


def test_no_pairs_report_is_explicit_and_does_not_add_data_records():
    book, counts = exported({**inputs(), "text": "Name\nSouthern Office Services"})
    assert counts["similar_pairs"] == 0
    assert book["Data"].max_row == 4
    assert list(book["Similarities"].iter_rows(min_row=4, values_only=True)) == [
        (None, None, None, None)
    ]
    assert "Zero pairs" in book["Similarities"]["A2"].value
