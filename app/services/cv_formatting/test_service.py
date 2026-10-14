import re
from io import BytesIO
from types import SimpleNamespace

import pdfplumber
import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.builders.word import build
from app.providers.documents.base import DocumentError
from app.providers.documents.rendering import WordFiles
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.cv_formatting.schema import Inputs, document, revision_schema
from app.services.cv_formatting.service import CvFormatting
from tests.cv_fixtures import inputs, planned, revision
from tests.word_pdf_fixtures import renderer


@pytest.mark.parametrize("language", ["ar", "en"])
def test_cv_retains_supplied_lines_escapes_templates_and_drafts_missing_summary(language):
    source = inputs(language)
    source.pop("summary")
    source["skills"] = "  Excel  \n\n{{ 7 * 7 }} <tag> & {% if True %}"
    CvFormatting.input_schema.validate_inputs(source)
    assert CvFormatting.needs_ai(source)
    plan = planned(source)
    doc = Document(BytesIO(build(plan, professional_template=True, format_dates=True)))
    assert not doc.tables and doc.paragraphs[0].text == source["full_name"]
    content = [p.text.replace("\u200e", "") for p in doc.paragraphs]
    for name in ("contact", "experience", "education", "skills", "additional"):
        for line in source[name].split("\n"):
            if line.strip():
                assert line.strip() in content
    assert ("نبذة مهنية" if language == "ar" else "Professional summary") in content
    assert doc.paragraphs[1].text.startswith("0500123456")


@pytest.mark.parametrize(
    "patch",
    [
        {"full_name": "\nInvalid"},
        {"contact": "\x00"},
        {"skills": "   "},
        {"experience": None, "education": None},
        {"language": "unknown"},
        {"summary": "x" * 1001},
        {"skills": "line\n" * 101},
        {"unexpected": "ignored?"},
        {"text": "not CV input"},
        {
            "experience": "x" * 4000,
            "education": "x" * 2000,
            "skills": "x" * 1500,
            "additional": "x" * 1500,
            "summary": "x" * 1000,
        },
    ],
)
def test_invalid_cv_is_rejected_before_reservation(patch):
    with pytest.raises(ServiceError, match="input_invalid"):
        CvFormatting.needs_ai(inputs() | patch)


def test_entry_level_cv_requires_no_invented_experience():
    source = inputs()
    source.pop("experience")
    plan = planned(source)
    assert all(section.heading != "الخبرات العملية" for section in plan.sections)
    assert any(section.heading == "التعليم" for section in plan.sections)


@pytest.mark.parametrize("language", ["ar", "en"])
async def test_actual_cv_native_pdf_has_readable_bounds_and_preserves_contacts_dates(language):
    result = await renderer().word_pdf(planned(inputs(language)))
    pages = PdfReader(BytesIO(result.pdf)).pages
    assert len(pages) == 1
    text = pages[0].extract_text()
    for token in (
        "0500123456",
        "sample@example.invalid",
        "2024-01-09",
        "2026-10-09",
        "00123",
        "125.50",
        "00017",
    ):
        assert len(re.findall(r"(?<!\d)" + re.escape(token) + r"(?!\d)", text)) == 1
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        assert all(
            0 <= char["x0"] < char["x1"] <= pdf.pages[0].width for char in pdf.pages[0].chars
        )


@pytest.mark.parametrize("language", ["ar", "en"])
def test_cv_document_edge_keeps_numeric_and_arabic_script_directions(language):
    plan = planned(inputs(language))
    doc = Document(BytesIO(build(plan, professional_template=True)))
    contact = doc.paragraphs[1]._p.pPr
    assert contact.find(qn("w:bidi")).get(qn("w:val")) == "0"
    assert contact.find(qn("w:jc")).get(qn("w:val")) == "center"
    title = doc.paragraphs[0]._p.pPr
    assert title.find(qn("w:bidi")).get(qn("w:val")) == ("1" if language == "ar" else "0")
    assert title.find(qn("w:jc")).get(qn("w:val")) == "center"


async def test_long_cv_retains_every_record_and_wraps_across_pages():
    source = inputs("en")
    source["experience"] = "\n".join(
        f"Record {index:05d} | milestone 2026-10-09" for index in range(100, 190)
    )
    result = await renderer().word_pdf(planned(source))
    pages = PdfReader(BytesIO(result.pdf)).pages
    assert 2 <= len(pages) <= 5
    text = "\n".join(page.extract_text() for page in pages)
    assert all(
        len(re.findall(r"(?<!\d)" + f"{index:05d}" + r"(?!\d)", text)) == 1
        for index in range(100, 190)
    )
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        for page in pdf.pages:
            assert all(
                0 <= char["x0"] < char["x1"] <= page.width
                and 0 <= char["top"] < char["bottom"] <= page.height
                for char in page.chars
            )


async def test_cv_uses_owned_storage_and_partial_save_cleanup(tmp_path, monkeypatch):
    class Renderer:
        async def word_pdf(self, document, **kwargs):
            return WordFiles(b"word", b"pdf")

    storage = OwnedStorage(LocalStorage(tmp_path), 1)
    service = CvFormatting()
    service.runtime = SimpleNamespace(storage=storage, renderer=Renderer(), ai=fixture_ai(inputs()))
    result = await service.run(inputs())
    assert [artifact.filename for artifact in result.artifacts] == ["cv.docx", "cv.pdf"]
    for artifact in result.artifacts:
        with pytest.raises(ServiceError, match="not_allowed"):
            LocalStorage(tmp_path).read(artifact.key, 2)
        storage.delete(artifact.key)
    original = storage.save

    def save(filename, *args):
        if filename.endswith(".pdf"):
            raise ServiceError("storage_quota")
        return original(filename, *args)

    monkeypatch.setattr(storage, "save", save)
    with pytest.raises(ServiceError, match="storage_quota"):
        await service.run(inputs())
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_cv_render_failure_does_not_save_files(tmp_path):
    class Renderer:
        async def word_pdf(self, *args, **kwargs):
            raise DocumentError("document_render_timeout")

    service = CvFormatting()
    service.runtime = SimpleNamespace(
        storage=OwnedStorage(LocalStorage(tmp_path), 1),
        renderer=Renderer(),
        ai=fixture_ai(inputs()),
    )
    with pytest.raises(ServiceError, match="document_render_timeout"):
        await service.run(inputs())
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize(
    "kind,first",
    [
        ("professional", "experience"),
        ("graduate", "education"),
        ("career_change", "skills"),
        ("academic", "education"),
    ],
)
@pytest.mark.parametrize("language", ["ar", "en"])
def test_cv_profiles_preserve_fields_and_order_sections(kind, first, language):
    from app.core.i18n import tr

    source = inputs(language) | {
        "cv_type": kind,
        "headline": "Operations 00123",
        "projects": "Project 00001",
        "languages": "Arabic | English",
        "research": "Research 00002",
        "publications": "Publication 00003 <tag> {{ 7 * 7 }}",
    }
    CvFormatting.input_schema.validate_inputs(source)
    assert CvFormatting.needs_ai(source)
    plan = planned(source)
    assert plan.layout == "resume"
    assert plan.sections[2].heading == tr("cv_" + first + "_heading", language)
    doc = Document(BytesIO(build(plan, professional_template=True, format_dates=True)))
    paragraphs = [p.text.replace("\u200e", "") for p in doc.paragraphs]
    for name in (
        "full_name",
        "headline",
        "contact",
        "summary",
        "experience",
        "education",
        "skills",
        "additional",
        "projects",
        "languages",
        "research",
        "publications",
    ):
        assert all(line in paragraphs for line in source[name].split("\n"))
    assert not doc.tables
    assert doc.styles["Heading 1"].font.size.pt == 12
    assert doc.styles["Normal"].font.size.pt == 11


@pytest.mark.parametrize(
    "patch",
    [
        {"cv_type": "ats_guaranteed"},
        {"headline": "two\nlines"},
        {"publications": "x" * 2001},
        {"research": "x\n" * 101},
    ],
)
def test_cv_profile_rejects_unknown_types_and_unbounded_fields(patch):
    with pytest.raises(ServiceError, match="input_invalid"):
        CvFormatting.needs_ai(inputs() | patch)


def test_resume_layout_is_opt_in_and_standard_word_stays_unchanged():
    plan = planned(inputs())
    resume = Document(BytesIO(build(plan, professional_template=True)))
    standard = Document(
        BytesIO(
            build(
                plan.model_copy(update={"layout": "standard", "design": None}),
                professional_template=True,
            )
        )
    )
    assert standard.styles["Heading 1"].font.size.pt == 14
    assert resume.styles["Heading 1"].font.size.pt == 12
    assert [p.text for p in resume.paragraphs] == [p.text for p in standard.paragraphs]


def fixture_ai(source):
    async def extract(schema, text, prompt):
        assert text == Inputs.parse(source).source()
        assert "professional CV" in prompt
        return schema.model_validate(revision(source))

    return SimpleNamespace(extract=extract)


def bilingual_inputs(kind="professional"):
    return inputs("ar") | {"language": "bilingual", "cv_type": kind}


@pytest.mark.parametrize("kind", ["professional", "graduate", "career_change", "academic"])
def test_bilingual_cv_uses_one_input_and_complete_language_bodies(kind):
    source = bilingual_inputs(kind)
    CvFormatting.input_schema.validate_inputs(source)
    assert CvFormatting.needs_ai(source)
    plan = planned(source)
    assert plan.design.theme == "modern" and plan.alignment == "auto"
    titles = [s for s in plan.sections if s.title]
    assert len(titles) == 1 and titles[0].title == "Sam Example" and titles[0].page_break_before
    doc = Document(BytesIO(build(plan, professional_template=True, format_dates=True)))
    lines = [p.text.replace("\u200e", "") for p in doc.paragraphs]
    split = lines.index("Sam Example")
    for language, body_lines in (("ar", lines[:split]), ("en", lines[split:])):
        for name in (
            "full_name",
            "contact",
            "summary",
            "experience",
            "education",
            "skills",
            "additional",
        ):
            assert all(line in body_lines for line in inputs(language)[name].split("\n"))
    assert doc.paragraphs[split].paragraph_format.page_break_before
    assert not doc.tables


@pytest.mark.parametrize("language", ["ar", "en", "bilingual"])
def test_one_intake_contains_no_second_language_form(language):
    source = inputs() | {"language": language}
    CvFormatting.input_schema.validate_inputs(source)
    assert CvFormatting.needs_ai(source)
    fields = CvFormatting.input_schema.conversation()
    assert len(fields) == 14 and not any(f.name.startswith("english") for f in fields)
    assert next(f for f in fields if f.name == "education").required
    assert not next(f for f in fields if f.name == "experience").required


def test_old_duplicate_language_body_is_rejected_in_new_contract():
    with pytest.raises(ServiceError, match="input_invalid"):
        CvFormatting.needs_ai(bilingual_inputs() | {"english": inputs("en")})


@pytest.mark.parametrize("language", ["ar", "en", "bilingual"])
def test_education_required_even_with_experience_and_graduate_can_skip_experience(language):
    source = inputs() | {"language": language}
    source.pop("experience")
    CvFormatting.input_schema.validate_inputs(source)
    assert CvFormatting.needs_ai(source)
    assert not any(
        s.heading in {"الخبرات العملية", "Work experience"} for s in planned(source).sections
    )
    source = inputs() | {"language": language}
    source.pop("education")
    with pytest.raises(ServiceError, match="input_invalid"):
        CvFormatting.input_schema.validate_inputs(source)
    with pytest.raises(ServiceError, match="input_invalid"):
        CvFormatting.needs_ai(source)


@pytest.mark.parametrize(
    "defect",
    [
        "missing_id",
        "duplicate_id",
        "unknown_id",
        "number",
        "email",
        "url",
        "phone",
        "name",
        "language",
        "extra_language",
        "section",
        "duplicate_section",
        "code",
        "empty",
        "newline",
        "summary_facts",
    ],
)
def test_model_output_cannot_bypass_source_and_design_contract(defect):
    source = bilingual_inputs()
    source["contact"] += "\n+966 50 012 3456"
    response = revision(source)
    body = response["bodies"][0]
    if defect == "missing_id":
        body["segments"].pop()
    elif defect == "duplicate_id":
        body["segments"].append(body["segments"][0])
    elif defect == "unknown_id":
        body["segments"][0]["id"] = "invented:1"
    elif defect == "number":
        next(s for s in body["segments"] if "00123" in s["text"])["text"] = "Record 123"
    elif defect == "email":
        body["segments"][0]["text"] = body["segments"][0]["text"].replace(
            "sample@example.invalid", "changed@example.invalid"
        )
    elif defect == "url":
        body["segments"][1]["text"] = "أبها | changed.example.invalid"
    elif defect == "phone":
        body["segments"][2]["text"] = "+966 50012 3456"
    elif defect == "name":
        body["full_name"] += " إضافي"
    elif defect == "language":
        body["language"] = "en"
    elif defect == "extra_language":
        response["bodies"].append(body)
    elif defect == "section":
        response["section_order"].remove("education")
    elif defect == "duplicate_section":
        response["section_order"].append("skills")
    elif defect == "code":
        response["design"]["theme"] = "<script>"
    elif defect == "empty":
        body["segments"][0]["text"] = " "
    elif defect == "newline":
        body["segments"][0]["text"] += "\nnew"
    else:
        body["summary"] = {"text": "Expert with 10 years", "source_ids": ["skills:1"]}
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        revision_schema(Inputs.parse(source)).model_validate(response)


@pytest.mark.parametrize(
    "defect",
    ["no_summary", "unknown_ref", "contact_ref", "duplicates", "new_number", "new_contact"],
)
def test_generated_summary_requires_known_career_facts(defect):
    source = inputs()
    source.pop("summary")
    response = revision(source)
    summary = response["bodies"][0]["summary"]
    if defect == "no_summary":
        response["bodies"][0]["summary"] = None
    elif defect == "unknown_ref":
        summary["source_ids"] = ["invented:1"]
    elif defect == "contact_ref":
        summary["source_ids"] = ["contact:1"]
    elif defect == "duplicates":
        summary["source_ids"] = ["skills:1", "skills:1"]
    elif defect == "new_number":
        summary["text"] += " 10"
    else:
        summary["text"] += " new@example.invalid"
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        revision_schema(Inputs.parse(source)).model_validate(response)


def test_revalidation_rejects_forged_model_copy_before_building():
    source = inputs()
    values = Inputs.parse(source)
    model = revision_schema(values).model_validate(revision(source))
    bad = model.model_copy(update={"section_order": ["skills"]})
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        document(values, bad)


@pytest.mark.parametrize("theme", ["modern", "classic", "editorial"])
@pytest.mark.parametrize("density", ["comfortable", "compact"])
def test_named_designs_keep_readable_editable_content_and_escape_templates(theme, density):
    source = inputs() | {"headline": "منسق عمليات", "skills": "{{ 7 * 7 }} <tag> & {% if True %}"}
    doc = Document(
        BytesIO(build(planned(source, theme=theme, density=density), professional_template=True))
    )
    assert source["skills"] in [p.text for p in doc.paragraphs]
    assert doc.styles["Normal"].font.size.pt == 11
    assert doc.styles["Title"].font.size.pt in (24, 26, 28)
    assert not doc.tables and not doc.sections[0].footer.paragraphs[0].text
    assert "PAGE" not in doc.sections[0].footer.paragraphs[0]._p.xml
    assert "w:pBdr" not in doc.styles["Title"].element.xml


async def test_bilingual_native_pdf_starts_english_on_second_page():
    rendered = await renderer().word_pdf(planned(bilingual_inputs()))
    pages = PdfReader(BytesIO(rendered.pdf)).pages
    assert len(pages) == 2
    assert "Sam Example" not in pages[0].extract_text() and "Sam Example" in pages[1].extract_text()
    for page in pages:
        assert page.extract_text().count("0500123456") == 1
        assert page.extract_text().count("sample@example.invalid") == 1
    with pdfplumber.open(BytesIO(rendered.pdf)) as pdf:
        assert all(
            0 <= char["x0"] < char["x1"] <= page.width
            and 0 <= char["top"] < char["bottom"] <= page.height
            for page in pdf.pages
            for char in page.chars
        )


async def test_missing_ai_fails_without_storage_or_render(tmp_path):
    service = CvFormatting()
    service.runtime = SimpleNamespace(ai=None, storage=OwnedStorage(LocalStorage(tmp_path), 1))
    with pytest.raises(ServiceError, match="provider_config"):
        await service.run(inputs())
    assert not list(tmp_path.glob("[0-9]*/*/*"))
