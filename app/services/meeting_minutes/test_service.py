from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pydantic import ValidationError

from app.builders.word import build
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.meeting_minutes.schema import Inputs, classification_schema, document
from app.services.meeting_minutes.service import MeetingMinutes
from tests.minutes_fixtures import inputs, plan
from tests.word_pdf_fixtures import renderer


@pytest.mark.parametrize(
    "assignments",
    [
        [{"note_id": 1, "category": "action"}] * 4,
        [{"note_id": 1, "category": "action"}],
        [{"note_id": 99, "category": "action"}],
        [{"note_id": True, "category": "action"}],
        [{"note_id": "1", "category": "action"}],
        [{"note_id": 1, "category": "invented"}],
        [{"note_id": 1, "category": "action", "assignee": "invented"}],
    ],
    ids=["duplicates", "omission", "unknown", "boolean", "string", "category", "extra-fact"],
)
def test_schema_rejects_invented_fields_and_missing_extra_or_non_integer_ids(assignments):
    with pytest.raises(ValidationError):
        classification_schema(Inputs.parse(inputs()).lines()).model_validate(
            {"assignments": assignments}
        )


def test_schemas_are_request_local_and_repeated_source_notes_are_retained():
    first = classification_schema(["Same note", "Same note"])
    second = classification_schema(["Only note"])
    assignments = [
        {"note_id": 1, "category": "review"},
        {"note_id": 2, "category": "review"},
    ]
    assert len(first.model_validate({"assignments": assignments}).assignments) == 2
    with pytest.raises(ValidationError):
        second.model_validate({"assignments": assignments})
    assert "Same note" not in str(first.model_json_schema())


@pytest.mark.parametrize("language", ["ar", "en"])
def test_source_paragraphs_remain_verbatim_with_ids_and_no_generated_facts(language):
    values = Inputs.parse(inputs(language))
    classification = classification_schema(values.lines()).model_validate(plan())
    word_plan = document(values, classification)
    doc = Document(BytesIO(build(word_plan, professional_template=True, format_dates=True)))
    paragraphs = [p.text.replace("\u200e", "") for p in doc.paragraphs]
    for index, source in enumerate(values.lines(), 1):
        assert paragraphs.count(f"[{index}] {source}") == 1
    assert not doc.tables
    assert len(word_plan.sections) == 4


@pytest.mark.parametrize(
    "patch",
    [
        {"notes": "   "},
        {"notes": "note\n" * 81},
        {"notes": "x" * 1201},
        {"notes": "\x00bad"},
        {"title": "two\nlines"},
        {"language": "bilingual"},
    ],
    ids=["empty", "too-many", "too-long", "control", "title", "language"],
)
def test_invalid_notes_rejected_before_reservation(patch):
    with pytest.raises(ServiceError, match="input_invalid"):
        MeetingMinutes.needs_ai(inputs() | patch)


async def test_review_precedes_render_and_resume_uses_no_second_ai_call(tmp_path):
    class AI:
        calls = 0

        async def extract(self, schema, source, prompt):
            self.calls += 1
            assert "numbered source" in prompt and "00123" in source
            return schema.model_validate(plan())

    class Renderer:
        async def word_pdf(self, *args, **kwargs):
            raise AssertionError("review must precede file creation")

    ai = AI()
    service = MeetingMinutes()
    service.runtime = SimpleNamespace(
        ai=ai, renderer=Renderer(), storage=OwnedStorage(LocalStorage(tmp_path), 1)
    )
    source = inputs()
    service.input_schema.validate_inputs(source)
    assert service.needs_ai(source)
    review = await service.run(source)
    assert review.needs_confirmation and not review.artifacts
    assert "[3]" in review.preview and "00017" in review.preview
    for lang in ("ar", "en"):
        assert (
            "[3]" in review.preview_localizations[lang]
            and "00017" in review.preview_localizations[lang]
        )
    assert not list(tmp_path.glob("[0-9]*/*/*"))
    resumed = source | {"__continuation": review.continuation}
    assert not service.needs_ai(resumed)
    service.runtime.ai = None
    service.runtime.renderer = renderer()
    result = await service.run(resumed)
    assert [file.filename for file in result.artifacts] == ["minutes.docx", "minutes.pdf"]
    assert ai.calls == 1


async def test_corrupt_continuation_does_not_make_a_new_provider_call_or_export(tmp_path):
    service = MeetingMinutes()
    service.runtime = SimpleNamespace(
        ai=None, renderer=None, storage=OwnedStorage(LocalStorage(tmp_path), 1)
    )
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run(inputs() | {"__continuation": {}})
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_review_lists_all_source_ids_even_when_only_two_samples_are_shown(tmp_path):
    class AI:
        async def extract(self, schema, *args):
            return schema.model_validate(
                {"assignments": [{"note_id": index, "category": "review"} for index in range(1, 5)]}
            )

    service = MeetingMinutes()
    service.runtime = SimpleNamespace(
        ai=AI(), renderer=None, storage=OwnedStorage(LocalStorage(tmp_path), 1)
    )
    result = await service.run(inputs())
    assert "1, 2, 3, 4" in result.preview
    assert "[1]" in result.preview and "[2]" in result.preview
    assert "[3]" not in result.preview and result.needs_confirmation


async def test_maximum_note_review_keeps_all_ids_in_both_ui_languages(tmp_path):
    class AI:
        async def extract(self, schema, *args):
            categories = ("discussion", "decision", "action", "review")
            return schema.model_validate(
                {
                    "assignments": [
                        {"note_id": index, "category": categories[(index - 1) % 4]}
                        for index in range(1, 81)
                    ]
                }
            )

    service = MeetingMinutes()
    service.runtime = SimpleNamespace(
        ai=AI(), renderer=None, storage=OwnedStorage(LocalStorage(tmp_path), 1)
    )
    source = {
        **inputs("en"),
        "notes": "\n".join(f"Source {i:02d} " + "x" * 130 for i in range(1, 81)),
    }
    result = await service.run(source)
    assert result.needs_confirmation
    for lang in ("ar", "en"):
        review = result.preview_localizations[lang]
        assert len(review) <= 2800
        assert ", 80" in review and "Source 01" in review
