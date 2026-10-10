from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.builders.pdf_tables import build
from app.providers.storage import LocalStorage, OwnedStorage
from app.providers.tables.base import Extraction
from app.services.base import ServiceError
from app.services.pdf_tables_excel.schema import direct_labels, labels_schema
from app.services.pdf_tables_excel.service import PdfTablesExcel

ROWS = [
    ["ID", "Amount", "Note"],
    ["00123", "125.50", "=1+1"],
    ["00017", "0", "2026-10-10"],
    ["", " +value", "a@example.invalid"],
]
LABELS = {
    "tables": [{"table_id": 1, "title": "Orders", "columns": ["Identifier", "Amount", "Note"]}]
}


def extraction():
    return Extraction.model_validate(
        {"pages": 1, "tables": [{"page": 1, "rows": ROWS, "accuracy": 100.0, "whitespace": 0.0}]}
    )


def inputs(**changes):
    return {
        "pdf": "1/" + "a" * 32 + "/source.pdf",
        "method": "lattice",
        "mode": "direct",
        "language": "en",
        **changes,
    }


@pytest.mark.parametrize(
    "changes",
    [{"method": "auto"}, {"mode": "invented"}, {"language": "fr"}, {"pdf": ""}, {"unknown": True}],
)
def test_invalid_admission(changes):
    with pytest.raises(ServiceError):
        PdfTablesExcel.needs_ai(inputs(**changes))


def test_pure_ai_admission_and_continuation_gate():
    assert not PdfTablesExcel.needs_ai(inputs())
    assert PdfTablesExcel.needs_ai(inputs(mode="labels"))
    assert not PdfTablesExcel.needs_ai({**inputs(mode="labels"), "__continuation": {}})
    assert not PdfTablesExcel.requires_ai and not PdfTablesExcel.enabled_by_default


@pytest.mark.parametrize(
    "plan",
    [
        {"tables": []},
        {"tables": [LABELS["tables"][0], LABELS["tables"][0]]},
        {"tables": [{**LABELS["tables"][0], "table_id": 2}]},
        {"tables": [{**LABELS["tables"][0], "table_id": True}]},
        {"tables": [{**LABELS["tables"][0], "table_id": "1"}]},
        {"tables": [{**LABELS["tables"][0], "columns": ["One", "Two"]}]},
        {"tables": [{**LABELS["tables"][0], "columns": ["One", "one", "Three"]}]},
        {"tables": [{**LABELS["tables"][0], "title": "Total 125.50"}]},
        {"tables": [{**LABELS["tables"][0], "columns": ["One", "Two", "Year 2026"]}]},
        {"tables": [{**LABELS["tables"][0], "title": "bad\nheading"}]},
        {"tables": [{**LABELS["tables"][0], "rows": [["changed"]]}]},
    ],
)
def test_labels_only_allow_exact_tables_columns_and_no_numeric_claims(plan):
    with pytest.raises(ValidationError):
        labels_schema(extraction()).model_validate(plan)


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("enhanced", [False, True])
def test_excel_all_cells_literal_rows_references_source_and_review(language, enhanced):
    source = extraction()
    labels = (
        labels_schema(source).model_validate(LABELS)
        if enhanced
        else direct_labels(source, language)
    )
    data = build(source, labels, language=language, enhanced=enhanced)
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Data1", "Source1", "Review"]
    for name in ("Data1", "Source1"):
        sheet = book[name]
        assert sheet.freeze_panes == "A4" and len(sheet.tables) == 1
        assert sheet.print_title_rows == "$1:$3"
        for r, row in enumerate(ROWS, 4):
            for c, value in enumerate(row, 1):
                assert sheet.cell(r, c).value == (value or None)
                assert sheet.cell(r, c).data_type != "f"
                assert sheet.cell(r, c).number_format == "@"
            assert sheet.cell(r, 4).value == f"T1:R{r - 3}"
        assert sheet["A5"].value == "00123" and sheet["C5"].value == "=1+1"
    assert book["Review"]["B4"].value == "1" and book["Review"]["C4"].value == "4"
    assert source == extraction()  # No source mutation when applying headings.


async def test_review_and_confirmed_export_do_not_reextract_or_reinvoke_ai(tmp_path):
    storage = OwnedStorage(LocalStorage(tmp_path), 1)
    file = storage.save("source.pdf", b"synthetic", "application/pdf")
    calls = []

    async def extract(data, method):
        calls.append((data, method))
        return extraction()

    async def ai_extract(schema, source, prompt):
        assert "Arabic" in prompt and "sample_rows" in source
        return schema.model_validate(LABELS)

    service = PdfTablesExcel()
    service.runtime = SimpleNamespace(
        storage=storage,
        tables=SimpleNamespace(extract=extract),
        ai=SimpleNamespace(extract=ai_extract),
    )
    request = inputs(pdf=file.key, mode="labels", language="ar")
    result = await service.run(request)
    assert result.needs_confirmation and not result.artifacts and len(calls) == 1
    service.runtime.tables = None
    service.runtime.ai = None
    final = await service.run({**request, "__continuation": result.continuation})
    assert final.artifacts[0].filename == "tables.xlsx"
    assert (
        load_workbook(BytesIO(storage.read(final.artifacts[0].key)))["Data1"]["A5"].value == "00123"
    )
    assert len(calls) == 1
    broken = {
        **result.continuation,
        "labels": {"tables": [{**LABELS["tables"][0], "columns": ["wrong", "count"]}]},
    }
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run({**request, "__continuation": broken})


@pytest.mark.parametrize("mode", ["direct", "labels"])
async def test_missing_runtime_capability_creates_no_output(tmp_path, mode):
    service = PdfTablesExcel()
    service.runtime = SimpleNamespace(
        storage=OwnedStorage(LocalStorage(tmp_path), 1), ai=None, tables=None
    )
    with pytest.raises(ServiceError):
        await service.run(inputs(mode=mode))
    assert not list(tmp_path.glob("[0-9]*/*/*"))


def test_oversized_heading_plan_cannot_hide_tables_from_confirmation():
    source = Extraction.model_validate(
        {
            "pages": 2,
            "tables": [
                {"page": page, "rows": [["x"] * 12] * 2, "accuracy": 100, "whitespace": 0}
                for page in (1, 2)
            ],
        }
    )
    with pytest.raises(ValidationError):
        labels_schema(source).model_validate(
            {
                "tables": [
                    {
                        "table_id": page,
                        "title": "Title",
                        "columns": [chr(65 + i) * 100 for i in range(12)],
                    }
                    for page in (1, 2)
                ]
            }
        )


async def test_file_heading_language_does_not_control_review_language(tmp_path):
    storage = OwnedStorage(LocalStorage(tmp_path), 1)
    file = storage.save("source.pdf", b"synthetic", "application/pdf")

    async def extract(*args):
        return extraction()

    service = PdfTablesExcel()
    service.runtime = SimpleNamespace(
        storage=storage, tables=SimpleNamespace(extract=extract), ai=None
    )
    request = inputs(pdf=file.key, language="en")
    result = await service.run(request)
    assert "راجع" in result.preview_localizations["ar"]
    assert "Review" in result.preview_localizations["en"]
    assert "Column 1" not in result.preview_localizations["ar"]
    final = await service.run({**request, "__continuation": result.continuation})
    assert "Excel" in final.preview_localizations["ar"]
    book = load_workbook(BytesIO(storage.read(final.artifacts[0].key)))
    assert book["Data1"]["A3"].value == "Column 1"
    assert book["Data1"]["A5"].value == "00123"
