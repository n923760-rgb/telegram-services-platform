from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from pptx import Presentation

from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.text_to_pptx.schema import DeckPlan
from app.services.text_to_pptx.service import TextToPptx


async def test_pptx_service_builds_editable_file(tmp_path):
    plan = DeckPlan.model_validate({"deck": {"slides": [{"title": "عنوان", "bullets": ["نقطة"]}]}})
    service = TextToPptx()
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service.runtime = SimpleNamespace(
        storage=store, ai=SimpleNamespace(extract=AsyncMock(return_value=plan))
    )
    result = await service.run({"text": "نص"})
    assert not result.needs_confirmation
    presentation = Presentation(BytesIO(store.read(result.artifacts[0].key)))
    assert any(
        shape.has_text_frame and "نقطة" in shape.text for shape in presentation.slides[0].shapes
    )


async def test_pptx_service_builds_native_table_with_count_and_preserved_ids(tmp_path):
    from tests.test_pptx_quality import sample_plan, sample_source

    plan = DeckPlan.model_validate({"deck": sample_plan()})
    service = TextToPptx()
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service.runtime = SimpleNamespace(
        storage=store, ai=SimpleNamespace(extract=AsyncMock(return_value=plan))
    )
    result = await service.run({"text": sample_source()})
    prs = Presentation(BytesIO(store.read(result.artifacts[0].key)))
    assert len(prs.slides) == 3 and not result.needs_confirmation
    native = next(s.table for s in prs.slides[2].shapes if s.has_table)
    assert native.cell(1, 0).text == "00123" and native.cell(1, 2).text == "125.50"
    assert "الجداول القابلة للتعديل: 1" in result.preview
    service.runtime.ai.extract.assert_awaited_once()


async def test_pptx_service_missing_information_releases(tmp_path):
    plan = DeckPlan.model_validate({"missing_information": True})
    service = TextToPptx()
    service.runtime = SimpleNamespace(
        storage=OwnedStorage(LocalStorage(tmp_path), 1),
        ai=SimpleNamespace(extract=AsyncMock(return_value=plan)),
    )
    with pytest.raises(ServiceError, match="needs_information"):
        await service.run({"text": "غير كافٍ"})


async def test_pptx_service_ambiguous_confirmation_then_resume(tmp_path):
    plan = DeckPlan.model_validate(
        {
            "ambiguous": True,
            "question": "اعتماد هذا الهيكل؟",
            "deck": {"slides": [{"title": "عنوان", "bullets": ["نقطة"]}]},
        }
    )
    service = TextToPptx()
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service.runtime = SimpleNamespace(
        storage=store, ai=SimpleNamespace(extract=AsyncMock(return_value=plan))
    )
    first = await service.run({"text": "نص"})
    assert first.needs_confirmation and first.continuation
    # Resuming with the continuation must not call the AI provider again.
    second = await service.run({"text": "نص", "__continuation": first.continuation})
    assert not second.needs_confirmation
    assert service.runtime.ai.extract.await_count == 1
    presentation = Presentation(BytesIO(store.read(second.artifacts[0].key)))
    assert any(
        shape.has_text_frame and "نقطة" in shape.text for shape in presentation.slides[0].shapes
    )


@pytest.mark.parametrize("text", ["", " \n", "\x00", "bad\x1f", "\ud800", "\uffff", "x" * 12001])
def test_pptx_invalid_source_rejected_before_reservation(text):
    with pytest.raises(ServiceError, match="input_invalid"):
        TextToPptx.needs_ai({"text": text})


def test_pptx_continuation_does_not_need_another_provider():
    assert TextToPptx.needs_ai({"text": "Content"})
    assert not TextToPptx.needs_ai({"text": "Content", "__continuation": {}})


@pytest.mark.parametrize("location", ["title", "bullet", "column", "cell"])
@pytest.mark.parametrize(
    "token",
    ["123", "125.5", "2026-10-10", "new@example.invalid", "https://other.invalid", "-00123", "12%"],
)
def test_source_schema_rejects_changed_or_invented_tokens_everywhere(location, token):
    from pydantic import ValidationError

    from app.services.text_to_pptx.schema import source_plan_schema

    slide = {"title": "Summary", "bullets": ["Information"]}
    if location == "title":
        slide["title"] = "Record " + token
    elif location == "bullet":
        slide["bullets"] = ["Record " + token]
    else:
        slide.pop("bullets")
        slide["table"] = {"columns": ["Record"], "rows": [["Information"]]}
        if location == "column":
            slide["table"]["columns"] = ["Record " + token]
        else:
            slide["table"]["rows"] = [[token]]
    with pytest.raises(ValidationError):
        source_plan_schema(
            "Record 00123 Total 125.50 Date 2026-10-09 old@example.invalid https://source.invalid"
        ).model_validate({"deck": {"slides": [slide]}})


def test_source_counter_bounds_context_reuse_and_ignores_instruction_only_slide_count():
    from pydantic import ValidationError

    from app.services.text_to_pptx.schema import source_plan_schema

    plan = {
        "deck": {
            "slides": [
                {"title": "Records", "bullets": ["Invoice 00123"]},
                {"title": "Follow up", "bullets": ["Invoice 00123"]},
            ]
        }
    }
    assert source_plan_schema("Prepare 2 slides. Invoice 00123").model_validate(plan)
    validated = source_plan_schema("Prepare 2 slides. Invoice 00123 twice: 00123").model_validate(
        plan
    )
    assert len(validated.deck.slides) == 2
    plan["deck"]["slides"].append({"title": "Again", "bullets": ["Invoice 00123"]})
    with pytest.raises(ValidationError):
        source_plan_schema("Prepare 3 slides. Invoice 00123").model_validate(plan)


@pytest.mark.parametrize("continuation", [False, True])
async def test_forged_response_or_continuation_cannot_save_invented_cost(tmp_path, continuation):
    plan = DeckPlan.model_validate(
        {"deck": {"slides": [{"title": "Invoice", "bullets": ["Amount 999.50"]}]}}
    )
    ai = SimpleNamespace(extract=AsyncMock(return_value=plan))
    service = TextToPptx()
    service.runtime = SimpleNamespace(ai=ai, storage=OwnedStorage(LocalStorage(tmp_path), 1))
    source = {"text": "Amount 125.50"}
    if continuation:
        source["__continuation"] = plan.model_dump()
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run(source)
    assert ai.extract.await_count == (0 if continuation else 1)
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_pptx_missing_ai_has_safe_configuration_error(tmp_path):
    service = TextToPptx()
    service.runtime = SimpleNamespace(ai=None, storage=OwnedStorage(LocalStorage(tmp_path), 1))
    with pytest.raises(ServiceError, match="provider_config"):
        await service.run({"text": "Content"})
    assert not list(tmp_path.glob("[0-9]*/*/*"))
