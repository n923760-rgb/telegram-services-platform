from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from app.ops.acceptance_cases import Case, cases, source_pdf
from app.ops.service_acceptance import ASSETS, inspect_output
from app.services.base import ServiceError


@pytest.fixture(autouse=True)
def database():
    yield


def pdf_bytes(count, rotation=0):
    writer = PdfWriter()
    source = PdfReader(BytesIO(source_pdf())).pages[0]
    for _ in range(count):
        writer.add_page(source).rotation = rotation
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("count", [1, 2, 3])
def test_exact_pdf_page_contract_rejects_extra_as_well_as_missing_pages(count):
    case = Case("exact", "pdf_tools", {}, exact_pdf_pages=2)
    if count == 2:
        assert inspect_output("result.pdf", pdf_bytes(count), case)["pages"] == 2
    else:
        with pytest.raises(ServiceError, match="acceptance_output_invalid"):
            inspect_output("result.pdf", pdf_bytes(count), case)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_rotation_contract_checks_native_flags_without_rewriting_text(rotation):
    case = Case("rotation", "pdf_tools", {}, ("00123",), pdf_rotations=(90,))
    if rotation == 90:
        assert inspect_output("result.pdf", pdf_bytes(1, rotation), case)["rotations"] == [90]
    else:
        with pytest.raises(ServiceError, match="acceptance_output_invalid"):
            inspect_output("result.pdf", pdf_bytes(1, rotation), case)


def test_rotation_tuple_checks_every_page_and_length():
    case = Case("rotation", "pdf_tools", {}, pdf_rotations=(90, 0))
    with pytest.raises(ServiceError, match="acceptance_output_invalid"):
        inspect_output("result.pdf", pdf_bytes(2, 90), case)
    with pytest.raises(ServiceError, match="acceptance_output_invalid"):
        inspect_output("result.pdf", pdf_bytes(1, 90), case)


@pytest.mark.parametrize("count", [1, 2, 3, 4])
def test_selected_slide_count_is_an_exact_output_contract(count):
    deck = Presentation()
    for _ in range(count):
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Synthetic"
    output = BytesIO()
    deck.save(output)
    case = Case("deck", "text_to_pptx", {}, pptx_slides=3)
    if count == 3:
        assert inspect_output("result.pptx", output.getvalue(), case)["slides"] == 3
    else:
        with pytest.raises(ServiceError, match="acceptance_output_invalid"):
            inspect_output("result.pptx", output.getvalue(), case)


@pytest.mark.parametrize("present", [True, False])
def test_powerpoint_acceptance_requires_source_values_including_native_table_cells(present):
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "00123" if present else "123"
    table.cell(0, 1).text = "125.50"
    output = BytesIO()
    deck.save(output)
    case = Case("pptx", "text_to_pptx", {}, ("00123", "125.50"), pptx_slides=1)
    if present:
        assert inspect_output("result.pptx", output.getvalue(), case)["slides"] == 1
    else:
        with pytest.raises(ServiceError, match="acceptance_source_token_missing"):
            inspect_output("result.pptx", output.getvalue(), case)


@pytest.mark.parametrize("missing", [None, "Data", "Source", "Review", "Columns"])
def test_required_workbook_sheets_reject_missing_but_allow_optional_sheets(missing):
    required = ("Data", "Source", "Review", "Columns")
    book = Workbook()
    book.remove(book.active)
    for name in (*required, "Optional"):
        if name != missing:
            book.create_sheet(name)["A1"] = "Synthetic"
    output = BytesIO()
    book.save(output)
    case = Case("csv", "csv_review", {}, required_sheets=required)
    if missing is None:
        assert inspect_output("review.xlsx", output.getvalue(), case)["sheets"] == [
            *required,
            "Optional",
        ]
    else:
        with pytest.raises(ServiceError, match="acceptance_output_invalid"):
            inspect_output("review.xlsx", output.getvalue(), case)


def test_legacy_minimum_pdf_contract_is_retained_for_variable_layouts():
    case = Case("cv", "cv_formatting", {}, pdf_pages=2)
    assert inspect_output("cv.pdf", pdf_bytes(3), case)["pages"] == 3


def test_plan_covers_document_contracts_alongside_smart_card_cases():
    plan = {case.name: case for case in cases(ASSETS)}
    assert len(plan) == 33
    assert plan["pptx"].inputs["slide_count"] == "3" and plan["pptx"].pptx_slides == 3
    assert plan["pptx"].tokens == ("00123", "125.50")
    assert plan["images_pdf"].inputs["page_order"] == "٢، ١"
    assert plan["images_pdf"].exact_pdf_pages == 2
    assert plan["pdf_rotate"].pdf_rotations == (90,)
    assert plan["pdf_merge"].exact_pdf_pages == 2
    assert "Columns" in plan["csv"].required_sheets
    assert plan["pdf_word"].docx_tables[0][1] == ("00123", "125.50", "Example")
    assert all("title" not in plan["minutes_" + lang].inputs for lang in ("ar", "en"))


@pytest.mark.parametrize("mutation", [None, "bullets", "swapped", "normalized", "duplicate"])
def test_pdf_word_acceptance_requires_exact_editable_cells(mutation):
    case = next(case for case in cases(ASSETS) if case.name == "pdf_word")
    document = Document()
    matrix = case.docx_tables[0]
    if mutation == "bullets":
        for row in matrix:
            document.add_paragraph(" ".join(row), style="List Bullet")
    else:
        table = document.add_table(rows=len(matrix), cols=len(matrix[0]))
        for cells, values in zip(table.rows, matrix, strict=True):
            for cell, value in zip(cells.cells, values, strict=True):
                cell.text = value
        if mutation == "swapped":
            table.cell(1, 0).text, table.cell(2, 0).text = "00017", "00123"
        if mutation == "normalized":
            table.cell(1, 0).text = "123"
        if mutation == "duplicate":
            document.add_table(rows=1, cols=1).cell(0, 0).text = "00123"
    output = BytesIO()
    document.save(output)
    if mutation is None:
        assert inspect_output("result.docx", output.getvalue(), case)["editable_tables"] == 1
    else:
        with pytest.raises(ServiceError, match="acceptance_output_invalid"):
            inspect_output("result.docx", output.getvalue(), case)
