from io import BytesIO

import arabic_reshaper
import pytest
from bidi.algorithm import get_display
from openpyxl import load_workbook

from app.builders.arabic_table_text import recover_visual_arabic
from app.builders.pdf_tables import build
from app.core.i18n import CATALOGS
from app.providers.tables.base import Extraction
from app.services.base import ServiceError
from app.services.pdf_tables_excel.schema import direct_labels
from app.services.pdf_tables_excel.service import SERVICE
from app.services.pdf_tables_excel.test_service import inputs


def visual(text):
    return get_display(arabic_reshaper.reshape(text))


@pytest.mark.parametrize(
    "text", ["بصل", "طماطم", "بطاطس", "رقم الطلب", "لا", "شركة الربيع", "بصل\nطماطم"]
)
def test_explicit_simple_visual_arabic_recovery(text):
    assert recover_visual_arabic(visual(text)) == text


@pytest.mark.parametrize(
    "text",
    [
        "00123",
        "0.00",
        "-125.50",
        "2026-10-13",
        "=1+1",
        "بصل",
        "رقم الطلب",
        "",
        "English",
        "ا\u0650",
    ],
)
def test_ordinary_arabic_numbers_formulas_and_empty_unchanged(text):
    assert recover_visual_arabic(text) == text


@pytest.mark.parametrize("suffix", [" 00123", " ٠٠١٢٣", " English", "@x", "،", "\t", "\u064e", "ب"])
def test_ambiguous_mixed_cells_unchanged(suffix):
    text = visual("بصل") + suffix
    assert recover_visual_arabic(text) == text


@pytest.mark.parametrize("language", ["ar", "en"])
def test_workbook_opt_in_data_only_source_exact_and_refs_stable(language):
    rows = [
        [visual("رقم الطلب"), visual("المنتج"), visual("السعر")],
        ["00123", visual("بصل"), "12.50"],
        ["00017", visual("بطاطس") + " 1", "0.00"],
    ]
    source = Extraction.model_validate(
        {"pages": 1, "tables": [{"page": 1, "rows": rows, "accuracy": 100, "whitespace": 0}]}
    )
    before = source.model_dump()
    for enabled in [False, True]:
        book = load_workbook(
            BytesIO(
                build(
                    source,
                    direct_labels(source, language),
                    language=language,
                    enhanced=False,
                    visual_arabic=enabled,
                )
            )
        )
        for r, row in enumerate(rows, 4):
            for c, value in enumerate(row, 1):
                assert book["Source1"].cell(r, c).value == value
                assert book["Data1"].cell(r, c).value == (
                    recover_visual_arabic(value) if enabled else value
                )
                assert book["Data1"].cell(r, c).data_type == "s"
            assert (
                book["Data1"].cell(r, 4).value == book["Source1"].cell(r, 4).value == f"T1:R{r - 3}"
            )
        assert book["Data1"]["A5"].value == "00123"
        assert book["Data1"]["C6"].value == "0.00"
        if enabled:
            assert "4" in book["Data1"]["A2"].value and "1" in book["Data1"]["A2"].value
        assert source.model_dump() == before


def test_intake_default_no_ai_and_catalogs():
    assert not SERVICE.needs_ai(inputs(arabic_order="visual"))
    for value in ["auto", "reverse", True]:
        with pytest.raises(ServiceError):
            SERVICE.needs_ai(inputs(arabic_order=value))
    for key in (
        "pdf_tables_arabic_order",
        "pdf_tables_arabic_visual",
        "pdf_tables_arabic_preserve",
        "pdf_tables_arabic_review",
        "pdf_tables_arabic_note",
    ):
        assert all(key in CATALOGS[lang] for lang in ("ar", "en"))
