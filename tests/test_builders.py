from io import BytesIO

from docx import Document as Word
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl import load_workbook
from pptx import Presentation

from app.builders import excel, pdf, pptx, word
from app.builders.direction import has_arabic, is_rtl
from app.builders.schema import Deck, Document, Section, Slide, Table
from app.builders.word_schema import WordDocument, WordSection, WordTable


def test_word_tables_are_editable_bounded_and_have_no_terminal_spacer():
    table = WordTable(columns=["Reference", "Description"], rows=[["00123", "Text"]])
    data = WordDocument(title="Brief", sections=[WordSection(tables=[table, table])])
    doc = Word(BytesIO(word.build(data, include_title=False)))
    assert len(doc.tables) == 2
    assert doc.element.body[-2].tag == qn("w:tbl")
    assert len(doc.paragraphs) == 1  # Only the separator between the two tables.
    available = (
        doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    )
    for native in doc.tables:
        assert native.cell(1, 0).text == "00123"
        assert abs(sum(c.width for c in native.columns) - available) <= 1270
        assert native.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in native.rows)
        assert native.autofit is False


def test_six_column_word_table_stays_inside_page_margins():
    table = WordTable(
        columns=list("ABCDEF"), rows=[["00123", "2026-10-08", "125.50", "", "نص", "X"]]
    )
    data = WordDocument(title="Table", sections=[WordSection(tables=[table])])
    doc = Word(BytesIO(word.build(data)))
    available = (
        doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    )
    assert abs(sum(c.width for c in doc.tables[0].columns) - available) <= 3810
    assert [c.text for c in doc.tables[0].rows[1].cells] == table.rows[0]


def test_direction_detects_arabic_english_and_mixed():
    assert is_rtl("تقرير سنوي") is True
    assert is_rtl("Annual Report") is False
    assert is_rtl("Annual Report 2024") is False
    assert is_rtl("تقرير Q3 2024") is True
    assert has_arabic("hello مرحبا") is True
    assert has_arabic("hello") is False


def test_word_builds_rtl_and_ltr_with_professional_layout():
    data = Document(
        title="Annual Report",
        sections=[
            Section(heading="Overview", paragraphs=["English content only."]),
            Section(heading="الملخص", paragraphs=["محتوى عربي."], bullets=["نقطة عربية"]),
        ],
    )
    doc = Word(BytesIO(word.build(data)))
    by_text = {paragraph.text: paragraph._p.get_or_add_pPr() for paragraph in doc.paragraphs}
    assert "English content only." in by_text
    assert "محتوى عربي." in by_text
    assert "نقطة عربية" in by_text
    for text, rtl in [("Overview", "0"), ("الملخص", "1")]:
        assert by_text[text].find(qn("w:jc")).get(qn("w:val")) == "start"
        assert by_text[text].find(qn("w:bidi")).get(qn("w:val")) == rtl
    assert abs(doc.sections[0].top_margin - Cm(2.2)) <= 635
    assert "PAGE" in doc.sections[0].footer.paragraphs[0]._p.xml
    assert "w:pBdr" not in doc.styles["Title"].element.xml


def test_excel_preserves_protection_and_adaptive_direction():
    table = Table(
        title="بيانات",
        columns=["الاسم", "القيمة"],
        rows=[["منتج", 12], ['=HYPERLINK("evil")', "00123"]],
    )
    sheet = load_workbook(BytesIO(excel.build(table))).active
    assert sheet.sheet_view.rightToLeft is True
    assert sheet["B2"].value == 12
    assert sheet["A3"].data_type == "s" and sheet["A3"].value.startswith("'")
    assert sheet["B3"].value == "00123"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions

    english = Table(title="Data", columns=["Name", "Value"], rows=[["Item", 1]])
    assert load_workbook(BytesIO(excel.build(english))).active.sheet_view.rightToLeft is False


def test_pptx_builds_title_slide_and_content_within_bounds():
    deck = Deck(
        slides=[
            Slide(title="تقرير", bullets=["نقطة أولى"]),
            Slide(title="الملخص", bullets=["نقطة ثانية"]),
        ]
    )
    prs = Presentation(BytesIO(pptx.build(deck)))
    assert len(prs.slides) == 2
    first = "\n".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "تقرير" in first and "نقطة أولى" in first
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height


def test_pdf_mixed_content_and_page_numbers():
    import pdfplumber

    data = Document(title="تقرير Report", sections=[Section(paragraphs=["Arabic اختبار 123"] * 40)])
    with pdfplumber.open(BytesIO(pdf.build(data))) as doc:
        assert len(doc.pages) >= 2
        text = "".join(page.extract_text() or "" for page in doc.pages)
        assert "123" in text


def test_pdf_does_not_add_blank_page_for_terminal_spacing():
    import pdfplumber

    text = "\n".join(
        f"السطر {i:02d}: محتوى عربي للمراجعة مع بيانات ثابتة 00123 وتاريخ 2026/10/08."
        if i % 2
        else f"Line {i:02d}: English content with unchanged facts 00123 and date 2026/10/08."
        for i in range(1, 46)
    )
    data = Document(
        title="مراجعة مستند متعدد الصفحات", sections=[Section(paragraphs=text.split("\n"))]
    )
    with pdfplumber.open(BytesIO(pdf.build(data))) as doc:
        assert len(doc.pages) == 2
        assert "45" in (doc.pages[-1].extract_text() or "")
