"""Column summaries use literal values after the customer's selected whitespace rule."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.builders.csv_review import build
from app.core.i18n import tr
from app.services.csv_review.schema import Inputs
from tests.csv_review_fixtures import inputs


@pytest.fixture(autouse=True)
def database():
    yield


def export(values, *, similarity=False):
    parsed = Inputs.parse(values)
    headers, rows = parsed.records()
    data, counts = build(
        headers,
        rows,
        trim=parsed.whitespace == "trim",
        language=parsed.language,
        similarity=similarity,
    )
    return load_workbook(BytesIO(data)), counts, data


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("similarity", [False, True])
def test_complete_source_column_summary_counts_and_native_layout(language, similarity):
    book, counts, _ = export(inputs(language), similarity=similarity)
    sheet = book["Columns"]
    assert book.sheetnames == ["Data", "Source", "Review", "Columns"] + (
        ["Similarities"] if similarity else []
    )
    assert [list(row[3:]) for row in sheet.iter_rows(min_row=4, values_only=True)] == [
        [4, 0, 3, 0],
        [3, 1, 2, 1],
        [4, 0, 3, 0],
        [3, 1, 2, 0],
    ]
    assert [sheet.cell(row, 1).value for row in range(4, 8)] == [1, 2, 3, 4]
    assert sum(sheet.cell(row, 5).value for row in range(4, 8)) == counts["blanks"]
    assert sum(sheet.cell(row, 7).value for row in range(4, 8)) == counts["changes"]
    assert sheet.tables["ColumnsRecords"].ref == "A3:G7"
    assert sheet.freeze_panes == "A4" and sheet.print_title_rows == "$1:$3"
    assert sheet.sheet_view.rightToLeft == (language == "ar")
    assert tr("csv_columns_rules", language) == sheet["A2"].value
    for row in sheet.iter_rows(min_row=4):
        assert row[1].data_type == row[2].data_type == "s"
        assert all(cell.data_type == "n" and type(cell.value) is int for cell in (row[0], *row[3:]))
        assert all(cell.alignment.wrap_text for cell in row)


@pytest.mark.parametrize(
    "mode,first,second",
    [
        ("trim", [4, 2, 3, 2], [4, 2, 2, 1]),
        ("preserve", [5, 1, 5, 0], [5, 1, 3, 0]),
    ],
)
def test_whitespace_case_nbsp_zeros_and_duplicates_have_explicit_literal_meaning(
    mode, first, second
):
    text = ' H , Value\n" \t ",0\n,00\nx,0\nX,\n x , \n\u00a0,0'
    book, _, _ = export({**inputs("en", whitespace=mode), "text": text})
    sheet = book["Columns"]
    assert list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))) == [
        1,
        " H ",
        "H",
        *first,
    ]
    assert list(next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))) == [
        2,
        " Value",
        "Value",
        *second,
    ]
    assert book["Source"]["A4"].value == " \t "
    assert book["Data"]["B5"].value == "00" and book["Data"]["B5"].data_type == "s"


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", " ="])
def test_customer_heading_values_remain_literal_without_formulas_links_or_macros(prefix):
    original = prefix + "Heading "
    book, _, data = export({**inputs("en"), "text": f"{original}\n00123\n0"})
    assert book["Columns"]["B4"].value == original
    assert book["Columns"]["C4"].value == original.strip()
    assert not any(
        cell.data_type == "f" or cell.hyperlink for sheet in book for row in sheet for cell in row
    )
    with ZipFile(BytesIO(data)) as archive:
        assert not any(
            "externallinks" in name.lower() or "vba" in name.lower() for name in archive.namelist()
        )
        assert all(
            b"<f>" not in archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/worksheets/")
        )


def test_missing_value_words_are_not_guessed_as_empty_or_normalized():
    book, _, _ = export({**inputs("en"), "text": "Value\nNULL\nN/A\nnan\nnull\nNULL"})
    assert [book["Columns"].cell(4, column).value for column in range(4, 8)] == [5, 0, 4, 0]


def test_maximum_empty_source_has_twenty_profile_rows_and_no_generated_reference_column():
    text = ",".join(f"C{i}" for i in range(20)) + "\n" + ("," * 19 + "\n") * 500
    book, counts, _ = export({**inputs("en"), "text": text})
    assert counts["rows"] == 500 and counts["blanks"] == 10000
    assert (
        book["Columns"].max_row == 23 and book["Columns"].tables["ColumnsRecords"].ref == "A3:G23"
    )
    assert all(
        list(row[3:]) == [0, 500, 0, 0]
        for row in book["Columns"].iter_rows(min_row=4, values_only=True)
    )
