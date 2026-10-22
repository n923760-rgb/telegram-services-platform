"""Resume hierarchy must preserve source text and leave standard documents alone."""

import re
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.builders.word import build
from tests.cv_fixtures import inputs, planned
from tests.word_pdf_fixtures import renderer


@pytest.fixture(autouse=True)
def database():
    """Pure renderer/schema checks do not use a database."""
    yield


@pytest.mark.parametrize("language", ["ar", "en", "bilingual"])
@pytest.mark.parametrize("theme", ["modern", "classic", "editorial"])
def test_resume_hierarchy_preserves_text_language_breaks_and_editable_body(language, theme):
    source = inputs("en" if language == "en" else "ar") | {"language": language}
    plan = planned(source, theme=theme)
    doc = Document(BytesIO(build(plan, professional_template=True)))
    standard = Document(
        BytesIO(build(plan.model_copy(update={"layout": "standard"}), professional_template=True))
    )
    assert [p.text for p in doc.paragraphs] == [p.text for p in standard.paragraphs]
    assert not doc.tables
    titles = [p for p in doc.paragraphs if p.style.name == "Title"]
    assert len(titles) == (2 if language == "bilingual" else 1)
    for title in titles:
        assert title.runs[0].font.size.pt < standard.styles["Title"].font.size.pt
        assert title.paragraph_format.keep_with_next
    if language == "bilingual":
        assert titles[1]._p.pPr.find(qn("w:pageBreakBefore")) is not None
    for section in plan.sections:
        if section.heading in {"Skills", "المهارات"}:
            assert all(block.kind == "paragraph" for block in section.blocks)
    assert all(
        p.paragraph_format.keep_with_next
        for p in doc.paragraphs
        if p.style.name in {"Heading 1", "Heading 2"}
    )


def test_long_qualification_does_not_compete_with_section_heading():
    source = inputs("en") | {
        "education": "Diploma in Business Administration from Example Institute of Business and Management — 2023"
    }
    doc = Document(BytesIO(build(planned(source), professional_template=True)))
    qualification = next(p for p in doc.paragraphs if p.text == source["education"])
    assert qualification.style.name == "Heading 2"
    assert all(not run.bold for run in qualification.runs)
    assert qualification.paragraph_format.keep_with_next


@pytest.mark.parametrize("language", ["ar", "en"])
async def test_dense_resume_paginates_without_losing_identifiers_or_overflow(language):
    import pdfplumber

    line = "Coordinated project records and appointments for department"
    if language == "ar":
        line = "تنسيق سجلات المشروع وتنظيم المواعيد ومتابعة الأعمال لدى القسم"
    source = inputs(language) | {
        "experience": "\n".join(f"{line} {index:05d}" for index in range(100, 140)),
    }
    result = await renderer().word_pdf(planned(source, density="compact"))
    pages = PdfReader(BytesIO(result.pdf)).pages
    assert 2 <= len(pages) <= 4
    text = "\n".join(page.extract_text() for page in pages)
    for index in range(100, 140):
        assert len(re.findall(r"(?<!\d)" + f"{index:05d}" + r"(?!\d)", text)) == 1
    assert text.count("0500123456") == text.count("sample@example.invalid") == 1
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        assert all(
            0 <= char["x0"] < char["x1"] <= page.width
            and 0 <= char["top"] < char["bottom"] <= page.height
            for page in pdf.pages
            for char in page.chars
        )
