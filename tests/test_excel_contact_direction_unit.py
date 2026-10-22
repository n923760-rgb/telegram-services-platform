"""Reopen real XLSX output from every shared builder; verify direction and fidelity."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.builders import csv_review, excel, excel_source, pdf_tables
from app.builders.schema import Table
from app.providers.tables.base import Extraction
from app.services.pdf_tables_excel.schema import direct_labels

CONTACTS = [
    "procurement.contracts.and.customer.accounts@example-company.invalid",
    "https://example.invalid/customer-services/contracts/department/reference",
    "www.example.invalid/customer-services/contracts/department/reference",
]


@pytest.fixture(autouse=True)
def database():
    yield


def source(lang, contact):
    columns = ["الملاحظات", "التواصل"] if lang == "ar" else ["Notes", "Contact"]
    text = ("يرجى التواصل عبر " if lang == "ar" else "Please contact ") + contact
    return columns, [[text, contact], ["00123", "125.50"]]


def assert_direction(sheet, row, lang, text, contact):
    note, identifier = sheet.cell(row, 1), sheet.cell(row, 2)
    assert note.value == text and identifier.value == contact
    assert note.alignment.horizontal == ("right" if lang == "ar" else "left")
    assert note.alignment.readingOrder == (2 if lang == "ar" else 1)
    assert identifier.alignment.horizontal == "left" and identifier.alignment.readingOrder == 1
    assert not any(c.data_type == "f" or c.hyperlink for cells in sheet for c in cells)
    assert sheet.freeze_panes == "A4" and len(sheet.tables) == 1
    assert not sheet.protection.sheet


@pytest.mark.parametrize("lang", ["ar", "en"])
@pytest.mark.parametrize("contact", CONTACTS)
def test_shared_excel_keeps_arabic_prose_and_ltr_contacts(lang, contact):
    columns, rows = source(lang, contact)
    table = Table(title="Source", columns=columns, rows=rows)
    original = table.model_dump_json()
    book = load_workbook(BytesIO(excel.build(table)))
    sheet = book.active
    assert sheet.sheet_view.rightToLeft == (lang == "ar")
    assert_direction(sheet, 4, lang, rows[0][0], contact)
    assert [c.value for c in sheet[5]] == rows[1]
    sheet["A4"] = rows[0][0] + " 00017"
    saved = BytesIO()
    book.save(saved)
    reopened = load_workbook(BytesIO(saved.getvalue())).active
    assert reopened["A4"].value.endswith("00017")
    assert reopened["A4"].alignment.readingOrder == (2 if lang == "ar" else 1)
    assert table.model_dump_json() == original


@pytest.mark.parametrize("lang", ["ar", "en"])
@pytest.mark.parametrize("contact", CONTACTS)
def test_source_excel_preserves_reference_and_contact_prefixed_source_direction(lang, contact):
    columns, rows = source(lang, contact)
    title = "طلبات العملاء" if lang == "ar" else "Customer orders"
    # A leading ASCII contact must not override the Arabic prose after it.
    lines = [contact + (" يرجى مراجعة الطلب" if lang == "ar" else " Please review the order")]
    table = Table(title=title, columns=columns, rows=rows[:1])
    book = load_workbook(BytesIO(excel_source.build(table, lines, [1], [[None, None]])))
    assert_direction(book["Data"], 4, lang, rows[0][0], contact)
    assert book["Data"]["C4"].value == "L001"
    raw = book["Source"]
    assert raw.sheet_view.rightToLeft == (lang == "ar")
    assert raw["C4"].value == lines[0]
    assert raw["C4"].alignment.horizontal == ("right" if lang == "ar" else "left")
    assert raw["C4"].alignment.readingOrder == (2 if lang == "ar" else 1)
    assert raw["A4"].value == "L001" and raw["A4"].data_type == "s"


@pytest.mark.parametrize("lang", ["ar", "en"])
@pytest.mark.parametrize("contact", CONTACTS)
def test_csv_data_and_source_keep_literal_values_and_direction(lang, contact):
    columns, rows = source(lang, contact)
    content, counts = csv_review.build(columns, rows, trim=False, language=lang)
    book = load_workbook(BytesIO(content))
    for name in ("Data", "Source"):
        sheet = book[name]
        assert sheet.sheet_view.rightToLeft == (lang == "ar")
        assert_direction(sheet, 4, lang, rows[0][0], contact)
        assert [sheet.cell(5, c).value for c in (1, 2)] == rows[1]
        assert sheet["C4"].value == "1" and sheet["C5"].value == "2"
    assert counts == {"rows": 2, "blanks": 0, "duplicates": 0, "changes": 0}


@pytest.mark.parametrize("lang", ["ar", "en"])
@pytest.mark.parametrize("contact", CONTACTS)
def test_pdf_table_data_and_source_keep_cells_references_and_direction(lang, contact):
    _, rows = source(lang, contact)
    extraction = Extraction.model_validate(
        {"pages": 1, "tables": [{"page": 1, "rows": rows, "accuracy": 100, "whitespace": 0}]}
    )
    labels = direct_labels(extraction, lang)
    book = load_workbook(
        BytesIO(pdf_tables.build(extraction, labels, language=lang, enhanced=False))
    )
    for name in ("Data1", "Source1"):
        sheet = book[name]
        assert sheet.sheet_view.rightToLeft == (lang == "ar")
        assert_direction(sheet, 4, lang, rows[0][0], contact)
        assert [sheet.cell(5, c).value for c in (1, 2)] == rows[1]
        assert sheet["C4"].value == "T1:R1" and sheet["C5"].value == "T1:R2"
    assert extraction.tables[0].rows == rows


def test_english_prose_and_standalone_arabic_digit_dates_remain_ltr():
    text = "Please confirm the delivery date ٢٠٢٦-١٠-٢١ for طلب " + CONTACTS[0]
    sheet = load_workbook(
        BytesIO(
            excel.build(
                Table(title="Orders", columns=["Note", "Date"], rows=[[text, "٢٠٢٦-١٠-٢١"]])
            )
        )
    ).active
    assert not sheet.sheet_view.rightToLeft
    for cell in sheet[4]:
        assert cell.alignment.horizontal == "left" and cell.alignment.readingOrder == 1
    assert sheet["A4"].value == text and sheet["B4"].value == "٢٠٢٦-١٠-٢١"
