"""One filter owner per native table, without overlapping worksheet filters."""

from datetime import datetime
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.builders import excel
from app.builders.schema import Table

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def assert_single_filter(content, ref):
    with ZipFile(BytesIO(content)) as package:
        sheet = ElementTree.fromstring(package.read("xl/worksheets/sheet1.xml"))
        table = ElementTree.fromstring(package.read("xl/tables/table1.xml"))
        book = ElementTree.fromstring(package.read("xl/workbook.xml"))
    assert sheet.find("s:autoFilter", NS) is None
    assert table.find("s:autoFilter", NS).get("ref") == ref
    assert table.get("ref") == ref
    assert not any(
        name.get("name") == "_xlnm._FilterDatabase"
        for name in book.findall("s:definedNames/s:definedName", NS)
    )


@pytest.mark.parametrize("arabic", [True, False])
def test_typed_table_owns_filter_and_preserves_customer_acceptance_records(arabic):
    columns = (
        ["رقم الطلب", "العميل", "المبلغ بالريال", "تاريخ الاستلام"]
        if arabic
        else ["Order ID", "Customer", "Amount SAR", "Received date"]
    )
    names = ["أحمد", "سارة", "خالد"] if arabic else ["Ahmed", "Sara", "Khaled"]
    data = Table(
        title="متابعة الطلبات" if arabic else "Order tracking",
        columns=columns,
        column_formats=[{"kind": k} for k in ("text", "text", "number", "date")],
        rows=[
            ["00123", names[0], 125.5, "2026-10-08"],
            ["00124", names[1], 0, "2026-10-09"],
            ["00125", names[2], None, "2026-10-10"],
            ["00123", names[0], 125.5, "2026-10-08"],
        ],
    )
    content = excel.build(data)
    assert_single_filter(content, "A3:D7")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert sheet.sheet_view.rightToLeft == arabic
    assert sheet.freeze_panes == "A4"
    assert sheet["A4"].value == "00123" and sheet["A4"].data_type == "s"
    assert sheet["A4"].number_format == "@"
    assert sheet["C4"].value == 125.5 and sheet["C5"].value == 0
    assert sheet["C6"].value is None
    assert sheet["D4"].value == datetime(2026, 10, 8)
    assert [c.value for c in sheet[4]] == [c.value for c in sheet[7]]
    assert sheet.tables["Records"].autoFilter.ref == "A3:D7"
    assert not sheet.protection.sheet
    sheet["B4"] = "Edited"
    saved = BytesIO()
    workbook.save(saved)
    assert_single_filter(saved.getvalue(), "A3:D7")
    assert load_workbook(BytesIO(saved.getvalue())).active["B4"].value == "Edited"


def test_legacy_table_also_has_one_filter_owner():
    data = Table(title="Legacy", columns=["ID", "Value"], rows=[["00123", 0], ["00123", None]])
    content = excel.build(data)
    assert_single_filter(content, "A3:B5")
    sheet = load_workbook(BytesIO(content)).active
    assert sheet["A4"].value == sheet["A5"].value == "00123"
    assert sheet["B4"].value == 0 and sheet["B5"].value is None
