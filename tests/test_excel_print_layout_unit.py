from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.builders import csv_review, excel, excel_source, pdf_tables
from app.builders.schema import Table
from tests.excel_print_fixtures import wide_table
from tests.excel_source_fixtures import typed_table


@pytest.fixture(autouse=True)
def database():
    yield


def wide_print(sheet):
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.scale == 100
    assert sheet.page_setup.fitToWidth == sheet.page_setup.fitToHeight == 0
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is False
    assert sheet.print_options.headings is True
    assert sheet.print_title_cols is None
    assert sheet.print_title_rows == "$1:$3"


@pytest.mark.parametrize("arabic", [False, True])
def test_wide_short_column_count_keeps_native_font_size_and_all_records(arabic):
    table = wide_table(arabic=arabic)
    before = table.model_dump()
    book = load_workbook(BytesIO(excel.build(table)))
    sheet = book.active
    wide_print(sheet)
    assert [list(r) for r in sheet.iter_rows(min_row=4, values_only=True)] == table.rows
    assert sheet.tables["Records"].ref == "A3:F5" and sheet.freeze_panes == "A4"
    assert all(c.font.sz == 11 for row in sheet.iter_rows(min_row=4) for c in row)
    assert table.model_dump() == before
    sheet["A4"] = "00017"
    saved = BytesIO()
    book.save(saved)
    reopened = load_workbook(BytesIO(saved.getvalue()))
    assert reopened.active["A4"].value == "00017"
    wide_print(reopened.active)


def test_compact_typed_table_retains_values_and_single_page_width():
    source = typed_table()
    sheet = load_workbook(BytesIO(excel.build(source))).active
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert not sheet.print_options.headings
    assert sheet["A4"].value == "00123"
    assert sheet["C4"].value == 125.5 and sheet["C5"].value == 0
    assert sheet["D4"].data_type == "d" and sheet["F4"].data_type == "b"


def test_landscape_is_selected_by_width_even_with_two_columns():
    table = Table(title="Notes", columns=["ID", "Details"], rows=[["X" * 80, "Y" * 80]])
    sheet = load_workbook(BytesIO(excel.build(table))).active
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1


def test_nine_compact_columns_do_not_require_horizontal_pagination():
    data = Table(title="Compact", columns=list("ABCDEFGHI"), rows=[[1] * 9])
    sheet = load_workbook(BytesIO(excel.build(data))).active
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1
    assert not sheet.print_options.headings


def test_source_column_padding_is_accounted_for_after_base_layout():
    table = Table(title="Rows", columns=list("ABCDEF"), rows=[["XXXXXXXXX"] * 6])
    augmented = Table(
        title=table.title,
        columns=[*table.columns, "Source reference"],
        rows=[table.rows[0] + ["L001"]],
    )
    base = load_workbook(BytesIO(excel.build(augmented))).active
    assert base.sheet_properties.pageSetUpPr.fitToPage is True
    book = load_workbook(
        BytesIO(excel_source.build(table, [" | ".join(table.rows[0])], [1], table.rows))
    )
    wide_print(book["Data"])
    assert [c.value for c in book["Data"][4]] == augmented.rows[0]
    assert book["Source"].page_setup.fitToWidth == 1


def test_csv_and_pdf_sheet_assembly_retains_wide_print_settings_and_literal_values():
    table = wide_table()
    table.rows[0][0] = "=1+1"
    csv_book, _ = csv_review.build(table.columns, table.rows, trim=False, language="en")
    # CSV returns serialized bytes; PDF uses an editable donor-sheet copy.
    csv_book = load_workbook(BytesIO(csv_book))
    for name in ("Data", "Source"):
        wide_print(csv_book[name])
        assert csv_book[name]["A4"].value == "=1+1"
        assert csv_book[name]["A4"].data_type == "s"
    sheet = pdf_tables.append_sheet(
        csv_book, table.title, table.columns, table.rows, "Printed", "Source note"
    )
    wide_print(sheet)
    assert sheet["A4"].value == "=1+1" and sheet["A4"].data_type == "s"
