"""Native Excel structure, source values and editable typed cells, not ZIP-only checks."""

from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.builders import excel
from app.builders.schema import Table


def typed_table():
    return Table.model_validate(
        {
            "title": "سجل الطلبات",
            "columns": ["المعرّف", "العميل", "المبلغ (ريال)", "التاريخ", "النسبة", "مؤكد"],
            "column_formats": [
                {"kind": kind} for kind in ("text", "text", "number", "date", "percent", "boolean")
            ],
            "rows": [
                ["00123", "أحمد", 125.5, "2026-10-08", 0.125, True],
                ["00124", "Example", 0, None, 0, False],
                ["00123", "أحمد", 125.5, "2026-10-08", 0.125, True],
            ],
        }
    )


def test_native_excel_table_preserves_types_missing_zero_duplicates_and_layout():
    data = typed_table()
    content = excel.build(data)
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert workbook.sheetnames == ["Data"]
    assert sheet["A1"].value == data.title
    assert [c.value for c in sheet[3]] == data.columns
    assert sheet.tables["Records"].ref == "A3:F6"
    assert sheet.tables["Records"].tableStyleInfo.showRowStripes
    assert sheet.freeze_panes == "A4" and sheet.sheet_view.rightToLeft
    assert not sheet.sheet_view.showGridLines
    assert sheet["A4"].value == "00123" and sheet["A4"].number_format == "@"
    assert sheet["C4"].value == 125.5 and sheet["C4"].data_type == "n"
    assert sheet["D4"].value == datetime(2026, 10, 8) and sheet["D4"].data_type == "d"
    assert sheet["D4"].number_format == "yyyy-mm-dd"
    assert sheet["E4"].value == 0.125 and sheet["E4"].number_format.endswith("%")
    assert sheet["F4"].value is True and sheet["F4"].data_type == "b"
    assert sheet["C5"].value == 0 and sheet["D5"].value is None
    assert sheet["E5"].value == 0 and sheet["F5"].value is False
    assert [c.value for c in sheet[4]] == [c.value for c in sheet[6]]
    assert sheet.print_title_rows == "$1:$3" and sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToHeight == 0
    assert sum(row[2] for row in data.rows) == sum(sheet.cell(r, 3).value for r in range(4, 7))
    with ZipFile(BytesIO(content)) as package:
        assert "xl/tables/table1.xml" in package.namelist()
        assert not any(
            "vba" in p.lower() or "externallinks" in p.lower() for p in package.namelist()
        )
    # A native cell edit survives save/reopen without converting the sheet to an image.
    sheet["C4"] = 250.75
    saved = BytesIO()
    workbook.save(saved)
    reopened = load_workbook(BytesIO(saved.getvalue()))
    assert reopened.active["C4"].value == 250.75
    assert reopened.active.tables["Records"].ref == "A3:F6"


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", " ="])
def test_excel_titles_headers_and_cells_are_never_executable(prefix):
    data = Table(title=prefix + "Title", columns=[prefix + "Column"], rows=[[prefix + "Payload"]])
    sheet = load_workbook(BytesIO(excel.build(data))).active
    assert sheet["A1"].value == data.title and sheet["A1"].data_type == "s"
    assert sheet["A3"].value == data.columns[0] and sheet["A3"].data_type == "s"
    assert sheet["A4"].value.startswith("'") and sheet["A4"].data_type == "s"
    assert not any(c.data_type == "f" or c.hyperlink for row in sheet for c in row)


@pytest.mark.parametrize(
    ("kind", "value"),
    [
        ("text", 123),
        ("number", "00123"),
        ("number", True),
        ("percent", "12.5%"),
        ("boolean", "true"),
        ("date", "08/10/2026"),
        ("date", "2026-02-30"),
        ("date", "1448-01-01"),
        ("date", "20261008"),
        ("date", "2026-W41-4"),
        ("number", 10**15),
        ("number", float("inf")),
        ("number", 0.12345678901234567),
    ],
)
def test_excel_rejects_mismatched_types_ambiguous_dates_and_precision_loss(kind, value):
    with pytest.raises(ValidationError):
        Table(title="Records", columns=["A"], rows=[[value]], column_formats=[{"kind": kind}])


def test_excel_legacy_plan_keeps_mixed_types_and_ambiguous_dates_as_text():
    table = Table(
        title="Mixed", columns=["Value"], rows=[["00123"], ["08/10/2026"], ["1448-01-01"], [12]]
    )
    sheet = load_workbook(BytesIO(excel.build(table))).active
    assert [sheet.cell(r, 1).value for r in range(4, 8)] == [
        "00123",
        "08/10/2026",
        "1448-01-01",
        12,
    ]
    assert sheet.sheet_view.rightToLeft is False


def test_excel_rejects_partial_format_list_and_sanitized_header_collisions():
    with pytest.raises(ValidationError):
        Table(title="Data", columns=["A", "B"], rows=[[1, 2]], column_formats=[{"kind": "number"}])
    for columns in (["Name", "name"], ["A", "A\x00"], ["\x00"]):
        with pytest.raises(ValueError):
            excel.build(Table(title="Data", columns=columns, rows=[["X"] * len(columns)]))


def test_excel_large_table_retains_all_rows_and_long_notes_without_extra_sheets():
    note = "ملاحظة مفصلة " * 50
    table = Table(
        title="Records",
        columns=["ID", "Notes"],
        rows=[[f"{i:05d}", note if i in (0, 999) else "ملاحظة"] for i in range(1000)],
    )
    sheet = load_workbook(BytesIO(excel.build(table))).active
    assert sheet.max_row == 1003 and sheet.tables["Records"].ref == "A3:B1003"
    assert sheet["A1003"].value == "00999" and sheet["B1003"].value == note.rstrip()
    assert sheet.row_dimensions[4].height > 120
    assert all(c.alignment.wrap_text for row in sheet.iter_rows(min_row=4) for c in row)
