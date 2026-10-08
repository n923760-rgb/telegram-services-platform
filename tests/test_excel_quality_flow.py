"""Excel quality failures retain the existing one-repair and settlement contract."""

import json
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.db import sessions
from app.core.models import Job, Order
from app.orders.engine import submit
from app.wallet.ledger import balance
from app.workers.runner import deliver_failures, execute_job
from tests.test_excel_quality import typed_table
from tests.test_foundation import job_for
from tests.test_real_services import setup


@pytest.mark.parametrize("repair", [False, True])
async def test_typed_excel_delivery_and_single_capture(tmp_path, monkeypatch, repair):
    plan = {"table": typed_table().model_dump(mode="json")}
    responses = [json.dumps(plan)]
    if repair:
        invalid = json.loads(responses[0])
        invalid["table"]["rows"][0][3] = "08/10/2026"
        responses.insert(0, json.dumps(invalid))
    _, provider, delivery, ctx, price = await setup(
        "text_to_office", responses, tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_office", {"text": "بيانات", "target": "excel"}, price, "typed-excel"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == len(responses) and len(delivery.files) == 1
    sheet = load_workbook(BytesIO(delivery.files[0][1])).active
    assert sheet.tables["Records"].ref == "A3:F6"
    assert sheet["A4"].value == "00123" and sheet["D4"].value == datetime(2026, 10, 8)
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("builder_failure", [False, True])
async def test_excel_invalid_plan_or_builder_failure_releases_without_delivery(
    tmp_path, monkeypatch, builder_failure
):
    invalid = {"table": typed_table().model_dump(mode="json")}
    if builder_failure:
        invalid["table"]["columns"][:2] = ["Name", "name"]
    else:
        invalid["table"]["rows"][0][3] = "2026-02-30"
    responses = [json.dumps(invalid)] * (1 if builder_failure else 2)
    _, provider, delivery, ctx, price = await setup(
        "text_to_office", responses, tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_office", {"text": "بيانات", "target": "excel"}, price, "bad-excel"
    )
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == len(responses) and not delivery.files and len(delivery.errors) == 1


async def test_excel_delivery_retry_uses_cached_native_table_without_second_ai_call(
    tmp_path, monkeypatch
):
    _, provider, delivery, ctx, price = await setup(
        "text_to_office",
        [json.dumps({"table": typed_table().model_dump(mode="json")})],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(
        1, "text_to_office", {"text": "بيانات", "target": "excel"}, price, "retry-excel"
    )
    jid = await job_for(oid)
    original_send = delivery.send

    async def unavailable(*args):
        raise ConnectionError("synthetic unavailable delivery")

    monkeypatch.setattr(delivery, "send", unavailable)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "delivering"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, price)
    monkeypatch.setattr(delivery, "send", original_send)
    async with sessions.begin() as db:
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC) - timedelta(seconds=1)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == 1 and len(delivery.files) == 1
    assert load_workbook(BytesIO(delivery.files[0][1])).active.tables["Records"].ref == "A3:F6"
