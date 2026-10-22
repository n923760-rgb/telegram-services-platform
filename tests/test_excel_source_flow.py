import json
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import CostHold, CostUsage, Job, Ledger, Order
from app.orders.engine import confirm_structure, submit
from app.providers.ai.gateway import AI
from app.providers.storage import OwnedStorage
from app.services.base import ServiceError
from app.wallet.ledger import balance
from app.workers.runner import deliver_confirmations, deliver_failures, execute_job
from tests.excel_source_fixtures import inputs, plan
from tests.test_foundation import job_for
from tests.test_real_services import setup


async def prepare(tmp_path, monkeypatch, responses):
    store, provider, delivery, ctx, price = await setup(
        "text_to_office", responses, tmp_path, monkeypatch
    )
    ctx["runtime_factory"] = lambda jid, uid: SimpleNamespace(
        ai=AI(provider, jid, uid, quality_review=True), storage=OwnedStorage(store, uid)
    )
    return store, provider, delivery, ctx, price


@pytest.mark.parametrize("repair", [False, True])
async def test_reviewed_source_workbook_single_capture_and_recorded_repair_cost(
    tmp_path, monkeypatch, repair
):
    responses = [json.dumps(plan()), '{"accepted":true,"issues":[]}']
    if repair:
        invalid = deepcopy(plan())
        invalid["table"]["rows"][0][2] = 999.50
        responses.insert(0, json.dumps(invalid))
    _, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, responses)
    oid = await submit(1, "text_to_office", inputs(), price, "source-excel")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004") * len(responses)
        assert await db.scalar(select(func.count(CostUsage.id))) == len(responses)
        assert await db.scalar(select(func.count(Ledger.id)).where(Ledger.kind == "capture")) == 1
        assert (await balance(db, 1)).reserved == 0
    book = load_workbook(BytesIO(delivery.files[0][1]))
    assert book.sheetnames == ["Data", "Source"] and book["Data"]["G4"].value == "L001"
    assert book["Data"]["C4"].number_format == "#,##0.00"
    assert provider.calls == len(responses) and len(delivery.files) == 1
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_bad_source_twice_releases_once_and_keeps_provider_usage(tmp_path, monkeypatch):
    invalid = plan()
    invalid["evidence"][0]["quotes"][1] = "Example"
    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(invalid)] * 2
    )
    oid = await submit(1, "text_to_office", inputs(), price, "bad-source")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        assert (await db.get(Job, jid)).cost_sar == Decimal(".008")
        assert await db.scalar(select(func.count(CostUsage.id))) == 2
        assert await db.scalar(select(func.count(Ledger.id)).where(Ledger.kind == "release")) == 1
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 2 and not delivery.files and len(delivery.errors) == 1


async def test_confirmed_source_plan_is_revalidated_and_does_not_recall_ai(tmp_path, monkeypatch):
    candidate = plan() | {"ambiguous": True, "question": "هل تعتمد الأعمدة المقترحة؟"}
    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(candidate), '{"accepted":true,"issues":[]}']
    )
    oid = await submit(1, "text_to_office", inputs(), price, "source-confirm")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_confirmations(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
        assert (await balance(db, 1)).reserved == price
    await confirm_structure(oid, 1, True)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    assert provider.calls == 2 and len(delivery.files) == 1
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"


@pytest.mark.parametrize("text", ["\n".join(["x"] * 201), "x" * 4001])
async def test_invalid_source_never_creates_order_hold_or_reservation(tmp_path, monkeypatch, text):
    _, provider, _, _, price = await prepare(tmp_path, monkeypatch, [])
    with pytest.raises(ServiceError, match="input_invalid"):
        await submit(1, "text_to_office", {"target": "excel", "text": text}, price, "limit")
    async with sessions() as db:
        assert await db.scalar(select(func.count(Order.id))) == 0
        assert await db.scalar(select(func.count(CostHold.id))) == 0
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 0
