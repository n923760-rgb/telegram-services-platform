from copy import deepcopy
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.builders.excel_source import build
from app.services.base import ServiceError
from app.services.text_to_office.excel_source import RowEvidence, source_excel_schema, value_matches
from app.services.text_to_office.schema import ExcelPlan, Inputs
from app.services.text_to_office.service import TextToOffice
from tests.excel_source_fixtures import inputs, plan


@pytest.fixture(autouse=True)
def database():
    yield


def planned():
    values = Inputs.parse(inputs())
    return values, source_excel_schema(values).model_validate(plan())


def test_source_records_and_native_precision_zero_duplicates_and_refs_survive_edit():
    values, result = planned()
    data = build(
        result.table,
        values.lines(),
        [e.source_id for e in result.evidence],
        [e.quotes for e in result.evidence],
    )
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Data", "Source"]
    sheet = book["Data"]
    assert sheet.tables["Records"].ref == "A3:G6"
    assert [sheet.cell(r, 7).value for r in range(4, 7)] == ["L001", "L002", "L003"]
    assert sheet["A4"].value == "00123" and sheet["A4"].data_type == "s"
    assert sheet["C4"].value == 125.5 and sheet["C4"].number_format == "#,##0.00"
    assert sheet["D4"].value == datetime(2026, 10, 8)
    assert sheet["E4"].value == 0.125 and sheet["E4"].number_format == "0.0%"
    assert sheet["C5"].value == 0 and sheet["D5"].value is None
    assert sheet["F5"].value is False
    assert [book["Source"].cell(r, 3).value for r in range(4, 7)] == values.lines()
    # Simulate a native row sort: source references travel with the complete row.
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    for r, row in enumerate(reversed(rows), 4):
        for c, value in enumerate(row, 1):
            sheet.cell(r, c).value = value
    saved = BytesIO()
    book.save(saved)
    reopened = load_workbook(BytesIO(saved.getvalue()))
    assert reopened["Data"]["G4"].value == "L003"
    assert reopened["Source"]["A6"].value == "L003"
    assert reopened["Data"].tables["Records"].ref == "A3:G6"


@pytest.mark.parametrize(
    "change",
    [
        "value",
        "cross_row",
        "missing_number",
        "missing_quote",
        "zero_for_null",
        "duplicate_id",
        "reverse",
        "forgot_context",
        "invent_label",
        "rounded",
    ],
)
def test_source_guard_rejects_cross_record_or_invented_data(change):
    data = deepcopy(plan())
    row = data["table"]["rows"][0]
    proof = data["evidence"][0]
    if change == "value":
        row[2] = 999.50
    elif change == "cross_row":
        row[1] = "Example"
        proof["quotes"][1] = "Example"
    elif change == "missing_number":
        row[2] = None
        proof["quotes"][2] = None
    elif change == "missing_quote":
        proof["quotes"][1] = None
    elif change == "zero_for_null":
        data["table"]["rows"][1][3] = "2026-10-08"
    elif change == "duplicate_id":
        data["evidence"][1]["source_id"] = 1
    elif change == "reverse":
        data["evidence"].reverse()
    elif change == "forgot_context":
        data["evidence"].pop()
        data["table"]["rows"].pop()
    elif change == "invent_label":
        data["table"]["title"] = "الطلبات 99999"
    elif change == "rounded":
        row[2] = 125.51
    with pytest.raises(ValidationError):
        source_excel_schema(Inputs.parse(inputs())).model_validate(data)


@pytest.mark.parametrize(
    "value,quote,kind,header,accepted",
    [
        (0.125, "12.5%", "percent", "Rate", True),
        (0.125, "0.125", "percent", "Rate", False),
        (125.5, "125.50", "number", "Amount", True),
        (123, "00123", "number", "Amount", False),
        (123, "123", "number", "ID", False),
        (966501234567, "966501234567", "number", "Phone", False),
        (True, "نعم", "boolean", "Confirmed", True),
        (False, "0", "boolean", "Confirmed", False),
        (0, "-0", "number", "Amount", False),
        (1.2, "1.20", "number", "Amount", True),
        (0, "0.000000000000000", "number", "Amount", False),
        ("2026-10-08", "2026/10/08", "date", "Date", True),
        ("2026-10-08", "08/10/2026", "date", "Date", False),
        ("2026-02-30", "2026/02/30", "date", "Date", False),
        ("08/10/2026", "08/10/2026", "text", "Date", True),
        ("00123", "00123", "text", "ID", True),
        ("٥٠", "٥٠", "text", "Amount", True),
        (50, "٥٠", "number", "Amount", False),
    ],
)
def test_only_defined_literal_type_conversions(value, quote, kind, header, accepted):
    assert value_matches(value, quote, kind, header) is accepted


async def test_injected_and_cached_plans_revalidated_without_storage_write():
    bad = plan()
    bad["table"]["rows"][0][2] = 999.5
    # Base plan cannot smuggle an unvalidated value through the injected provider.
    base = ExcelPlan.model_validate({"table": bad["table"]})
    candidate = source_excel_schema(Inputs.parse(inputs())).model_construct(
        table=base.table,
        evidence=[RowEvidence.model_validate(e) for e in bad["evidence"]],
        context_ids=[],
    )
    service = TextToOffice()
    store = SimpleNamespace(save=AsyncMock())
    service.runtime = SimpleNamespace(
        storage=store, ai=SimpleNamespace(extract=AsyncMock(return_value=candidate))
    )
    for cached in (False, True):
        request = inputs() | ({"__continuation": bad} if cached else {})
        with pytest.raises(ServiceError, match="provider_invalid"):
            await service.run(request)
    store.save.assert_not_called()


@pytest.mark.parametrize("text", ["\n".join(["x"] * 201), "x" * 4001, "bad\x00"])
def test_source_limits_rejected_before_admission(text):
    with pytest.raises(ServiceError, match="input_invalid"):
        TextToOffice.needs_ai({"target": "excel", "text": text})


def test_literal_source_parts_reconstruct_long_formula_like_source_without_formulas():
    source = "=SUM(1,2) " + "W" * 600
    table = ExcelPlan.model_validate(
        {"table": {"title": "Literal", "columns": ["Text"], "rows": [[source]]}}
    ).table
    book = load_workbook(BytesIO(build(table, [source], [1], [[source]])))
    sheet = book["Source"]
    assert book["Data"]["A4"].value == source and book["Data"]["A4"].data_type == "s"
    assert "".join(sheet.cell(r, 3).value for r in range(4, sheet.max_row + 1)) == source
    assert all(sheet.cell(r, 1).value == "L001" for r in range(4, sheet.max_row + 1))
    assert all(row.height <= 409 for row in sheet.row_dimensions.values())
    assert not any(
        cell.data_type == "f" or cell.hyperlink for tab in book for row in tab for cell in row
    )


def test_context_is_explicit_and_original_lines_retained():
    source = inputs() | {"text": "اعتمد الحقول الموجودة\n" + inputs()["text"]}
    data = plan()
    for row in data["evidence"]:
        row["source_id"] += 1
    data["context_ids"] = [1]
    values = Inputs.parse(source)
    result = source_excel_schema(values).model_validate(data)
    assert result.context_ids == [1] and [e.source_id for e in result.evidence] == [2, 3, 4]


def test_explicit_hijri_context_cannot_be_normalized_as_gregorian():
    source = inputs() | {"text": "التواريخ هجري\n" + inputs()["text"]}
    data = plan()
    for row in data["evidence"]:
        row["source_id"] += 1
    data["context_ids"] = [1]
    with pytest.raises(ValidationError):
        source_excel_schema(Inputs.parse(source)).model_validate(data)


def test_contact_occurrence_cannot_be_removed_or_changed():
    values = Inputs.parse({"target": "excel", "text": "تواصل وsample@example.invalid"})
    data = {
        "table": {"title": "التواصل", "columns": ["البريد"], "rows": [["sample@example.invalid"]]},
        "evidence": [{"source_id": 1, "quotes": ["sample@example.invalid"]}],
    }
    source_excel_schema(values).model_validate(data)
    data["table"]["rows"][0][0] = "other@example.invalid"
    with pytest.raises(ValidationError):
        source_excel_schema(values).model_validate(data)
