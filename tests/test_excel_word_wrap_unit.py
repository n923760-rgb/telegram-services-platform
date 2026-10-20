"""Prevent clipped whole-word paragraphs in native Excel cells."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.builders import excel
from app.builders.schema import Table


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("word", ["WWWWWWWWWWWWWWWWWWWW", "بياناتمحاسبيةماليةمفصلةللعملاء"])
def test_whole_words_keep_unused_line_space_in_native_row_height(word):
    # Each word fits a 60-unit cell but two words cannot share a line.
    width = 60
    word_pixels = excel._text_pixels(word) * 1.12
    available = width * 7 - 12
    assert available / 2 < word_pixels <= available
    text = " ".join([word] * 6)
    book = load_workbook(
        BytesIO(excel.build(Table(title="Notes", columns=["Details"], rows=[[text]])))
    )
    sheet = book.active
    assert sheet["A4"].value == text
    assert sheet.column_dimensions["A"].width == width
    assert sheet.row_dimensions[4].height >= 6 * 16 + 6
    assert sheet["A4"].alignment.wrap_text


def test_wrap_preserves_explicit_blank_lines_and_unbroken_values(monkeypatch):
    monkeypatch.setattr(excel, "_measurement_font", lambda *args: None)
    assert excel._wrapped_lines("first\n\nlast", 60) == 3
    assert excel._wrapped_lines("W" * 200, 60) >= 4
    assert excel._wrapped_lines(0, 60) == 1


def test_wrapping_keeps_excel_height_limit_and_all_long_source_text():
    text = "WWWWWWWWWWWWWWWWWWWW " * 150
    sheet = load_workbook(
        BytesIO(excel.build(Table(title="Notes", columns=["Details"], rows=[[text]])))
    ).active
    assert sheet["A4"].value == text.rstrip()
    assert sheet.row_dimensions[4].height == 409
