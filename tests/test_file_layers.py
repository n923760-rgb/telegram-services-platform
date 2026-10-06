import json
from decimal import Decimal
from io import BytesIO

import httpx
import pytest
from docx import Document as Word
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pydantic import ValidationError
from sqlalchemy import select

from app.builders import excel, pdf, pptx, word
from app.builders.schema import Deck, Document, Section, Slide, Table
from app.core.db import sessions
from app.core.models import CostUsage, Job
from app.core.settings import config
from app.files.validation import image
from app.providers.ai.base import AIProvider, AIResponse, Usage
from app.providers.ai.gateway import AI
from app.providers.ai.openai_compatible import OpenAICompatible
from app.providers.storage import LocalStorage
from app.services.base import ServiceError
from tests.test_operations import running


class Provider(AIProvider):
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = 0

    def upper_bound(self, *args):
        return Decimal(".01")

    async def extract_from_text(self, *args):
        self.calls += 1
        return AIResponse(self.responses.pop(0), Usage(10, 20, Decimal(".004")))

    async def extract_from_image(self, *args):
        return await self.extract_from_text(*args)

    async def translate(self, *args):
        return await self.extract_from_text(*args)


async def test_ai_schema_retry_once_and_failed_cost_recorded():
    _, jid = await running()
    provider = Provider(["not json", '{"title":"وثيقة","sections":[{"paragraphs":["مرحبًا"]}]}'])
    result = await AI(provider, jid, 1).extract(Document, "input")
    assert result.title == "وثيقة" and provider.calls == 2
    async with sessions() as db:
        assert (await db.get(Job, jid)).cost_sar == Decimal(".008")
        assert len((await db.scalars(select(CostUsage))).all()) == 2


async def test_ai_invalid_twice_clean_error():
    _, jid = await running()
    provider = Provider(["{}", "{}", "unused"])
    with pytest.raises(ServiceError, match="provider_invalid"):
        await AI(provider, jid, 1).extract(Document, "input")
    assert provider.calls == 2
    async with sessions() as db:
        assert (await db.get(Job, jid)).cost_sar == Decimal(".008")


def test_office_files_open_and_contain_expected_text():
    data = Document(
        title="تقرير Report",
        sections=[Section(heading="العنوان", paragraphs=["مرحبًا 123"], bullets=["نقطة"])],
    )
    doc = Word(BytesIO(word.build(data)))
    assert "مرحبًا 123" in [p.text for p in doc.paragraphs]
    assert "w:bidi" in doc.element.xml
    table = Table(
        title="بيانات",
        columns=["الاسم", "القيمة"],
        rows=[["منتج", 12], ['=HYPERLINK("evil")', "00123"]],
    )
    workbook = load_workbook(BytesIO(excel.build(table)))
    assert workbook.active.sheet_view.rightToLeft
    assert workbook.active["B2"].value == 12
    assert workbook.active["A3"].data_type == "s" and workbook.active["A3"].value.startswith("'")
    assert workbook.active["B3"].value == "00123"
    deck = Deck(slides=[Slide(title="عنوان", bullets=["محتوى"])])
    presentation = Presentation(BytesIO(pptx.build(deck)))
    assert any(
        shape.has_text_frame and "محتوى" in shape.text for shape in presentation.slides[0].shapes
    )
    assert all(
        shape.left + shape.width <= presentation.slide_width
        and shape.top + shape.height <= presentation.slide_height
        for shape in presentation.slides[0].shapes
    )


def test_pdf_opens_contains_content_and_multiple_pages():
    import pdfplumber

    data = Document(title="Report تقرير", sections=[Section(paragraphs=["Arabic اختبار 123"] * 60)])
    with pdfplumber.open(BytesIO(pdf.build(data))) as doc:
        assert len(doc.pages) >= 2
        assert "123" in "".join(page.extract_text() or "" for page in doc.pages)


@pytest.mark.parametrize("rows", [[["short"]], [[1, 2, 3]]])
def test_table_rejects_ragged_rows(rows):
    with pytest.raises(ValidationError):
        Table(title="x", columns=["a", "b"], rows=rows)


def test_file_limits_blank_images_and_path_ownership(tmp_path):
    with pytest.raises(ServiceError):
        image(b"not an image", 100)
    blank = BytesIO()
    Image.new("RGB", (100, 100), "white").save(blank, format="PNG")
    with pytest.raises(ServiceError, match="image_unclear"):
        image(blank.getvalue(), 10000)
    original = Image.new("RGB", (3000, 2000), "white")
    ImageDraw.Draw(original).text((20, 20), "TEST TEXT", fill="black")
    buf = BytesIO()
    original.save(buf, format="PNG")
    compressed = image(buf.getvalue(), 1000000)
    assert max(Image.open(BytesIO(compressed)).size) <= 1024
    store = LocalStorage(tmp_path, 100)
    file = store.save(1, "result.txt", b"hello", "text/plain")
    assert store.read(file.key, 1) == b"hello"
    with pytest.raises(ServiceError, match="not_allowed"):
        store.read(file.key, 2)
    with pytest.raises(ServiceError):
        store.path("../../etc/passwd")
    store.delete(file.key, 1)
    with pytest.raises(ServiceError, match="file_expired"):
        store.read(file.key, 1)


async def test_concrete_provider_contract_and_usage(monkeypatch):
    cfg = config()
    monkeypatch.setattr(cfg, "ai_enabled", True)
    from pydantic import SecretStr

    monkeypatch.setattr(cfg, "ai_api_key", SecretStr("test-secret"))
    monkeypatch.setattr(cfg, "ai_input_usd_per_million", Decimal(".4"))
    monkeypatch.setattr(cfg, "ai_output_usd_per_million", Decimal("1.6"))

    def request(req):
        body = json.loads(req.content)
        assert body["response_format"]["type"] == "json_object"
        assert req.headers["authorization"] == "Bearer test-secret"
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": '{"ok":true}'}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 20},
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(request)) as client:
        result = await OpenAICompatible(client).extract_from_text({}, "hello", "extract")
    assert result.usage.cost_sar == Decimal(".000270")
    assert result.usage.input_tokens == 100


@pytest.mark.parametrize(
    "body",
    [
        {"choices": [], "usage": {"prompt_tokens": True, "completion_tokens": 1}},
        {"choices": [], "usage": []},
        ["unexpected"],
    ],
)
async def test_provider_rejects_malformed_usage_cleanly(monkeypatch, body):
    from pydantic import SecretStr

    cfg = config()
    monkeypatch.setattr(cfg, "ai_enabled", True)
    monkeypatch.setattr(cfg, "ai_api_key", SecretStr("fixture"))
    monkeypatch.setattr(cfg, "ai_input_usd_per_million", Decimal("1"))
    monkeypatch.setattr(cfg, "ai_output_usd_per_million", Decimal("1"))
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(lambda req: httpx.Response(200, json=body))
    ) as client:
        with pytest.raises(ServiceError, match="provider_usage"):
            await OpenAICompatible(client).extract_from_text({}, "input", "extract")


def test_file_validation_rejects_forged_pdf_and_accepts_real_pdf_wav():
    import wave

    from app.files.validation import validate_file

    with pytest.raises(ServiceError, match="file_invalid"):
        validate_file(b"%PDF-1.7\nfake", "application/pdf", 10000)
    rendered = pdf.build(Document(title="ملف", sections=[Section(paragraphs=["123"])]))
    assert validate_file(rendered, "application/pdf", 100000) == rendered
    wav = BytesIO()
    with wave.open(wav, "wb") as sound:
        sound.setnchannels(1)
        sound.setsampwidth(2)
        sound.setframerate(8000)
        sound.writeframes(b"\x00\x00" * 100)
    assert validate_file(wav.getvalue(), "audio/wav", 10000) == wav.getvalue()


async def test_ai_service_enable_requires_valid_provider_configuration(monkeypatch):
    from app.core.models import Service
    from app.ops.admin import set_service

    monkeypatch.setattr(config(), "ai_enabled", False)
    with pytest.raises(ServiceError, match="provider_config"):
        await set_service("image_to_text", enabled=True)
    async with sessions() as db:
        assert not (await db.get(Service, "image_to_text")).enabled


@pytest.mark.parametrize(
    "values",
    [
        {"max_concurrent_jobs_per_user": 0},
        {"ai_image_token_bound": 0},
        {"circuit_min_samples": 11},
        {"circuit_failure_threshold": 1.1},
        {"ai_max_output_tokens": 40000},
        {"ai_enabled": True},
    ],
)
def test_configuration_rejects_invalid_execution_bounds(values):
    from app.core.settings import Config

    with pytest.raises(ValidationError):
        Config(_env_file=None, **values)


async def test_telegram_delivery_uses_storage_interface_and_utf16_chunks(tmp_path):
    from unittest.mock import AsyncMock

    from aiogram.types import BufferedInputFile

    from app.providers.delivery import TelegramDelivery
    from app.services.base import Result

    store = LocalStorage(tmp_path)
    file = store.save(1, "result.txt", b"content", "text/plain")
    bot = AsyncMock()
    text = "🙂" * 6000
    await TelegramDelivery(bot, store).send(1, "order", Result(text=text, artifacts=[file]))
    chunks = [call.args[1] for call in bot.send_message.await_args_list[1:]]
    assert "".join(chunks) == text
    assert all(len(chunk.encode("utf-16-le")) // 2 <= 3500 for chunk in chunks)
    document = bot.send_document.await_args.args[1]
    assert isinstance(document, BufferedInputFile) and document.data == b"content"
