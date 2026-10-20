from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pydantic import ValidationError

from app.builders.word import build
from app.core.i18n import tr
from app.services.meeting_minutes.schema import Inputs, classification_schema, document
from app.services.meeting_minutes.service import MeetingMinutes
from tests.minutes_fixtures import inputs, plan


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("language", ["ar", "en"])
def test_natural_paragraph_preserves_decimals_ids_and_dates(language):
    values = Inputs.parse(
        inputs(language) | {"notes": inputs(language)["notes"].replace("\n", " ")}
    )
    assert values.lines() == Inputs.parse(inputs(language)).lines()
    assert "125.50" in values.lines()[0] and "00017" in values.lines()[2]
    assert MeetingMinutes.version == "4"


@pytest.mark.parametrize(
    "language,source,expected",
    [
        (
            "ar",
            "د. أحمد يرسل التقرير 00017 بتاريخ 2026-10-12. اعتمد الطلب 00123.",
            ["د. أحمد يرسل التقرير 00017 بتاريخ 2026-10-12.", "اعتمد الطلب 00123."],
        ),
        (
            "en",
            "Dr. Sam sends report 00017 on 2026-10-12. Approved order 00123.",
            ["Dr. Sam sends report 00017 on 2026-10-12.", "Approved order 00123."],
        ),
        (
            "en",
            "Prof. A. B. Smith sends the report. Budget is 125.50.",
            ["Prof. A. B. Smith sends the report.", "Budget is 125.50."],
        ),
        (
            "ar",
            "م. سام يراجع الطلب. أ. أحمد يرسل التقرير؛ اعتمد الموعد؟ نراجع الميزانية!",
            ["م. سام يراجع الطلب.", "أ. أحمد يرسل التقرير؛", "اعتمد الموعد؟", "نراجع الميزانية!"],
        ),
        (
            "en",
            "Use e.g. report 00123 for the U.S. team. Send it at 10 a.m. tomorrow.",
            ["Use e.g. report 00123 for the U.S. team.", "Send it at 10 a.m. tomorrow."],
        ),
        (
            "en",
            "1. Dr. Sam sends the report.\n2. Mr. Lee reviews it.",
            ["1. Dr. Sam sends the report.", "2. Mr. Lee reviews it."],
        ),
        (
            "ar",
            "١. د. أحمد يرسل التقرير.\n٢. سام يراجع الطلب.",
            ["١. د. أحمد يرسل التقرير.", "٢. سام يراجع الطلب."],
        ),
        (
            "en",
            "Unknown xyz. Keep ordinary sentence boundaries.",
            ["Unknown xyz.", "Keep ordinary sentence boundaries."],
        ),
        (
            "en",
            "U.S.\nExplicit newline is authoritative.",
            ["U.S.", "Explicit newline is authoritative."],
        ),
        (
            "en",
            "Dr. Sam sends the report.\t  Dr. Sam sends the report.",
            ["Dr. Sam sends the report.", "Dr. Sam sends the report."],
        ),
        (
            "ar",
            "أحمد (د. سام) يراجع التقرير. اعتمد الطلب.",
            ["أحمد (د. سام) يراجع التقرير.", "اعتمد الطلب."],
        ),
    ],
)
def test_name_abbreviation_and_list_boundaries_preserve_source(language, source, expected):
    assert Inputs.parse(inputs(language) | {"notes": source}).lines() == expected


@pytest.mark.parametrize("language,name", [("ar", "د. أحمد"), ("en", "Dr. Sam")])
async def test_titled_owner_stays_in_one_source_unit_through_approval_and_word(language, name):
    note = f"{name} sends report 00017 on 2026-10-12."
    source = inputs(language) | {"notes": note + " Approved budget 125.50."}
    candidate = {
        "assignments": [
            {"note_id": 1, "category": "action", "assignee": name, "deadline": "2026-10-12"},
            {"note_id": 2, "category": "decision"},
        ]
    }

    class AI:
        async def extract(self, schema, serialized, prompt):
            import json

            assert json.loads(serialized)["notes"][0] == {"id": 1, "text": note}
            return schema.model_validate(candidate)

    service = MeetingMinutes()
    service.runtime = SimpleNamespace(ai=AI(), storage=None, renderer=None)
    result = await service.run(source)
    assert result.needs_confirmation and not result.artifacts
    assert name in result.preview and "2026-10-12" in result.preview
    values = Inputs.parse(source)
    classified = classification_schema(values.lines()).model_validate(result.continuation)
    doc = Document(BytesIO(build(document(values, classified), professional_template=True)))
    assert sum(p.text == f"[1] {note}" for p in doc.paragraphs) == 1
    assert [cell.text for cell in doc.tables[0].rows[1].cells] == ["[1]", name, "2026-10-12"]


@pytest.mark.parametrize("notes", ["Dr. " + "x" * 1197, "Task done. " * 81])
def test_boundary_protection_does_not_bypass_unit_limits(notes):
    from app.services.base import ServiceError

    with pytest.raises(ServiceError, match="minutes_input_invalid"):
        Inputs.parse(inputs("en") | {"notes": notes})


@pytest.mark.parametrize(
    "field,value",
    [
        ("assignee", "Invented person"),
        ("deadline", "2026-10-99"),
        ("deadline", "12"),
        ("deadline", "2026/10/12"),
        ("assignee", True),
        ("assignee", "Sam\n"),
    ],
)
def test_metadata_requires_same_source_span_and_whole_numeric_tokens(field, value):
    values = Inputs.parse(inputs("en"))
    candidate = plan()
    candidate["assignments"][2][field] = value
    with pytest.raises(ValidationError):
        classification_schema(values.lines()).model_validate(candidate)


def test_task_metadata_is_for_actions_only_and_request_local():
    candidate = plan()
    candidate["assignments"][0]["deadline"] = "00123"
    with pytest.raises(ValidationError):
        classification_schema(Inputs.parse(inputs("en")).lines()).model_validate(candidate)
    candidate = plan()
    candidate["assignments"][2].update(assignee="Sam", deadline="2026-10-12")
    schema = classification_schema(Inputs.parse(inputs("en")).lines())
    assert schema.model_validate(candidate)
    assert "Sam" not in str(schema.model_json_schema())
    with pytest.raises(ValidationError):
        classification_schema(Inputs.parse(inputs("ar")).lines()).model_validate(candidate)


@pytest.mark.parametrize("language", ["ar", "en"])
def test_native_word_register_has_exact_metadata_and_verbatim_source(language):
    values = Inputs.parse(inputs(language))
    candidate = plan()
    name = "سام" if language == "ar" else "Sam"
    candidate["assignments"][2].update(assignee=name, deadline="2026-10-12")
    classified = classification_schema(values.lines()).model_validate(candidate)
    doc = Document(BytesIO(build(document(values, classified), professional_template=True)))
    for index, note in enumerate(values.lines(), 1):
        assert sum(p.text == f"[{index}] {note}" for p in doc.paragraphs) == 1
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows][1] == [
        "[3]",
        name,
        "2026-10-12",
    ]
    missing = classification_schema(values.lines()).model_validate(plan())
    assert document(values, missing).sections[2].tables[0].rows[0][1:] == [
        tr("minutes_unspecified", language),
        tr("minutes_unspecified", language),
    ]


async def test_review_metadata_is_bounded_and_exports_nothing_before_confirmation(tmp_path):
    candidate = plan()
    candidate["assignments"][2].update(assignee="Sam", deadline="2026-10-12")

    class AI:
        async def extract(self, schema, *args):
            return schema.model_validate(deepcopy(candidate))

    service = MeetingMinutes()
    service.runtime = SimpleNamespace(ai=AI(), storage=None, renderer=None)
    result = await service.run(inputs("en"))
    assert result.needs_confirmation and not result.artifacts
    assert "Sam" in result.preview and "2026-10-12" in result.preview
    assert result.continuation["assignments"][2]["deadline"] == "2026-10-12"
    assert len(result.preview.encode("utf-16-le")) // 2 < 4096
