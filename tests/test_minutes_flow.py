import json
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import Job, Order, Service
from app.core.settings import config
from app.orders.engine import confirm_structure, submit
from app.services.base import ServiceError
from app.wallet.ledger import balance
from app.workers.runner import deliver_confirmations, deliver_failures, execute_job
from tests.minutes_fixtures import inputs, plan
from tests.test_foundation import job_for
from tests.test_navigation import flow  # noqa: F401
from tests.test_real_services import setup
from tests.word_pdf_fixtures import renderer


async def prepare(tmp_path, monkeypatch, responses):
    store, provider, delivery, ctx, price = await setup(
        "meeting_minutes", responses, tmp_path, monkeypatch
    )
    original = ctx["runtime_factory"]

    def factory(jid, uid):
        runtime = original(jid, uid)
        return SimpleNamespace(ai=runtime.ai, storage=runtime.storage, renderer=renderer())

    ctx["runtime_factory"] = factory
    return store, provider, delivery, ctx, price


async def review(ctx, oid):
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_confirmations(ctx)
    await deliver_confirmations(ctx)
    return jid


@pytest.mark.parametrize("language", ["ar", "en"])
async def test_telegram_title_skip_keeps_language_and_source_until_confirmation(
    flow,  # noqa: F811
    tmp_path,
    monkeypatch,
    language,
):
    from app.core.i18n import tr
    from app.core.users import set_language
    from tests.test_bounded_fixes import button_by_prefix

    transport, message, callback, state = flow
    _, provider, _, _, price = await prepare(tmp_path, monkeypatch, [])
    await set_language(1, language)
    await message("/start")
    await callback("service:meeting_minutes")
    assert transport.messages[-1].text.endswith(tr("minutes_title", language))
    skip = button_by_prefix(transport.messages[-1].reply_markup, "skip:")
    draft = await state.get_data()
    await message("Title\nnot a single line")
    assert transport.messages[-1].text == tr("input_single_line", language)
    assert await state.get_data() == draft
    await callback(skip)
    assert transport.messages[-1].text.endswith(tr("minutes_language", language))
    choice = button_by_prefix(transport.messages[-1].reply_markup, "choice:")
    await callback(choice.rsplit(":", 1)[0] + (":0" if language == "ar" else ":1"))
    source = inputs(language)["notes"]
    await message(source)
    await callback(button_by_prefix(transport.messages[-1].reply_markup, "confirm:"))
    async with sessions() as db:
        order = await db.scalar(select(Order))
        assert order.inputs == {"language": language, "notes": source}
        assert order.service_version == "4"
        assert (await balance(db, 1)).reserved == price
    assert provider.calls == 0


@pytest.mark.parametrize("language", ["ar", "en"])
async def test_order_details_reopen_all_tasks_in_current_ui_language(
    flow,  # noqa: F811
    tmp_path,
    monkeypatch,
    language,
):
    from app.core.i18n import tr
    from app.core.users import set_language

    transport, message, callback, _ = flow
    owner = "المسؤول الكامل " + "أحمد " * 12
    deadline = "2026-10-12 only after written approval from the full project committee"
    notes = [f"{owner} sends report {i:05d} by {deadline}." for i in range(1, 13)]
    candidate = {
        "assignments": [
            {"note_id": i, "category": "action", "assignee": owner, "deadline": deadline}
            for i in range(1, 13)
        ]
    }
    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(candidate)]
    )
    oid = await submit(
        1, "meeting_minutes", inputs() | {"notes": "\n".join(notes)}, price, "full-review"
    )
    await execute_job(ctx, str(await job_for(oid)))
    await set_language(1, language)
    await message("/start")
    before = len(transport.messages)
    await callback(f"order:{oid}")
    rendered = transport.messages[before:]
    assert len(rendered) > 1
    text = "".join(msg.text for msg in rendered)
    assert tr("minutes_review_complete", language) in text
    for i, note in enumerate(notes, 1):
        assert text.count(f"[{i}] {note}") == 1
    assert text.count(tr("minutes_task_preview", language, assignee=owner, deadline=deadline)) == 12
    assert all(msg.reply_markup is None for msg in rendered[:-1])
    assert all(len(msg.text.encode("utf-16-le")) // 2 <= 3500 for msg in rendered)
    assert rendered[-1].reply_markup.inline_keyboard[0][0].callback_data == f"approve:{oid}"
    assert not delivery.files and provider.calls == 1
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
        assert (await balance(db, 1)).reserved == price


@pytest.mark.parametrize("approved", [True, False])
async def test_multipart_review_notification_recovers_without_new_ai_or_early_capture(
    tmp_path, monkeypatch, approved
):
    from unittest.mock import AsyncMock

    from app.providers.delivery import TelegramDelivery

    notes = [f"Source {i:05d} " + "🙂" * 110 for i in range(1, 31)]
    candidate = {"assignments": [{"note_id": i, "category": "review"} for i in range(1, 31)]}
    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(candidate)]
    )
    oid = await submit(
        1, "meeting_minutes", inputs() | {"notes": "\n".join(notes)}, price, "review-retry"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    bot = AsyncMock()
    # The first message succeeds and the next fails; no approval button was delivered.
    bot.send_message.side_effect = [None, OSError("synthetic transport offline")]
    notifier = TelegramDelivery(bot)
    notifier._lang = AsyncMock(return_value="en")
    await deliver_confirmations({"delivery": notifier})
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == "waiting_confirmation" and not order.confirmation_notified
        assert (await balance(db, 1)).reserved == price
    assert bot.send_message.await_args_list[0].kwargs["reply_markup"] is None
    bot.reset_mock()
    bot.send_message.side_effect = None
    await deliver_confirmations({"delivery": notifier})
    await deliver_confirmations({"delivery": notifier})
    calls = bot.send_message.await_args_list
    text = "".join(call.args[1] for call in calls)
    for i, note in enumerate(notes, 1):
        assert text.count(f"[{i}] {note}") == 1
    assert all(call.kwargs["reply_markup"] is None for call in calls[:-1])
    assert calls[-1].kwargs["reply_markup"].inline_keyboard[0][0].callback_data == f"approve:{oid}"
    assert provider.calls == 1 and not delivery.files
    async with sessions() as db:
        assert (await db.get(Order, oid)).confirmation_notified
    await confirm_structure(oid, 1, approved)
    if approved:
        await execute_job(ctx, str(jid))
        await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == ("completed" if approved else "cancelled")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if approved else 1000, 0)
    assert provider.calls == 1 and len(delivery.files) == (2 if approved else 0)


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("title_mode", ["custom", "omitted", "null"])
async def test_confirmed_source_notes_native_pair_single_capture_and_provider_call(
    tmp_path, monkeypatch, language, title_mode
):
    store, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(plan())]
    )
    source = inputs(language)
    if title_mode != "custom":
        source.pop("title")
    if title_mode == "null":
        source["title"] = None
    expected_title = source.get("title") or (
        "محضر اجتماع" if language == "ar" else "Meeting minutes"
    )
    oid = await submit(1, "meeting_minutes", source, price, "minutes-native")
    jid = await review(ctx, oid)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
        assert (await balance(db, 1)).reserved == price
        assert await db.scalar(select(func.count(Order.id))) == 1
    assert not delivery.files and len(delivery.confirmations) == 1
    assert delivery.confirmations[0][1].startswith(expected_title + "\n")
    for i, note in enumerate(source["notes"].split("\n"), 1):
        assert f"[{i}] {note}" in delivery.confirmations[0][1]
    with pytest.raises(ServiceError, match="not_allowed"):
        await confirm_structure(oid, 2, True)
    await confirm_structure(oid, 1, True)
    monkeypatch.setattr(config(), "ai_enabled", False)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004")
    assert provider.calls == 1
    assert [name for name, _ in delivery.files] == ["minutes.docx", "minutes.pdf"]
    text = [
        p.text.replace("\u200e", "") for p in Document(BytesIO(delivery.files[0][1])).paragraphs
    ]
    assert text[0] == expected_title
    for index, line in enumerate(source["notes"].split("\n"), 1):
        assert text.count(f"[{index}] {line}") == 1
    pdf_text = PdfReader(BytesIO(delivery.files[1][1])).pages[0].extract_text()
    assert "00123" in pdf_text and "125.50" in pdf_text and "00017" in pdf_text
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("repair_success", [False, True])
async def test_incomplete_classification_one_repair_or_release(
    tmp_path, monkeypatch, repair_success
):
    invalid = json.dumps({"assignments": [{"note_id": 1, "category": "action"}]})
    responses = [invalid, json.dumps(plan()) if repair_success else invalid]
    _, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, responses)
    oid = await submit(1, "meeting_minutes", inputs(), price, "minutes-repair")
    await review(ctx, oid)
    if repair_success:
        await confirm_structure(oid, 1, True)
        await execute_job(ctx, str(await job_for(oid)))
    else:
        await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == ("completed" if repair_success else "failed")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (
            1000 - price if repair_success else 1000,
            0,
        )
        assert (await db.get(Job, await job_for(oid))).cost_sar == Decimal(".008")
    assert provider.calls == 2
    assert len(delivery.files) == (2 if repair_success else 0)


async def test_declined_review_releases_without_render_but_keeps_actual_ai_cost(
    tmp_path, monkeypatch
):
    _, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [json.dumps(plan())])
    oid = await submit(1, "meeting_minutes", inputs(), price, "minutes-decline")
    await review(ctx, oid)
    await confirm_structure(oid, 1, False)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        assert (await balance(db, 1)).available == 1000
        assert (await balance(db, 1)).reserved == 0
        assert (await db.get(Job, await job_for(oid))).cost_sar == Decimal(".004")
    assert provider.calls == 1 and not delivery.files


async def test_confirmed_render_failure_releases_credit(tmp_path, monkeypatch):
    from app.providers.documents.base import DocumentError
    from app.providers.documents.libreoffice import LibreOfficeDocuments

    _, _, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [json.dumps(plan())])
    oid = await submit(1, "meeting_minutes", inputs(), price, "minutes-render-fail")
    await review(ctx, oid)
    await confirm_structure(oid, 1, True)

    async def fail(*args, **kwargs):
        raise DocumentError("document_render_unavailable")

    monkeypatch.setattr(LibreOfficeDocuments, "word_pdf", fail)
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        assert (await balance(db, 1)).available == 1000
        assert (await balance(db, 1)).reserved == 0
    assert not delivery.files


@pytest.mark.parametrize("old_version", [None, "1", "2", "3"])
async def test_confirmed_delivery_retry_reuses_pair_without_ai_or_export(
    tmp_path, monkeypatch, old_version
):
    from app.providers.documents.libreoffice import LibreOfficeDocuments

    _, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [json.dumps(plan())])
    oid = await submit(1, "meeting_minutes", inputs(), price, "minutes-retry")
    jid = await review(ctx, oid)
    await confirm_structure(oid, 1, True)
    original = delivery.send

    async def offline(*args):
        raise ConnectionError("synthetic delivery failure")

    monkeypatch.setattr(delivery, "send", offline)
    await execute_job(ctx, str(jid))
    async with sessions.begin() as db:
        assert len((await db.get(Order, oid)).result["artifacts"]) == 2
        assert (await balance(db, 1)).reserved == price
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC)
        if old_version:
            (await db.get(Order, oid)).service_version = old_version

    async def no_render(*args, **kwargs):
        raise AssertionError("prepared minutes must be reused")

    monkeypatch.setattr(LibreOfficeDocuments, "word_pdf", no_render)
    monkeypatch.setattr(delivery, "send", original)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 1 and len(delivery.files) == 2
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("old_version", ["1", "2", "3"])
@pytest.mark.parametrize("approved", [False, True])
async def test_boundary_upgrade_retains_admin_settings_and_cancels_unprepared_old_work(
    tmp_path, monkeypatch, old_version, approved
):
    from app.services.registry import registry

    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(plan())] if approved else []
    )
    oid = await submit(1, "meeting_minutes", inputs(), price, "minutes-old-boundaries")
    jid = await job_for(oid)
    if approved:
        await review(ctx, oid)
        await confirm_structure(oid, 1, True)
    async with sessions.begin() as db:
        (await db.get(Order, oid)).service_version = old_version
        row = await db.get(Service, "meeting_minutes")
        row.version, row.price_halala, row.price_stars = old_version, 777, 42
    async with sessions.begin() as db:
        await registry.sync(db)
        row = await db.get(Service, "meeting_minutes")
        assert (row.version, row.enabled, row.price_halala, row.price_stars) == ("4", True, 777, 42)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004" if approved else "0")
    assert provider.calls == int(approved) and not delivery.files
    assert len(delivery.errors) == 1 and delivery.errors[0][2] == "service_changed"


@pytest.mark.parametrize("language,name", [("ar", "د. أحمد"), ("en", "Dr. Sam")])
async def test_titled_action_owner_native_export_after_approval_only(
    tmp_path, monkeypatch, language, name
):
    source = inputs(language) | {"notes": f"{name} sends report 00017 on 2026-10-12."}
    candidate = {
        "assignments": [
            {"note_id": 1, "category": "action", "assignee": name, "deadline": "2026-10-12"}
        ]
    }
    _, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(candidate)]
    )
    oid = await submit(1, "meeting_minutes", source, price, "minutes-titled-owner")
    jid = await review(ctx, oid)
    assert not delivery.files and name in delivery.confirmations[0][1]
    await confirm_structure(oid, 1, True)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    doc = Document(BytesIO(dict(delivery.files)["minutes.docx"]))
    assert [cell.text.replace("\u200e", "") for cell in doc.tables[0].rows[1].cells] == [
        "[1]",
        name,
        "2026-10-12",
    ]
    pdf = PdfReader(BytesIO(dict(delivery.files)["minutes.pdf"]))
    assert "00017" in pdf.pages[0].extract_text() and "2026-10-12" in pdf.pages[0].extract_text()
    assert provider.calls == 1 and len(delivery.files) == 2
    async with sessions() as db:
        funds = await balance(db, 1)
        assert (await db.get(Order, oid)).status == "completed"
        assert (funds.available, funds.reserved) == (1000 - price, 0)


async def test_note_limits_before_reservation(tmp_path, monkeypatch):
    _, provider, _, _, price = await prepare(tmp_path, monkeypatch, [])
    with pytest.raises(ServiceError, match="minutes_input_invalid"):
        await submit(
            1, "meeting_minutes", inputs() | {"notes": "line\n" * 81}, price, "minutes-invalid"
        )
    async with sessions() as db:
        assert await db.scalar(select(func.count(Order.id))) == 0
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 0


@pytest.mark.parametrize("accepted", [True, False])
@pytest.mark.parametrize("defect", ["owner", "deadline"])
async def test_action_metadata_quality_review_repairs_or_refuses_before_export(
    tmp_path, monkeypatch, accepted, defect
):
    from app.providers.ai.gateway import AI
    from app.providers.storage import OwnedStorage
    from app.services.meeting_minutes.schema import Inputs, classification_schema

    source = inputs("en") | {
        "notes": (
            "Sam reported that Sara must submit report 00017 on 2026-10-12, "
            "not the earlier proposed 2026-10-10."
        )
    }
    good = {
        "assignments": [
            {"note_id": 1, "category": "action", "assignee": "Sara", "deadline": "2026-10-12"}
        ]
    }
    bad = {
        "assignments": [
            {
                **good["assignments"][0],
                "assignee": "Sam" if defect == "owner" else "Sara",
                "deadline": "2026-10-10" if defect == "deadline" else "2026-10-12",
            }
        ]
    }
    # Lexical validation alone cannot prove which quoted owner/date is correct.
    classification_schema(Inputs.parse(source).lines()).model_validate(bad)
    reject = '{"accepted":false,"issues":["meaning"]}'
    responses = [
        json.dumps(bad),
        reject,
        json.dumps(good if accepted else bad),
        '{"accepted":true,"issues":[]}' if accepted else reject,
    ]
    store, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, responses)
    ctx["runtime_factory"] = lambda jid, uid: SimpleNamespace(
        ai=AI(provider, jid, uid, quality_review=True),
        storage=OwnedStorage(store, uid),
        renderer=renderer(),
    )
    oid = await submit(1, "meeting_minutes", source, price, "action-review")
    jid = await review(ctx, oid)
    assert provider.calls == 4 and not delivery.files
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == (
            "waiting_confirmation" if accepted else "failed"
        )
        assert (await db.get(Job, jid)).cost_sar == Decimal(".016")
    if accepted:
        await confirm_structure(oid, 1, True)
        await execute_job(ctx, str(jid))
        await execute_job(ctx, str(jid))
        document = Document(BytesIO(dict(delivery.files)["minutes.docx"]))
        cells = [cell.text.replace("\u200e", "") for cell in document.tables[0].rows[1].cells]
        assert cells == ["[1]", "Sara", "2026-10-12"]
        assert provider.calls == 4 and len(delivery.files) == 2
    async with sessions() as db:
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if accepted else 1000, 0)
    assert not list(tmp_path.glob("[0-9]*/*/*"))
