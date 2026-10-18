from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.builders.word import build
from app.services.base import ServiceError
from app.services.meeting_minutes.schema import Inputs, classification_schema, document
from app.services.meeting_minutes.service import MeetingMinutes
from tests.minutes_fixtures import inputs, plan


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("language,default", [("ar", "محضر اجتماع"), ("en", "Meeting minutes")])
@pytest.mark.parametrize("mode", ["omitted", "null", "custom"])
async def test_title_is_shared_by_complete_reviews_and_native_document(language, default, mode):
    source = inputs(language)
    original = source.pop("title")
    if mode != "omitted":
        source["title"] = original if mode == "custom" else None
    expected = original if mode == "custom" else default
    MeetingMinutes.input_schema.validate_inputs(source)
    values = Inputs.parse(source)
    assert values.document_title() == expected
    assert values.title == source.get("title")
    classification = classification_schema(values.lines()).model_validate(plan())
    word_plan = document(values, classification)
    assert word_plan.title == expected
    doc = Document(BytesIO(build(word_plan, professional_template=True, format_dates=True)))
    assert doc.paragraphs[0].text.replace("\u200e", "") == expected
    for index, note in enumerate(values.lines(), 1):
        assert sum(p.text.replace("\u200e", "") == f"[{index}] {note}" for p in doc.paragraphs) == 1

    class AI:
        calls = 0

        async def extract(self, schema, body, prompt):
            self.calls += 1
            assert expected not in body
            return schema.model_validate(plan())

    ai = AI()
    service = MeetingMinutes()
    service.runtime = SimpleNamespace(ai=ai)
    result = await service.run(source)
    assert result.needs_confirmation and not result.artifacts and ai.calls == 1
    assert result.continuation == classification.model_dump(mode="json")
    for ui_language in ("ar", "en"):
        assert result.review_details_localizations[ui_language].startswith(expected + "\n")


@pytest.mark.parametrize(
    "title",
    [
        "",
        "  ",
        "x" * 201,
        "two\nlines",
        "two\rlines",
        "a\u2028b",
        "a\u2029b",
        "\x00",
        True,
        12,
        [],
        {},
    ],
)
def test_invalid_supplied_title_never_becomes_a_default(title):
    with pytest.raises(ServiceError, match="minutes_input_invalid"):
        MeetingMinutes.needs_ai(inputs() | {"title": title})


def test_optional_field_uses_existing_skip_and_keeps_other_fields_required():
    title, language, notes = MeetingMinutes.input_schema.conversation()
    assert title.name == "title" and not title.required and title.single_line
    assert title.max_length == 200 and title.skip_key == "skip_field"
    assert language.required and notes.required
    custom = "  تقرير 00017 =SUM(A1)  "
    # Existing title whitespace normalization is unchanged, without inference/date synthesis.
    assert Inputs.parse(inputs() | {"title": custom}).document_title() == custom.strip()
