from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app.providers.storage import LocalStorage, OwnedStorage
from app.services.image_to_text.schema import Extracted
from app.services.image_to_text.service import ImageToText
from tests.ocr_fixtures import printed


@pytest.fixture(autouse=True)
def database():
    yield


async def output(tmp_path, source, native=False):
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    image = store.save("source.png", printed(), "image/png")
    extract = AsyncMock(return_value=Extracted(readable=True, confidence=0.99, text=source))
    service = ImageToText()
    service.runtime = SimpleNamespace(
        storage=store,
        ai=None if native else SimpleNamespace(extract=extract),
        documents=SimpleNamespace(recognize=extract) if native else None,
    )
    result = await service.run({"images": [image.key]})
    assert result.text == source.strip()
    assert extract.await_count == 1
    return {artifact.filename: store.read(artifact.key) for artifact in result.artifacts}


@pytest.mark.parametrize("native", [False, True])
@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("separator", [" ", "   ", "\n\n  ", " \t\n  ", "\u00a0 "])
async def test_long_word_chunks_retain_every_internal_source_character(
    tmp_path, language, separator, native
):
    word = "السجل" if language == "ar" else "Record"
    source = (word + " 00123\t125.50\n") * 800
    # Place whitespace exactly at a chunk boundary, including leading whitespace
    # in the next chunk. Joining editable body paragraphs must reproduce the TXT.
    source = source[:11999] + separator + source[11999:]
    source = source.strip()
    files = await output(tmp_path, source, native)
    assert files["result.txt"].decode("utf-8") == source
    document = Document(BytesIO(files["result.docx"]))
    paragraphs = [p.text for p in document.paragraphs[1:]]
    assert len(paragraphs) > 1 and all(0 < len(p) <= 12000 for p in paragraphs)
    assert "".join(paragraphs) == source
    # Native editing remains available after save/reopen, not a screenshot.
    document.paragraphs[1].text = "Updated 00123"
    saved = BytesIO()
    document.save(saved)
    assert Document(BytesIO(saved.getvalue())).paragraphs[1].text == "Updated 00123"


@pytest.mark.parametrize("size", [3500, 3501, 11999, 12000, 12001, 24001, 50000])
async def test_thresholds_and_unbroken_text_keep_existing_artifact_and_chunk_limits(tmp_path, size):
    source = "x" * size
    files = await output(tmp_path, source)
    if size <= 3500:
        assert not files
        return
    assert set(files) == {"result.txt", "result.docx"}
    assert files["result.txt"].decode() == source
    paragraphs = Document(BytesIO(files["result.docx"])).paragraphs[1:]
    assert "".join(p.text for p in paragraphs) == source
    assert len(paragraphs) == (size + 11999) // 12000
    assert all(0 < len(p.text) <= 12000 for p in paragraphs)


@pytest.mark.parametrize("separator", [" " * 24000, " \t\n" * 4000])
async def test_whitespace_spanning_whole_chunks_is_not_discarded(tmp_path, separator):
    source = "First 00123" + separator + "Last 125.50"
    files = await output(tmp_path, source)
    paragraphs = Document(BytesIO(files["result.docx"])).paragraphs[1:]
    assert "".join(p.text for p in paragraphs) == files["result.txt"].decode() == source
    assert all(0 < len(p.text) <= 12000 for p in paragraphs)
