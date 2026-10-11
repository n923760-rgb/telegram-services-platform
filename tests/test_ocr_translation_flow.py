import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from docx import Document

from app.core.db import sessions
from app.core.models import Job, Order, Service
from app.orders.engine import submit
from app.services.base import ServiceError
from app.services.registry import registry
from app.wallet.ledger import balance
from app.workers.runner import execute_job
from tests.test_foundation import job_for
from tests.test_ocr_translation_quality import GOOD, SOURCE
from tests.test_real_services import setup


def extracted(text=SOURCE):
    return json.dumps({"readable": True, "confidence": 0.99, "text": text})


async def create_order(store, price):
    image = store.save(1, "input.jpg", b"synthetic test image", "image/jpeg")
    inputs = {"images": [image.key]}
    return await submit(
        1,
        "image_to_text",
        inputs,
        price,
        "numeric-translation",
    )


async def test_ocr_preserves_original_numbers_and_never_translates(tmp_path, monkeypatch):
    responses = [extracted()]
    responses.append(json.dumps({"text": GOOD}))
    store, provider, delivery, ctx, price = await setup(
        "image_to_text", responses, tmp_path, monkeypatch
    )
    oid = await create_order(store, price)
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004")
    assert provider.calls == 1 and len(delivery.results) == 1
    assert delivery.results[0].text == SOURCE
    assert len(provider.responses) == 1


@pytest.mark.parametrize("legacy", [{"language": "none"}, {"language": "en", "style": "formal"}])
async def test_embedded_translation_inputs_rejected_before_reservation(
    tmp_path, monkeypatch, legacy
):
    store, provider, delivery, _, price = await setup("image_to_text", [], tmp_path, monkeypatch)
    image = store.save(1, "input.jpg", b"synthetic", "image/jpeg")
    with pytest.raises(ServiceError):
        await submit(1, "image_to_text", {"images": [image.key], **legacy}, price, "legacy")
    async with sessions() as db:
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 0
    assert not delivery.results and not delivery.files


async def test_ocr_long_original_cached_delivery_retains_editable_numeric_output(
    tmp_path, monkeypatch
):
    records = 60
    source = (SOURCE + "\n") * records
    assert len(source.strip()) > 3500
    store, provider, delivery, ctx, price = await setup(
        "image_to_text",
        [extracted(source)],
        tmp_path,
        monkeypatch,
    )
    oid = await create_order(store, price)
    jid = await job_for(oid)
    send = delivery.send

    async def unavailable(*args):
        raise ConnectionError("synthetic delivery failure")

    monkeypatch.setattr(delivery, "send", unavailable)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "delivering"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, price)
    monkeypatch.setattr(delivery, "send", send)
    async with sessions.begin() as db:
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC) - timedelta(seconds=1)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == 1 and len(delivery.results) == 1
    files = dict(delivery.files)
    assert files["result.txt"].decode() == source.strip()
    document = Document(BytesIO(files["result.docx"]))
    body = "\n".join(p.text for p in document.paragraphs[1:])
    assert body.count("00123") == records and body.count("125.50") == records
    # Arabic dates/numbers are separate direction-aware runs; edit the whole paragraph.
    document.paragraphs[1].text = "Updated editable text 00123"
    saved = BytesIO()
    document.save(saved)
    assert Document(BytesIO(saved.getvalue())).paragraphs[1].text == "Updated editable text 00123"


async def test_ocr_without_translation_has_one_call_and_keeps_original_numbers(
    tmp_path, monkeypatch
):
    store, provider, delivery, ctx, price = await setup(
        "image_to_text", [extracted()], tmp_path, monkeypatch
    )
    oid = await create_order(store, price)
    await execute_job(ctx, str(await job_for(oid)))
    assert provider.calls == 1 and delivery.results[0].text == SOURCE


async def test_ocr_version_upgrade_retains_admin_settings_and_releases_old_work(
    tmp_path, monkeypatch
):
    store, provider, delivery, ctx, price = await setup("image_to_text", [], tmp_path, monkeypatch)
    oid = await create_order(store, price)
    async with sessions.begin() as db:
        row = await db.get(Service, "image_to_text")
        row.version, row.price_halala, row.enabled = "2", 777, True
        (await db.get(Order, oid)).service_version = "2"
    async with sessions.begin() as db:
        await registry.sync(db)
    async with sessions() as db:
        row = await db.get(Service, "image_to_text")
        assert row.version == "3" and row.price_halala == 777 and row.enabled
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 0 and not delivery.results
