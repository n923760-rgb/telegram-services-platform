"""Explicit, source-owned PDF headers without losing extracted rows or calling AI."""

from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import arabic_reshaper
import pytest
from bidi.algorithm import get_display
from openpyxl import load_workbook

from app.builders.pdf_tables import build
from app.core.i18n import CATALOGS
from app.providers.storage import LocalStorage, OwnedStorage
from app.providers.tables.base import Extraction
from app.services.base import ServiceError
from app.services.pdf_tables_excel.schema import direct_labels, source_labels
from app.services.pdf_tables_excel.service import PdfTablesExcel
from app.services.pdf_tables_excel.test_service import extraction, inputs


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("visual", [False, True])
def test_source_headers_and_stable_refs_keep_every_source_cell(language, visual):
    rows = [["رقم الطلب", "المنتج", "2026"], ["00123", "بصل", "0.00"]]
    if visual:
        rows = [[get_display(arabic_reshaper.reshape(cell)) for cell in row] for row in rows]
    original = deepcopy(rows)
    source = Extraction.model_validate(
        {
            "pages": 2,
            "tables": [
                {"page": page, "rows": rows, "accuracy": 100, "whitespace": 0} for page in (1, 2)
            ],
        }
    )
    labels = source_labels(source, language, visual_arabic=visual)
    book = load_workbook(
        BytesIO(
            build(
                source,
                labels,
                language=language,
                enhanced=False,
                source_headers=True,
                visual_arabic=visual,
            )
        )
    )
    for number in (1, 2):
        data, raw = book[f"Data{number}"], book[f"Source{number}"]
        assert [data.cell(3, c).value for c in range(1, 4)] == ["رقم الطلب", "المنتج", "2026"]
        assert [data.cell(4, c).value for c in range(1, 4)] == ["00123", "بصل", "0.00"]
        assert data["D4"].value == f"T{number}:R2"
        assert data.max_row == 4 and raw.max_row == 5
        assert data.tables[f"Data{number}Records"].ref == "A3:D4"
        assert data.tables[f"Data{number}Records"].autoFilter.ref == "A3:D4"
        assert [[raw.cell(r, c).value for c in range(1, 4)] for r in (4, 5)] == original
        assert raw["D4"].value == f"T{number}:R1"
        assert book["Review"].cell(number + 3, 3).value == "2"
    assert source.tables[0].rows == original
    assert not any(
        cell.data_type == "f" or cell.hyperlink for s in book for row in s for cell in row
    )


@pytest.mark.parametrize(
    "headers",
    [
        ["", "Amount"],
        ["ID", "id"],
        [" ID", "Amount"],
        ["ID", "Amount "],
        ["ID", "Two\nlines"],
        ["ID", "x" * 101],
    ],
)
def test_invalid_headers_never_silently_rename_or_drop_source(headers):
    source = extraction()
    source.tables[0].rows = [headers, ["00123", "0"]]
    with pytest.raises(ValueError):
        source_labels(source, "en")
    assert direct_labels(source, "en").tables[0].columns == ["Column 1", "Column 2"]


def test_numeric_and_formula_shaped_source_headers_stay_literal_and_editable():
    source = extraction()
    source.tables[0].rows[0] = ["=1+1", "00123", "Source row reference"]
    labels = source_labels(source, "en")
    book = load_workbook(
        BytesIO(
            build(
                source,
                labels,
                language="en",
                enhanced=False,
                source_headers=True,
            )
        )
    )
    sheet = book["Data1"]
    assert [cell.value for cell in sheet[3]] == [
        "=1+1",
        "00123",
        "Source row reference",
        "Source row reference_",
    ]
    assert sheet["A3"].data_type == "s" and sheet["B3"].data_type == "s"
    sheet["A4"] = "Edited"
    saved = BytesIO()
    book.save(saved)
    assert load_workbook(BytesIO(saved.getvalue()))["Data1"]["A4"].value == "Edited"
    with pytest.raises(ValueError):
        build(
            source, direct_labels(source, "en"), language="en", enhanced=False, source_headers=True
        )


def test_source_headers_reject_oversized_complete_review_and_visual_collisions():
    columns = [f"{i:02d}" + "x" * 98 for i in range(12)]
    source = Extraction.model_validate(
        {
            "pages": 5,
            "tables": [
                {"page": page, "rows": [columns, ["value"] * 12], "accuracy": 100, "whitespace": 0}
                for page in range(1, 6)
            ],
        }
    )
    with pytest.raises(ValueError):
        source_labels(source, "en")
    visual = get_display(arabic_reshaper.reshape("بصل"))
    source = Extraction.model_validate(
        {
            "pages": 1,
            "tables": [
                {"page": 1, "rows": [[visual, "بصل"], ["1", "2"]], "accuracy": 100, "whitespace": 0}
            ],
        }
    )
    with pytest.raises(ValueError):
        source_labels(source, "ar", visual_arabic=True)


async def test_explicit_mode_reviews_headers_and_reuses_source_without_ai(tmp_path):
    storage = OwnedStorage(LocalStorage(tmp_path), 1)
    file = storage.save("source.pdf", b"synthetic", "application/pdf")
    extractor = AsyncMock(return_value=extraction())
    ai = AsyncMock(side_effect=AssertionError("source headings never use AI"))
    service = PdfTablesExcel()
    service.runtime = SimpleNamespace(
        storage=storage, tables=SimpleNamespace(extract=extractor), ai=SimpleNamespace(extract=ai)
    )
    request = inputs(pdf=file.key, mode="source")
    assert not service.needs_ai(request)
    result = await service.run(request)
    assert result.needs_confirmation and not result.artifacts
    for lang in ("ar", "en"):
        review = result.preview_localizations[lang]
        assert all(header in review for header in extraction().tables[0].rows[0])
        assert "R1" in review and "Source" in review
        assert "pdf_tables_source_headers" in CATALOGS[lang]
    service.runtime.tables = service.runtime.ai = None
    final = await service.run({**request, "__continuation": result.continuation})
    assert final.artifacts[0].filename == "tables.xlsx"
    assert (
        load_workbook(BytesIO(storage.read(final.artifacts[0].key)))["Data1"]["A4"].value == "00123"
    )
    assert extractor.await_count == 1 and ai.await_count == 0
    broken = deepcopy(result.continuation)
    broken["labels"]["tables"][0]["columns"][0] = "Invented"
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run({**request, "__continuation": broken})


async def test_invalid_source_headers_create_no_result_and_do_not_call_ai(tmp_path):
    storage = OwnedStorage(LocalStorage(tmp_path), 1)
    file = storage.save("source.pdf", b"synthetic", "application/pdf")
    source = extraction()
    source.tables[0].rows[0][0] = ""
    service = PdfTablesExcel()
    service.runtime = SimpleNamespace(
        storage=storage, tables=SimpleNamespace(extract=AsyncMock(return_value=source)), ai=None
    )
    with pytest.raises(ServiceError, match="input_invalid"):
        await service.run(inputs(pdf=file.key, mode="source"))
    assert len(list(tmp_path.glob("[0-9]*/*/*"))) == 1
