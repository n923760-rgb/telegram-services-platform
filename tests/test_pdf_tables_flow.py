import json
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.db import sessions
from app.core.models import Job, Order
from app.orders.engine import confirm_structure, submit
from app.providers.tables.camelot import CamelotTables
from app.services.base import ServiceError
from app.services.pdf_tables_excel.test_service import LABELS, inputs
from app.wallet.ledger import balance
from app.workers.runner import deliver_confirmations, deliver_failures, execute_job
from tests.pdf_tables_fixtures import table_pdf
from tests.test_foundation import job_for
from tests.test_real_services import setup


async def prepare(tmp_path, monkeypatch, responses):
    store, provider, delivery, ctx, price = await setup(
        "pdf_tables_excel", responses, tmp_path, monkeypatch
    )
    factory = ctx["runtime_factory"]

    def runtime(jid, uid):
        value = factory(jid, uid)
        value.tables = CamelotTables()
        return value

    ctx["runtime_factory"] = runtime
    return store, provider, delivery, ctx, price


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("mode", ["direct", "labels"])
async def test_native_confirmation_capture_once_and_preserved_cells(
    tmp_path, monkeypatch, language, mode
):
    store, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, [json.dumps(LABELS)]
    )
    data, rows = table_pdf(language)
    file = store.save(1, "source.pdf", data, "application/pdf")
    oid = await submit(
        1, "pdf_tables_excel", inputs(pdf=file.key, mode=mode, language=language), price, "tables"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_confirmations(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
        assert (await balance(db, 1)).reserved == price
    assert not delivery.files
    with pytest.raises(ServiceError, match="not_allowed"):
        await confirm_structure(oid, 2, True)
    await confirm_structure(oid, 1, True)

    async def forbidden(*args):
        raise AssertionError("confirmation must reuse extracted source")

    monkeypatch.setattr(CamelotTables, "extract", forbidden)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == (1 if mode == "labels" else 0)
    assert len(delivery.files) == 1
    book = load_workbook(BytesIO(delivery.files[0][1]))
    for name in ("Data1", "Source1"):
        assert [[book[name].cell(r, c).value for c in range(1, 4)] for r in range(4, 7)] == rows
        assert book[name]["C5"].data_type == "s"
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_rejected_review_releases_without_export(tmp_path, monkeypatch):
    store, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [])
    file = store.save(1, "source.pdf", table_pdf()[0], "application/pdf")
    oid = await submit(1, "pdf_tables_excel", inputs(pdf=file.key), price, "reject")
    await execute_job(ctx, str(await job_for(oid)))
    await confirm_structure(oid, 1, False)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert not delivery.files and provider.calls == 0


@pytest.mark.parametrize("repair", [True, False])
async def test_ai_one_repair_or_release(tmp_path, monkeypatch, repair):
    store, provider, delivery, ctx, price = await prepare(
        tmp_path, monkeypatch, ["{}", json.dumps(LABELS) if repair else "{}"]
    )
    file = store.save(1, "source.pdf", table_pdf()[0], "application/pdf")
    oid = await submit(1, "pdf_tables_excel", inputs(pdf=file.key, mode="labels"), price, "repair")
    await execute_job(ctx, str(await job_for(oid)))
    assert provider.calls == 2 and not delivery.files
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == ("waiting_confirmation" if repair else "failed")
        funds = await balance(db, 1)
        assert funds.reserved == (price if repair else 0)
        if not repair:
            assert order.error_key == "provider_invalid" and funds.available == 1000


async def test_invalid_pdf_releases_and_notifies_once(tmp_path, monkeypatch):
    store, _, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [])
    file = store.save(1, "source.pdf", b"not pdf", "application/pdf")
    oid = await submit(1, "pdf_tables_excel", inputs(pdf=file.key), price, "bad")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert not delivery.files and len(delivery.errors) == 1
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_delivery_retry_uses_saved_workbook(tmp_path, monkeypatch):
    from app.services.pdf_tables_excel import service

    store, _, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [])
    file = store.save(1, "source.pdf", table_pdf()[0], "application/pdf")
    oid = await submit(1, "pdf_tables_excel", inputs(pdf=file.key), price, "retry")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await confirm_structure(oid, 1, True)
    original = delivery.send

    async def offline(*args):
        raise ConnectionError("synthetic delivery failure")

    monkeypatch.setattr(delivery, "send", offline)
    await execute_job(ctx, str(jid))
    async with sessions.begin() as db:
        assert (await db.get(Order, oid)).result
        assert (await balance(db, 1)).reserved == price
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC)

    def forbidden(*args, **kwargs):
        raise AssertionError("retry must reuse saved XLSX")

    monkeypatch.setattr(service, "build", forbidden)
    monkeypatch.setattr(delivery, "send", original)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert len(delivery.files) == 1


async def test_native_visual_arabic_confirmed_export_preserves_source_zero_ai(
    tmp_path, monkeypatch
):
    from app.builders.arabic_table_text import recover_visual_arabic

    store, provider, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [])
    data, rows = table_pdf("ar", visual_arabic=True)
    file = store.save(1, "source.pdf", data, "application/pdf")
    oid = await submit(
        1,
        "pdf_tables_excel",
        inputs(pdf=file.key, language="ar", arabic_order="visual"),
        price,
        "visual-arabic",
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_confirmations(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
    await confirm_structure(oid, 1, True)

    async def forbidden(*args):
        raise AssertionError("must reuse extraction")

    monkeypatch.setattr(CamelotTables, "extract", forbidden)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    book = load_workbook(BytesIO(delivery.files[0][1]))
    for r, row in enumerate(rows, 4):
        for c, value in enumerate(row, 1):
            assert book["Source1"].cell(r, c).value == value
            assert book["Data1"].cell(r, c).value == recover_visual_arabic(value)
    assert book["Data1"]["C5"].value == "مرحبا" and book["Data1"]["A5"].value == "00123"
    assert provider.calls == 0 and len(delivery.files) == 1
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).available == 1000 - price
        assert (await balance(db, 1)).reserved == 0


@pytest.mark.parametrize("cached", [False, True])
async def test_arabic_order_upgrade_keeps_prices_gates_and_legacy_cache(
    tmp_path, monkeypatch, cached
):
    from app.core.models import Service
    from app.services.registry import registry

    store, _, delivery, ctx, price = await prepare(tmp_path, monkeypatch, [])
    file = store.save(1, "source.pdf", table_pdf()[0], "application/pdf")
    oid = await submit(1, "pdf_tables_excel", inputs(pdf=file.key), price, "upgrade-arabic")
    jid = await job_for(oid)
    original = delivery.send
    if cached:
        await execute_job(ctx, str(jid))
        await confirm_structure(oid, 1, True)

        async def offline(*args):
            raise ConnectionError("synthetic offline")

        monkeypatch.setattr(delivery, "send", offline)
        await execute_job(ctx, str(jid))
    async with sessions.begin() as db:
        service = await db.get(Service, "pdf_tables_excel")
        old = {
            **service.input_schema,
            "fields": [f for f in service.input_schema["fields"] if f["name"] != "arabic_order"],
        }
        service.version, service.input_schema = "2", old
        service.price_halala, service.price_stars = 777, 25
        order = await db.get(Order, oid)
        order.service_version, order.input_schema_snapshot = "2", old
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC)
    async with sessions.begin() as db:
        await registry.sync(db)
        service = await db.get(Service, "pdf_tables_excel")
        assert service.version == "3" and service.enabled
        assert service.price_halala == 777 and service.price_stars == 25

    async def forbidden(*args):
        raise AssertionError("upgrade must not reextract")

    monkeypatch.setattr(CamelotTables, "extract", forbidden)
    monkeypatch.setattr(delivery, "send", original)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == ("completed" if cached else "cancelled")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if cached else 1000, 0)
    assert len(delivery.files) == int(cached)
