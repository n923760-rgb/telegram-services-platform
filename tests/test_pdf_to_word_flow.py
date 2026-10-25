import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import CostHold, CostUsage, Job, Ledger, Order, Service
from app.orders.engine import submit
from app.services.pdf_to_word.test_service import SOURCE, pdf_bytes
from app.services.registry import registry
from app.wallet.ledger import balance
from app.workers.runner import deliver_failures, execute_job
from tests.pdf_form_fixtures import form_pdf
from tests.pdf_image_fixtures import mixed_pdf
from tests.pdf_tables_fixtures import table_pdf
from tests.test_foundation import job_for
from tests.test_real_services import setup

PLAN = '{"document":{"title":"Invoice","sections":[{"paragraphs":["' + SOURCE + '"]}]}}'


@pytest.mark.parametrize("mode", ["valid", "repaired", "failed"])
async def test_editable_table_delivery_or_invalid_cell_release_once(tmp_path, monkeypatch, mode):
    data, rows = table_pdf()
    plan = json.dumps(
        {
            "document": {
                "title": "Document",
                "sections": [{"tables": [{"columns": rows[0], "rows": rows[1:]}]}],
            }
        }
    )
    bad = plan.replace("00123", "123")
    responses = [plan] if mode == "valid" else [bad, plan if mode == "repaired" else bad]
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", responses, tmp_path, monkeypatch
    )
    source = store.save(1, "source.pdf", data, "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-editable-table")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    success = mode != "failed"
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == ("completed" if success else "failed")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if success else 1000, 0)
        kinds = list(await db.scalars(select(Ledger.kind).where(Ledger.order_id == oid)))
        assert sorted(kinds) == sorted(["reserve", "capture" if success else "release"])
        assert list(await db.scalars(select(CostHold.status))) == ["settled"] * len(responses)
    assert provider.calls == len(responses)
    assert len(delivery.files) == int(success) and len(delivery.errors) == int(not success)
    if success:
        document = Document(BytesIO(delivery.files[0][1]))
        assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == rows


@pytest.mark.parametrize("recover", [True, False])
async def test_missing_context_refusal_repairs_or_releases_once_with_settled_costs(
    tmp_path, monkeypatch, recover
):
    bad = '{"missing_information":true}'
    responses = [bad, PLAN, '{"accepted":true,"issues":[]}'] if recover else [bad, bad]
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", responses, tmp_path, monkeypatch
    )
    factory = ctx["runtime_factory"]

    def reviewed(jid, uid):
        runtime = factory(jid, uid)
        runtime.ai.quality_review = True
        return runtime

    ctx["runtime_factory"] = reviewed
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "missing-context-v6")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    calls = 3 if recover else 2
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.service_version == "8"
        assert order.status == ("completed" if recover else "failed")
        assert order.error_key == (None if recover else "provider_invalid")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if recover else 1000, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004") * calls
        assert await db.scalar(select(func.count(CostUsage.id))) == calls
        assert list(await db.scalars(select(CostHold.status))) == ["settled"] * calls
        kinds = list(await db.scalars(select(Ledger.kind).where(Ledger.order_id == oid)))
        assert sorted(kinds) == sorted(["reserve", "capture" if recover else "release"])
    assert provider.calls == calls and len(delivery.files) == int(recover)
    assert len(delivery.errors) == int(not recover)
    if recover:
        assert SOURCE in [p.text for p in Document(BytesIO(delivery.files[0][1])).paragraphs]
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("reviewed", [False, True])
@pytest.mark.parametrize("repaired", [False, True])
async def test_source_fidelity_repair_and_failure_settlement(
    tmp_path, monkeypatch, reviewed, repaired
):
    bad = PLAN.replace("12345", "54321")
    responses = [bad, PLAN if repaired else bad]
    if reviewed and repaired:
        responses.append('{"accepted":true,"issues":[]}')
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", responses, tmp_path, monkeypatch
    )
    original_factory = ctx["runtime_factory"]

    def runtime(jid, uid):
        result = original_factory(jid, uid)
        result.ai.quality_review = reviewed
        return result

    ctx["runtime_factory"] = runtime
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-fidelity-repair")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    expected_calls = 3 if reviewed and repaired else 2
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == ("completed" if repaired else "failed")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if repaired else 1000, 0)
        entries = list(await db.scalars(select(Ledger).where(Ledger.order_id == oid)))
        assert sorted(entry.kind for entry in entries) == sorted(
            ["reserve", "capture" if repaired else "release"]
        )
        usage = list(await db.scalars(select(CostUsage).where(CostUsage.job_id == jid)))
        assert len(usage) == expected_calls
    assert provider.calls == expected_calls
    if repaired:
        assert len(delivery.files) == 1 and not delivery.errors
        assert SOURCE in [p.text for p in Document(BytesIO(delivery.files[0][1])).paragraphs]
    else:
        assert not delivery.files
        assert delivery.errors == [(1, oid, "provider_invalid")]
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("kind", ["xobject", "inline", "nested", "later_page"])
async def test_mixed_pdf_rejection_releases_credit_and_notifies_once(tmp_path, monkeypatch, kind):
    store, provider, delivery, ctx, price = await setup("pdf_to_word", [], tmp_path, monkeypatch)
    source = store.save(1, "source.pdf", mixed_pdf(kind), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "mixed-page-failure")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 0 and not delivery.files
    assert delivery.errors == [(1, oid, "pdf_image_content")]
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize(
    "kind", ["acroform", "widget", "widget_indirect", "permissions", "empty_form"]
)
async def test_pdf_form_rejection_has_no_provider_cost_releases_and_cleans_files(
    tmp_path, monkeypatch, kind
):
    store, provider, delivery, ctx, price = await setup("pdf_to_word", [], tmp_path, monkeypatch)
    source = store.save(1, "source.pdf", form_pdf(kind, later=True), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "form-rejection")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
        assert (await db.get(Job, jid)).cost_sar == 0
        assert not list(await db.scalars(select(CostUsage).where(CostUsage.job_id == jid)))
        entries = list(await db.scalars(select(Ledger).where(Ledger.order_id == oid)))
        assert sorted(entry.kind for entry in entries) == ["release", "reserve"]
    assert provider.calls == 0 and not delivery.files
    assert delivery.errors == [(1, oid, "pdf_word_form")]
    assert not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("old_version", ["2", "3", "4", "5", "6", "7"])
async def test_pdf_word_upgrade_keeps_admin_settings_and_releases_old_pending_work(
    tmp_path, monkeypatch, old_version
):
    store, provider, delivery, ctx, price = await setup("pdf_to_word", [], tmp_path, monkeypatch)
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "old-pdf-work")
    async with sessions.begin() as db:
        row = await db.get(Service, "pdf_to_word")
        row.version, row.price_halala, row.enabled = old_version, 777, True
        row.price_stars = 143
        (await db.get(Order, oid)).service_version = old_version
    async with sessions.begin() as db:
        await registry.sync(db)
    async with sessions() as db:
        row = await db.get(Service, "pdf_to_word")
        assert row.version == "8" and row.price_halala == 777 and row.enabled
        assert row.price_stars == 143
    await execute_job(ctx, str(await job_for(oid)))
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 0 and not delivery.files


async def test_pdf_to_word_delivers_editable_docx_then_captures_once(tmp_path, monkeypatch):
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", [PLAN], tmp_path, monkeypatch
    )
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-word-success")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert funds.available == 1000 - price and funds.reserved == 0
    assert provider.calls == 1 and len(delivery.files) == 1
    filename, data = delivery.files[0]
    assert filename == "result.docx"
    assert SOURCE in [p.text for p in Document(BytesIO(data)).paragraphs]


@pytest.mark.parametrize(
    ("data", "key"),
    [
        (b"not a pdf", "file_invalid"),
        (pdf_bytes(encrypted=True), "file_invalid"),
        (pdf_bytes(text=""), "pdf_text_unreadable"),
        (pdf_bytes(drawn_page=True), "pdf_text_unreadable"),
        (pdf_bytes(text="x" * 50001), "pdf_text_too_long"),
    ],
    ids=["malformed", "encrypted", "blank", "mixed", "oversized-text"],
)
async def test_pdf_rejection_releases_credit_and_notifies_once(tmp_path, monkeypatch, data, key):
    store, provider, delivery, ctx, price = await setup("pdf_to_word", [], tmp_path, monkeypatch)
    source = store.save(1, "source.pdf", data, "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-word-failure")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert funds.available == 1000 and funds.reserved == 0
    assert provider.calls == 0
    assert delivery.errors == [(1, oid, key)] and not delivery.files


async def test_invalid_pdf_plan_releases_credit_without_delivering(tmp_path, monkeypatch):
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", ["{}", "{}"], tmp_path, monkeypatch
    )
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-word-invalid-plan")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert funds.available == 1000 and funds.reserved == 0
    assert provider.calls == 2 and not delivery.files
    assert delivery.errors == [(1, oid, "provider_invalid")]


@pytest.mark.parametrize("old_version", ["3", "4", "5", "6", "7"])
async def test_pdf_delivery_retry_reuses_output_and_captures_after_delivery(
    tmp_path, monkeypatch, old_version
):
    store, provider, delivery, ctx, price = await setup(
        "pdf_to_word", [PLAN], tmp_path, monkeypatch
    )
    source = store.save(1, "source.pdf", pdf_bytes(), "application/pdf")
    oid = await submit(1, "pdf_to_word", {"pdf": source.key}, price, "pdf-word-retry")
    jid = await job_for(oid)
    original_send = delivery.send

    async def unavailable(*args):
        raise ConnectionError("test delivery unavailable")

    monkeypatch.setattr(delivery, "send", unavailable)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "delivering"
        funds = await balance(db, 1)
        assert funds.available == 1000 - price and funds.reserved == price
    monkeypatch.setattr(delivery, "send", original_send)
    async with sessions.begin() as db:
        # Prepared historical results replay; only unprepared work is cancelled.
        (await db.get(Order, oid)).service_version = old_version
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC) - timedelta(seconds=1)

    def no_build(*args, **kwargs):
        raise AssertionError("prepared Word output must not be regenerated")

    monkeypatch.setattr("app.services.pdf_to_word.service.word.build", no_build)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert funds.available == 1000 - price and funds.reserved == 0
    assert provider.calls == 1 and len(delivery.files) == 1
