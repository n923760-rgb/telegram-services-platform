from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app.builders.word import build
from app.builders.word_schema import WordDocument
from app.services.document_fidelity import visible_text
from app.services.pdf_to_word.layout import opening_title
from app.services.pdf_to_word.schema import WordPlan, source_plan_schema
from app.services.pdf_to_word.test_service import configured_service
from tests.pdf_tables_fixtures import table_pdf


@pytest.mark.parametrize("title", ["جدول متابعة الطلبات", "Order records"])
@pytest.mark.parametrize("field", ["heading", "title"])
def test_single_source_title_renders_once_and_preserves_editable_cells(title, field):
    source = f"{title}\nID\nAmount\n00123\n125.50\nNotes\nKeep pending."
    values = {
        "title": title,
        "sections": [
            {
                field: title,
                "tables": [{"columns": ["ID", "Amount"], "rows": [["00123", "125.50"]]}],
            },
            {"heading": "Notes", "paragraphs": ["Keep pending."]},
        ],
    }
    reviewed = source_plan_schema(source).model_validate({"document": values})
    snapshot = reviewed.model_dump()
    cleaned = opening_title(reviewed.document, source)
    assert reviewed.model_dump() == snapshot
    assert visible_text(cleaned).splitlines().count(title) == 1
    # All existing source-fidelity gates still pass after presentation cleanup.
    source_plan_schema(source).model_validate({"document": cleaned.model_dump()})
    output = Document(BytesIO(build(cleaned)))
    assert [p.text for p in output.paragraphs].count(title) == 1
    assert [p.text for p in output.paragraphs if p.text][-2:] == ["Notes", "Keep pending."]
    assert [[c.text for c in r.cells] for r in output.tables[0].rows] == [
        ["ID", "Amount"],
        ["00123", "125.50"],
    ]


def document(title="Report", **section):
    return WordDocument(
        title=title, sections=[{"heading": title, "paragraphs": ["Body"], **section}]
    )


@pytest.mark.parametrize("source", ["Report\nBody\nReport", "Other report\nBody", "Report Body"])
def test_repeated_or_unconfirmed_source_title_is_not_removed(source):
    reviewed = document()
    assert opening_title(reviewed, source) is reviewed


@pytest.mark.parametrize(
    "title", ["Report 00123", "تقرير ٠٠١٢٣", "a@example.com", "https://example.com"]
)
def test_protected_title_values_are_never_removed(title):
    reviewed = document(title)
    assert opening_title(reviewed, f"{title}\nBody") is reviewed


@pytest.mark.parametrize(
    "first",
    [
        {"heading": "Overview", "paragraphs": ["Body"]},
        {"subtitle": "Cover note", "heading": "Report", "paragraphs": ["Body"]},
        {"paragraphs": ["Body"]},
        {"tables": [{"columns": ["Name"], "rows": [["Body"]]}]},
    ],
)
def test_later_matching_heading_is_kept(first):
    reviewed = WordDocument(
        title="Report", sections=[first, {"heading": "Report", "paragraphs": ["More"]}]
    )
    assert opening_title(reviewed, "Report\nBody\nMore") is reviewed


def test_explicit_page_break_heading_only_and_additional_occurrences_are_kept():
    reviewed = document(page_break_before=True)
    assert opening_title(reviewed, "Report\nBody") is reviewed
    reviewed = document()
    reviewed.layout = "resume"
    assert opening_title(reviewed, "Report\nBody") is reviewed
    reviewed = WordDocument(title="Report", sections=[{"heading": "Report"}])
    assert opening_title(reviewed, "Report") is reviewed
    reviewed = document(bullets=["Report"])
    assert opening_title(reviewed, "Report\nBody") is reviewed


def test_empty_sections_and_equivalent_whitespace_keep_the_actual_main_title():
    reviewed = WordDocument(
        title="Order records", sections=[{}, {"heading": "Order  records", "paragraphs": ["Body"]}]
    )
    cleaned = opening_title(reviewed, "Order   records\nBody")
    assert cleaned.title == "Order records" and cleaned.sections[1].heading == ""
    assert reviewed.sections[1].heading == "Order  records"


async def test_service_cleans_opening_heading_after_validation_without_another_ai_call(tmp_path):
    data, rows = table_pdf()
    overlay = BytesIO()
    page = canvas.Canvas(overlay, pagesize=(600, 800))
    page.drawString(60, 775, "Order records")
    page.save()
    writer = PdfWriter()
    source_page = writer.add_page(PdfReader(BytesIO(data)).pages[0])
    source_page.merge_page(PdfReader(BytesIO(overlay.getvalue())).pages[0], over=False)
    source = BytesIO()
    writer.write(source)
    service, artifact, store, ai = configured_service(tmp_path, source.getvalue())
    values = {
        "document": {
            "title": "Order records",
            "sections": [
                {"heading": "Order records", "tables": [{"columns": rows[0], "rows": rows[1:]}]}
            ],
        }
    }
    plan = WordPlan.model_validate(deepcopy(values))
    ai.extract.return_value = plan
    result = await service.run({"pdf": artifact.key})
    output = Document(BytesIO(store.read(result.artifacts[0].key)))
    assert [p.text for p in output.paragraphs] == ["Order records"]
    assert [[c.text for c in r.cells] for r in output.tables[0].rows] == rows
    assert plan.document.sections[0].heading == "Order records"
    assert ai.extract.await_count == 1
