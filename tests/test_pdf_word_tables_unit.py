"""Editable PDF conversion tables retain literal data through the real builder."""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError
from pypdf import PdfReader

from app.builders.word import build
from app.files.pdf_text import extract_page_text
from app.services.pdf_to_word.schema import WordPlan, source_plan_schema
from app.services.pdf_to_word.test_service import configured_service
from tests.pdf_tables_fixtures import table_pdf


def table_plan(rows):
    return {
        "document": {
            "title": "Document",
            "sections": [
                {"tables": [{"columns": rows[0], "rows": rows[1:]}]},
            ],
        }
    }


@pytest.mark.parametrize("language", ["ar", "en"])
async def test_real_pdf_table_is_editable_with_exact_cell_order(tmp_path, language):
    data, rows = table_pdf(language)
    service, artifact, store, ai = configured_service(tmp_path, data)
    ai.extract.return_value = WordPlan.model_validate(table_plan(rows))
    result = await service.run({"pdf": artifact.key})
    document = Document(BytesIO(store.read(result.artifacts[0].key)))
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == rows
    assert not any("00123" in paragraph.text for paragraph in document.paragraphs)
    assert ai.extract.await_count == 1
    # The actual request-local schema accepts exactly the extracted source values.
    schema = ai.extract.call_args.args[0]
    schema.model_validate(table_plan(rows))


def test_table_only_document_round_trips_without_a_placeholder_paragraph():
    rows = [["ID", "Amount", "Contact"], ["00123", "125.50", "a@example.com"]]
    schema = source_plan_schema("\n".join(cell for row in rows for cell in row))
    parsed = schema.model_validate(table_plan(rows))
    assert not parsed.document.sections[0].paragraphs
    document = Document(BytesIO(build(parsed.document)))
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == rows


@pytest.mark.parametrize(
    ("source", "value"),
    [
        ("00123", "123"),
        ("125.50", "125.5"),
        ("٠٠١٢٣", "00123"),
        ("١٢٥٫٥٠", "١٢٥.٥٠"),
        ("10%", "10"),
        ("a@example.com", "other@example.com"),
        ("https://example.com/id", "https://example.com/other"),
    ],
)
@pytest.mark.parametrize("location", ["caption", "column", "cell", "block", "subtitle"])
def test_protected_tokens_in_every_new_visible_field(source, value, location):
    candidate = table_plan([["Value"], ["Text"]])
    section = candidate["document"]["sections"][0]
    if location == "caption":
        section["tables"][0]["caption"] = source
    elif location == "column":
        section["tables"][0]["columns"] = [source]
    elif location == "cell":
        section["tables"][0]["rows"] = [[source]]
    elif location == "block":
        section["blocks"] = [{"text": source}]
    else:
        section["subtitle"] = source
    schema = source_plan_schema(source)
    schema.model_validate(candidate)
    changed = deepcopy(candidate)
    # Change only one protected field, preserving all shape/metadata.
    import json

    changed = json.loads(json.dumps(changed, ensure_ascii=False).replace(source, value))
    with pytest.raises(ValidationError):
        schema.model_validate(changed)


@pytest.mark.parametrize("mutation", ["duplicate", "omit", "numbered_caption"])
def test_rejects_duplicate_missing_or_invented_numeric_table_content(mutation):
    candidate = table_plan([["ID"], ["00123"]])
    section = candidate["document"]["sections"][0]
    if mutation == "duplicate":
        section["paragraphs"] = ["00123"]
    elif mutation == "omit":
        section["tables"][0]["rows"] = [["Text"]]
    else:
        section["tables"][0]["caption"] = "Table 1"
    with pytest.raises(ValidationError):
        source_plan_schema("00123").model_validate(candidate)


def test_table_schema_is_bounded_and_does_not_disclose_private_source():
    source = "private@example.com 00123"
    schema = source_plan_schema(source)
    assert source not in str(schema.model_json_schema())
    for rows in (
        [["ID"], *[["00123"]] * 101],
        [["ID"], ["x" * 241]],
        [["ID"], ["00123", "extra"]],
        [["ID"], ["12\x0034"]],
    ):
        with pytest.raises(ValidationError):
            schema.model_validate(table_plan(rows))


def test_legacy_paragraph_conversion_remains_valid():
    data, rows = table_pdf()
    text = extract_page_text(PdfReader(BytesIO(data)).pages[0], limit=50000)
    candidate = {"document": {"title": "Document", "sections": [{"paragraphs": [text]}]}}
    assert not source_plan_schema(text).model_validate(candidate).document.sections[0].tables
