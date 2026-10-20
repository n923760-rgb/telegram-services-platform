from io import BytesIO

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from app.builders.direction import is_rtl
from app.builders.pptx import _reading_rtl, build
from app.builders.schema import Deck


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize(
    "text,rtl",
    [
        ("تواصل: help@demo.invalid", True),
        ("راجع https://example.invalid/report", True),
        ("زر www.example.invalid", True),
        ("موعد 2026-10-18", True),
        ("عنوان عربي واضح", True),
        ("English source text مرحبا", False),
        ("Contact help@demo.invalid", False),
        ("help@demo.invalid", False),
        ("https://example.invalid/مسار/عربي/طويل", False),
        ("", False),
    ],
)
def test_identifiers_do_not_vote_as_prose_in_direction_detection(text, rtl):
    assert _reading_rtl(text) is rtl


@pytest.mark.parametrize("text", ["تواصل: help@demo.invalid", "زر www.example.invalid"])
def test_reproduced_dominant_contact_regression(text):
    assert not is_rtl(text) and _reading_rtl(text)


@pytest.mark.parametrize("cover", [False, True])
@pytest.mark.parametrize(
    "text,rtl",
    [
        ("تواصل: help@demo.invalid", True),
        ("زر www.example.invalid", True),
        ("Contact help@demo.invalid", False),
    ],
)
def test_native_titles_and_bullets_keep_text_and_use_correct_direction(cover, text, rtl):
    slides = [{"title": text, "bullets": [text]}]
    if cover:
        slides.append({"title": "Next", "bullets": ["Source text"]})
    data = Deck.model_validate({"slides": slides})
    deck = Presentation(BytesIO(build(data)))
    for name in ("Slide title", "Slide content"):
        frame = next(shape.text_frame for shape in deck.slides[0].shapes if shape.name == name)
        paragraph = frame.paragraphs[0]
        assert paragraph.text == text
        assert paragraph._p.pPr.get("rtl") == str(int(rtl))
        assert paragraph.alignment == (
            PP_ALIGN.CENTER if cover else PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
        )
        if name == "Slide content" and not cover:
            assert paragraph._p.pPr.get("marR" if rtl else "marL") is not None
    assert (
        next(shape.text for shape in deck.slides[0].shapes if shape.name == "Slide number")
        == f"1 / {len(slides)}"
    )
    assert data.slides[0].title == text and data.slides[0].bullets == [text]


def test_native_table_direction_is_per_prose_and_per_cell_with_exact_contacts():
    columns = ["مرجع: help@demo.invalid", "بيان"]
    rows = [
        ["للتواصل: help@demo.invalid", "راجع www.example.invalid"],
        ["help@demo.invalid", "English record"],
    ]
    deck = Presentation(
        BytesIO(
            build(
                Deck.model_validate(
                    {"slides": [{"title": "السجل", "table": {"columns": columns, "rows": rows}}]}
                )
            )
        )
    )
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert table._tbl.tblPr.get("rtl") == "1"
    assert [[cell.text for cell in row.cells] for row in table.rows] == [columns, *rows]
    for row, flags in zip(table.rows, [(1, 1), (1, 1), (0, 0)], strict=True):
        assert [cell.text_frame.paragraphs[0]._p.pPr.get("rtl") for cell in row.cells] == list(
            map(str, flags)
        )
