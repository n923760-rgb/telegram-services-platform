"""One contextual numeric reuse is different from invented or unbounded values."""

import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.ops.service_acceptance import pptx_token_diagnostics
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.protected_values import ProtectedValues
from app.services.text_editing.schema import CONTACT_TOKEN
from app.services.text_to_pptx.schema import source_plan_schema
from app.services.text_to_pptx.service import TextToPptx


@pytest.fixture(autouse=True)
def database():
    yield


SOURCE = "Order 00123\nAmount 125.50 SAR\nThe payment has not been approved."


def closing_plan():
    return {
        "deck": {
            "slides": [
                {"title": "Order", "bullets": ["Order 00123"]},
                {"title": "Amount", "bullets": ["Amount 125.50 SAR"]},
                {
                    "title": "Status",
                    "bullets": [
                        "Order 00123: amount 125.50 SAR.",
                        "The payment has not been approved.",
                    ],
                },
            ]
        }
    }


def test_three_slide_closing_context_passes_with_exact_values_and_safe_diagnostics():
    plan = closing_plan()
    schema = source_plan_schema(SOURCE, 3)
    result = schema.model_validate(plan)
    assert schema.model_validate(result.model_dump(mode="json")) == result
    diagnostics = pptx_token_diagnostics(json.dumps(plan), SOURCE)
    assert set(diagnostics["counts"].values()) == {0}
    assert diagnostics["locations"] == []
    assert "00123" not in json.dumps(schema.model_json_schema())


@pytest.mark.parametrize(
    "source,token",
    [
        ("Record ٠٠١٢٣", "٠٠١٢٣"),
        ("Change -12.50%", "-12.50%"),
        ("Date 2026-10-20", "2026-10-20"),
        ("Value 0", "0"),
    ],
)
def test_exact_arabic_digits_sign_precision_date_and_zero_can_recur_on_another_slide(source, token):
    data = {
        "deck": {
            "slides": [
                {"title": "Detail", "bullets": [token]},
                {"title": "Context", "bullets": [token]},
            ]
        }
    }
    assert source_plan_schema(source, 2).model_validate(data)


@pytest.mark.parametrize("change", ["third_use", "same_slide", "new", "precision", "leading_zero"])
def test_excess_repetition_and_new_or_altered_values_stay_rejected(change):
    plan = closing_plan()
    slides = plan["deck"]["slides"]
    if change == "third_use":
        slides[1]["bullets"].append("Order 00123")
    elif change == "same_slide":
        slides[0]["title"] = "Order 00123"
    elif change == "new":
        slides[2]["bullets"].append("99999")
    elif change == "precision":
        slides[2]["bullets"][0] = "Amount 125.5 SAR"
    else:
        slides[2]["bullets"][0] = "Order 123"
    with pytest.raises(ValidationError, match="bounded reuse"):
        source_plan_schema(SOURCE, 3).model_validate(plan)
    assert sum(pptx_token_diagnostics(json.dumps(plan), SOURCE)["counts"].values()) > 0


@pytest.mark.parametrize("token", ["a@example.invalid", "https://demo.invalid"])
def test_contacts_keep_the_original_global_count_limit(token):
    data = {
        "deck": {
            "slides": [
                {"title": "Contact", "bullets": [token]},
                {"title": "Again", "bullets": [token]},
            ]
        }
    }
    with pytest.raises(ValidationError, match="bounded reuse"):
        source_plan_schema(token, 2).model_validate(data)


def test_multiple_original_occurrences_add_only_one_global_reuse_not_one_per_slide():
    data = {
        "deck": {
            "slides": [
                {"title": "First", "bullets": ["00123 00123"]},
                {"title": "Second", "bullets": ["00123"]},
            ]
        }
    }
    schema = source_plan_schema("00123\n00123", 2)
    assert schema.model_validate(data)
    data["deck"]["slides"][1]["bullets"].append("00123")
    with pytest.raises(ValidationError):
        schema.model_validate(data)


@pytest.mark.parametrize("continuation", [False, True])
async def test_real_editable_deck_restores_reused_references_and_keeps_exact_count(
    tmp_path, continuation
):
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    encoded = json.dumps(closing_plan())
    for marker, literal in protected.values.items():
        encoded = encoded.replace(literal, marker)
    calls = []

    async def extract(schema, source, prompt):
        calls.append(source)
        assert "one extra contextual use" in prompt.lower()
        return schema.model_validate_json(encoded)

    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service = TextToPptx()
    service.runtime = SimpleNamespace(storage=store, ai=SimpleNamespace(extract=extract))
    inputs = {"text": SOURCE, "slide_count": "3"}
    if continuation:
        inputs["__continuation"] = closing_plan()
    result = await service.run(inputs)
    native = Presentation(BytesIO(store.read(result.artifacts[0].key)))
    texts = "\n".join(
        shape.text for slide in native.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert len(native.slides) == 3
    assert texts.count("00123") == texts.count("125.50") == 2
    assert "not been approved" in texts and "[[value-" not in texts
    assert len(calls) == int(not continuation)
