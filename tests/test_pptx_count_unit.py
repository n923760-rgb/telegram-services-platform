"""Explicit count is a request-local output constraint, never a new source fact."""

import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.builders import pptx
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.text_to_pptx.prompt import DECK
from app.services.text_to_pptx.schema import DeckPlan, Inputs, source_plan_schema
from app.services.text_to_pptx.service import TextToPptx


@pytest.fixture(autouse=True)
def database():
    yield


def plan(count):
    return {
        "deck": {
            "slides": [{"title": "Section", "bullets": ["Source content"]} for _ in range(count)]
        }
    }


@pytest.mark.parametrize(
    "value,expected", [(1, 1), (30, 30), ("1", 1), ("٣", 3), ("۳۰", 30), (" 5 ", 5), (None, None)]
)
def test_count_accepts_optional_bounded_western_and_arabic_digits(value, expected):
    assert Inputs.parse({"text": "Content", "slide_count": value}).slide_count == expected


@pytest.mark.parametrize(
    "value", [0, 31, True, 3.0, "", "auto", "3.0", "1e1", "1-3", "-1", "3\n4", [], {}]
)
def test_invalid_count_is_rejected_at_admission(value):
    with pytest.raises(ServiceError, match="input_invalid"):
        TextToPptx.needs_ai({"text": "Content", "slide_count": value})


@pytest.mark.parametrize("count", [1, 3, 30])
def test_schema_exposes_exact_count_without_source_values_and_native_output_matches(count):
    schema = source_plan_schema("Private source contact@example.invalid 00123", count)
    shape = schema.model_json_schema()
    slides = shape["$defs"]["RequestedDeck"]["properties"]["slides"]
    assert slides["minItems"] == slides["maxItems"] == count
    assert "contact@example.invalid" not in json.dumps(shape) and "00123" not in json.dumps(shape)
    validated = schema.model_validate(plan(count))
    native = Presentation(BytesIO(pptx.build(validated.deck)))
    assert len(native.slides) == count
    assert all(
        any(shape.has_text_frame and shape.text == "Source content" for shape in slide.shapes)
        for slide in native.slides
    )
    assert any(
        shape.has_text_frame and shape.text == f"{count} / {count}"
        for shape in native.slides[-1].shapes
    )
    with pytest.raises(ValidationError):
        schema.model_validate(plan(count + 1))
    if count > 1:
        with pytest.raises(ValidationError):
            schema.model_validate(plan(count - 1))


def test_auto_mode_retains_existing_bounds_and_request_schemas_do_not_share_counts():
    auto, two, three = (
        source_plan_schema("Content"),
        source_plan_schema("Content", 2),
        source_plan_schema("Content", 3),
    )
    assert auto.model_validate(plan(1)) and auto.model_validate(plan(3))
    assert three.model_validate(plan(3)) and two.model_validate(plan(2))
    with pytest.raises(ValidationError):
        two.model_validate(plan(3))
    assert Inputs.parse({"text": "Content"}).slide_count is None
    TextToPptx.input_schema.validate_inputs({"text": "Content"})
    fields = TextToPptx.input_schema.conversation()
    assert fields[0].name == "slide_count" and not fields[0].required and fields[0].single_line
    assert fields[1].name == "text"


def test_count_metadata_cannot_introduce_source_numbers_or_weaken_fidelity():
    schema = source_plan_schema("Source content", 3)
    bad = plan(3)
    bad["deck"]["slides"][0]["bullets"] = ["Source content 3"]
    with pytest.raises(ValidationError, match="numeric/contact"):
        schema.model_validate(bad)
    for token in ("999", "changed@example.invalid"):
        bad = plan(3)
        bad["deck"]["slides"][1]["bullets"] = [token]
        with pytest.raises(ValidationError, match="numeric/contact"):
            source_plan_schema("00123 source@example.invalid", 3).model_validate(bad)


@pytest.mark.parametrize("count", [None, "1", "٣"])
async def test_service_uses_one_request_and_passes_trusted_count_separately(tmp_path, count):
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    calls = []

    async def extract(schema, source, prompt):
        calls.append((source, prompt))
        return schema.model_validate(plan(1 if count is None else int(count)))

    service = TextToPptx()
    service.runtime = SimpleNamespace(storage=store, ai=SimpleNamespace(extract=extract))
    values = {"text": "Source content", **({"slide_count": count} if count is not None else {})}
    service.input_schema.validate_inputs(values)
    result = await service.run(values)
    assert len(calls) == 1 and calls[0][0] == "Source content"
    assert len(Presentation(BytesIO(store.read(result.artifacts[0].key))).slides) == (
        1 if count is None else int(count)
    )
    if count is None:
        assert calls[0][1] == DECK
    else:
        assert f"exactly {int(count)} slides" in calls[0][1]


@pytest.mark.parametrize("actual", [2, 3])
async def test_continuation_enforces_count_without_another_ai_call(tmp_path, actual):
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    ai = SimpleNamespace(extract=AsyncMock(side_effect=AssertionError("no provider call")))
    service = TextToPptx()
    service.runtime = SimpleNamespace(storage=store, ai=ai)
    values = {"text": "Source content", "slide_count": "3", "__continuation": plan(actual)}
    assert not service.needs_ai(values)
    if actual == 2:
        with pytest.raises(ServiceError, match="provider_invalid"):
            await service.run(values)
        assert not list(tmp_path.glob("[0-9]*/*/*"))
    else:
        result = await service.run(values)
        assert len(Presentation(BytesIO(store.read(result.artifacts[0].key))).slides) == 3
    ai.extract.assert_not_called()


async def test_ignored_schema_by_injected_adapter_cannot_save_wrong_count(tmp_path):
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service = TextToPptx()
    service.runtime = SimpleNamespace(
        storage=store,
        ai=SimpleNamespace(extract=AsyncMock(return_value=DeckPlan.model_validate(plan(2)))),
    )
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run({"text": "Source content", "slide_count": "3"})
    assert not list(tmp_path.glob("[0-9]*/*/*"))


def test_missing_information_still_has_no_forced_placeholder_deck():
    assert (
        source_plan_schema("Content", 3).model_validate({"missing_information": True}).deck is None
    )
