"""Date display controls must never change the characters inside source contacts."""

from io import BytesIO

import pytest
from pptx import Presentation

from app.builders.pptx import build
from app.builders.schema import Deck


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("cover", [False, True])
@pytest.mark.parametrize(
    "contact",
    [
        "https://example.invalid/2026-10-24",
        "www.example.invalid/2026-10-24",
        "r.2026-10-24@e.invalid",
        "2026-10-24@e.invalid",
    ],
)
def test_arabic_titles_bullets_and_cells_preserve_dated_contacts(cover, contact):
    text = "راجع " + contact
    slides = [{"title": text, "bullets": [text]}]
    if cover:
        slides.append({"title": "التالي", "bullets": ["تفاصيل المصدر"]})
    slides.append({"title": "المراجع", "table": {"columns": ["الرابط"], "rows": [[text]]}})
    # Three bullets keep the first slide on the non-cover title/body path.
    if not cover:
        slides[0]["bullets"] = [text, "تفاصيل المصدر", "ملاحظة المصدر"]
    source = Deck.model_validate({"slides": slides})
    original = source.model_dump_json()
    deck = Presentation(BytesIO(build(source)))
    for name in ("Slide title", "Slide content"):
        shape = next(shape for shape in deck.slides[0].shapes if shape.name == name)
        assert shape.text_frame.paragraphs[0].text == text
    table = next(shape.table for shape in deck.slides[-1].shapes if shape.has_table)
    assert table.cell(1, 0).text == text
    saved = BytesIO()
    deck.save(saved)
    reopened = Presentation(BytesIO(saved.getvalue()))
    assert (
        next(shape.table for shape in reopened.slides[-1].shapes if shape.has_table).cell(1, 0).text
        == text
    )
    assert source.model_dump_json() == original


@pytest.mark.parametrize(
    "date", ["2026-10-24", "2026/10/24", "24-10-2026", "24/10/2026", "٢٠٢٦/١٠/٢٤"]
)
def test_arabic_prose_isolates_complete_date_tokens_without_reformatting(date):
    text = "الموعد " + date
    deck = Presentation(
        BytesIO(build(Deck.model_validate({"slides": [{"title": text, "bullets": [text]}]})))
    )
    for name in ("Slide title", "Slide content"):
        shape = next(shape for shape in deck.slides[0].shapes if shape.name == name)
        assert shape.text == "الموعد \u200e" + date + "\u200e"


def test_english_dates_and_standalone_contacts_have_no_new_display_controls():
    texts = ["Due 2026/10/24", "https://example.invalid/2026-10-24", "2026/10/24"]
    deck = Presentation(
        BytesIO(build(Deck.model_validate({"slides": [{"title": "Dates", "bullets": texts}]})))
    )
    body = next(shape for shape in deck.slides[0].shapes if shape.name == "Slide content")
    assert [paragraph.text for paragraph in body.text_frame.paragraphs] == texts


def test_same_arabic_paragraph_keeps_contacts_exact_and_isolates_only_free_dates():
    contact = "https://e.invalid/2026-10-24"
    text = f"راجع {contact} قبل 24/10/2026"
    source = Deck.model_validate({"slides": [{"title": "الموعد", "bullets": [text]}]})
    deck = Presentation(BytesIO(build(source)))
    body = next(shape for shape in deck.slides[0].shapes if shape.name == "Slide content")
    assert body.text == f"راجع {contact} قبل \u200e24/10/2026\u200e"
    assert source.slides[0].bullets == [text]
