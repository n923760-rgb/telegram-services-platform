"""Wrong-language decks use one repair and settle/release an order once."""

import json
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import CostHold, CostUsage, Job, Ledger, Order
from app.orders.engine import submit
from app.providers.ai.gateway import AI
from app.providers.storage import OwnedStorage
from app.wallet.ledger import balance
from app.workers.runner import deliver_failures, execute_job
from tests.test_foundation import job_for
from tests.test_pptx_language_unit import ARABIC, ENGLISH, candidate
from tests.test_real_services import setup


@pytest.mark.parametrize("source,arabic", [(ENGLISH, False), (ARABIC, True)])
@pytest.mark.parametrize("repair", [True, False])
async def test_language_repair_capture_release_cleanup_and_replay(
    tmp_path, monkeypatch, source, arabic, repair
):
    wrong = candidate(not arabic)
    correct = {
        "deck": {
            "slides": [
                {
                    "title": "حالة الدفع" if arabic else "Payment status",
                    "bullets": source.splitlines(),
                }
            ]
        }
    }
    responses = [json.dumps(wrong), json.dumps(correct if repair else wrong)]
    if repair:
        responses.append('{"accepted": true, "issues": []}')
    store, provider, delivery, ctx, price = await setup(
        "text_to_pptx", responses, tmp_path, monkeypatch
    )
    ctx["runtime_factory"] = lambda jid, uid: SimpleNamespace(
        ai=AI(provider, jid, uid, quality_review=True), storage=OwnedStorage(store, uid)
    )
    oid = await submit(1, "text_to_pptx", {"text": source, "slide_count": "1"}, price, "language")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    calls = 3 if repair else 2
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == ("completed" if repair else "failed")
        assert order.error_key == (None if repair else "provider_invalid")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if repair else 1000, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004") * calls
        assert await db.scalar(select(func.count(CostUsage.id))) == calls
        assert list(await db.scalars(select(CostHold.status))) == ["settled"] * calls
        for kind, expected in (
            ("reserve", 1),
            ("capture", int(repair)),
            ("release", int(not repair)),
        ):
            assert (
                await db.scalar(
                    select(func.count(Ledger.id)).where(Ledger.order_id == oid, Ledger.kind == kind)
                )
                == expected
            )
    assert provider.calls == calls
    assert len(delivery.files) == int(repair)
    assert len(delivery.errors) == int(not repair)
    assert not list(tmp_path.glob("[0-9]*/*/*"))
    if repair:
        deck = Presentation(BytesIO(delivery.files[0][1]))
        body = next(shape.text for shape in deck.slides[0].shapes if shape.name == "Slide content")
        assert body == source
