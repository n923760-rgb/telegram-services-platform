"""Whole-deck script switches are not authorized translations."""

import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.providers.ai.gateway import AI
from app.providers.ai.validation import validation_categories
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.services.text_to_pptx.schema import source_plan_schema
from app.services.text_to_pptx.service import TextToPptx
from tests.test_schema_repair_context_unit import Provider
from tests.test_schema_repair_context_unit import ledger as ledger


@pytest.fixture(autouse=True)
def database():
    yield


ENGLISH = "Order 00123\nAmount 125.50 SAR\nThe payment has not been approved."
ARABIC = "الطلب 00123\nالمبلغ 125.50\nلم تتم الموافقة على الدفعة."


def candidate(arabic):
    return {
        "deck": {
            "slides": [
                {
                    "title": "حالة الدفع" if arabic else "Payment status",
                    "bullets": [
                        "لم تتم الموافقة على الدفعة."
                        if arabic
                        else "The payment has not been approved."
                    ],
                }
            ]
        }
    }


@pytest.mark.parametrize("source,arabic", [(ENGLISH, True), (ARABIC, False)])
def test_whole_deck_translation_rejects_with_content_free_repair_category(source, arabic):
    schema = source_plan_schema(source)
    with pytest.raises(ValidationError) as error:
        schema.model_validate(candidate(arabic))
    assert validation_categories(schema, error.value) == ["pptx_source_language"]


@pytest.mark.parametrize("source,arabic", [(ENGLISH, False), (ARABIC, True)])
def test_original_source_language_is_accepted(source, arabic):
    assert source_plan_schema(source).model_validate(candidate(arabic)).deck


@pytest.mark.parametrize(
    "source",
    [
        ENGLISH + "\nملاحظة للمراجعة",
        ARABIC + "\nPending payment",
        "00123",
        "Short",
        "客户订单 " + ENGLISH,
    ],
)
@pytest.mark.parametrize("arabic", [False, True])
def test_mixed_other_script_short_and_numeric_sources_are_left_to_semantic_review(source, arabic):
    assert source_plan_schema(source).model_validate(candidate(arabic)).deck


def test_contacts_are_not_language_evidence_and_missing_information_is_not_a_deck():
    assert source_plan_schema(ENGLISH + "\ncontact@example.invalid").model_validate(
        candidate(False)
    )
    assert source_plan_schema(ARABIC + "\ncontact@example.invalid").model_validate(candidate(True))
    with pytest.raises(ValidationError):
        source_plan_schema(ARABIC + "\ncontact@example.invalid").model_validate(candidate(False))
    assert source_plan_schema(ENGLISH).model_validate({"missing_information": True}).deck is None


def test_guard_covers_table_headers_cells_and_titles_but_not_clarification_metadata():
    translated = {
        "deck": {
            "slides": [
                {
                    "title": "حالة الدفع",
                    "table": {
                        "columns": ["الحالة"],
                        "rows": [["لم تتم الموافقة على الدفعة."]],
                    },
                }
            ]
        }
    }
    with pytest.raises(ValidationError):
        source_plan_schema(ENGLISH).model_validate(translated)
    original = candidate(False) | {"ambiguous": True, "question": "هل تفضل هذا التنظيم؟"}
    assert source_plan_schema(ENGLISH).model_validate(original).deck
    # The language expectation is available to the provider without embedding source text.
    schema = json.dumps(source_plan_schema(ENGLISH).model_json_schema())
    assert "English" in schema and "00123" not in schema


@pytest.mark.parametrize("source,arabic", [(ENGLISH, False), (ARABIC, True)])
@pytest.mark.parametrize("repair", [True, False])
async def test_service_repairs_once_or_stores_no_file_and_accounts_for_each_call(
    tmp_path, ledger, source, arabic, repair
):
    wrong = candidate(not arabic)
    correct = {
        "deck": {
            "slides": [
                {
                    "title": "حالة الدفع" if arabic else "Payment status",
                    "bullets": source.splitlines(),
                }
            ]
        }
    }
    replies = [json.dumps(wrong), json.dumps(correct if repair else wrong)]
    if repair:
        replies.append('{"accepted": true, "issues": []}')
    provider = Provider(replies)
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service = TextToPptx()
    service.runtime = SimpleNamespace(ai=AI(provider, 1, 1, quality_review=True), storage=store)
    if repair:
        result = await service.run({"text": source, "slide_count": "1"})
        deck = Presentation(BytesIO(store.read(result.artifacts[0].key)))
        body = next(shape.text for shape in deck.slides[0].shapes if shape.name == "Slide content")
        assert body == source
        assert store.storage.usage(1)[1] == 1
    else:
        with pytest.raises(ServiceError, match="provider_invalid"):
            await service.run({"text": source, "slide_count": "1"})
        assert store.storage.usage(1)[1] == 0
    calls = 3 if repair else 2
    reserve, settle, uncertain = ledger
    assert provider.calls == reserve.await_count == settle.await_count == calls
    uncertain.assert_not_awaited()
    assert "pptx_source_language" in provider.requests[1][0]
    assert (
        "Keep the source prose in " + ("Arabic" if arabic else "English") in provider.requests[0][1]
    )


@pytest.mark.parametrize("source,arabic", [(ENGLISH, True), (ARABIC, False)])
async def test_unprepared_translated_continuation_rejects_without_provider_or_file(
    tmp_path, source, arabic
):
    service = TextToPptx()
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service.runtime = SimpleNamespace(ai=None, storage=store)
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run({"text": source, "__continuation": candidate(arabic)})
    assert store.storage.usage(1)[1] == 0
