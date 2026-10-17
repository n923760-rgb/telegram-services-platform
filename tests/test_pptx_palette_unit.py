from io import BytesIO

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from app.builders import pptx
from app.builders.schema import Deck
from tests.test_pptx_quality import (
    sample_plan,
)
from tests.test_pptx_quality import (
    test_pptx_native_text_table_source_values_fonts_and_editability as assert_editable_structure,
)


@pytest.fixture(autouse=True)
def database():
    yield


def luminance(color):
    values = [value / 255 for value in color]
    return sum(
        weight * (value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4)
        for weight, value in zip((0.2126, 0.7152, 0.0722), values, strict=True)
    )


@pytest.mark.parametrize(
    "foreground,background",
    [
        (pptx.COVER_TEXT, pptx.COVER),
        (pptx.COVER_MUTED, pptx.COVER),
        (pptx.DARK, pptx.PAPER),
        (pptx.ACCENT, pptx.PAPER),
        (pptx.MUTED, pptx.PAPER),
        (pptx.COVER_TEXT, pptx.ACCENT),
    ],
)
def test_native_palette_retains_readable_contrast(foreground, background):
    high, low = sorted([luminance(foreground), luminance(background)], reverse=True)
    assert (high + 0.05) / (low + 0.05) >= 4.5


def test_backgrounds_table_borders_and_all_original_values_are_native():
    data = Deck.model_validate(sample_plan())
    prs = Presentation(BytesIO(pptx.build(data)))
    assert [str(s.background.fill.fore_color.rgb) for s in prs.slides] == [
        str(pptx.COVER),
        str(pptx.PAPER),
        str(pptx.PAPER),
    ]
    title = next(s for s in prs.slides[0].shapes if s.name == "Slide title")
    assert title.text_frame.paragraphs[0].runs[0].font.color.rgb == pptx.COVER_TEXT
    table = next(s.table for s in prs.slides[2].shapes if s.has_table)
    for row in table.rows:
        for cell in row.cells:
            assert all(
                cell._tc.tcPr.find(qn("a:ln" + edge)) is not None for edge in ("L", "R", "T", "B")
            )
            children = list(cell._tc.tcPr)
            fill_index = children.index(cell._tc.tcPr.find(qn("a:solidFill")))
            assert all(
                children.index(cell._tc.tcPr.find(qn("a:ln" + edge))) < fill_index
                for edge in ("L", "R", "T", "B")
            )
    assert not any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)
    assert_editable_structure()


def test_single_slide_uses_paper_instead_of_an_unrequested_cover():
    prs = Presentation(
        BytesIO(pptx.build(Deck.model_validate({"slides": [sample_plan()["slides"][1]]})))
    )
    assert prs.slides[0].background.fill.fore_color.rgb == pptx.PAPER


@pytest.mark.parametrize(
    "text,marked", [("تم التسليم في 2026-10-02", True), ("Delivered 2026-10-02", False)]
)
def test_iso_date_direction_marks_preserve_exact_visible_source(text, marked):
    data = Deck.model_validate({"slides": [{"title": "Records", "bullets": [text]}]})
    prs = Presentation(BytesIO(pptx.build(data)))
    body = next(s.text for s in prs.slides[0].shapes if s.name == "Slide content")
    assert body.replace("\u200e", "") == text
    assert ("\u200e2026-10-02\u200e" in body) is marked
