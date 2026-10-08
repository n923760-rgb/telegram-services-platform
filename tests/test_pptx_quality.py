"""Readability planning plus reopened native text/table structure and editing."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pydantic import ValidationError

from app.builders import pptx
from app.builders.pptx_layout import BODY_HEIGHT, BODY_WIDTH, plan_slide, text_height
from app.builders.schema import Deck


def sample_plan():
    return {
        "slides": [
            {"title": "مراجعة الطلبات", "bullets": ["بيانات الطلبات من 2026-10-01 إلى 2026-10-08"]},
            {
                "title": "ملخص المتابعة",
                "bullets": [
                    "الطلب 00123 بمبلغ 125.50 ريال، تم تسليمه في 2026-10-02.",
                    "الطلب 00124 يحتاج مراجعة البيانات. لا يوجد موعد تسليم مؤكد.",
                    "Customer note: retain all source identifiers and dates.",
                ],
            },
            {
                "title": "الطلبات",
                "table": {
                    "columns": ["الرقم", "العميل", "المبلغ (ريال)", "التاريخ"],
                    "rows": [
                        ["00123", "أحمد", "125.50", "2026-10-02"],
                        ["00124", "Example", "0", ""],
                        ["00123", "أحمد", "125.50", "2026-10-02"],
                    ],
                },
            },
        ]
    }


def test_pptx_native_text_table_source_values_fonts_and_editability():
    data = Deck.model_validate(sample_plan())
    content = pptx.build(data)
    prs = Presentation(BytesIO(content))
    assert len(prs.slides) == 3
    assert plan_slide(data.slides[0], cover_candidate=True).kind == "cover"
    for index, slide in enumerate(prs.slides):
        shapes = {shape.name: shape for shape in slide.shapes}
        assert shapes["Slide title"].text == data.slides[index].title
        assert shapes["Slide number"].text == f"{index + 1} / 3"
        footer = shapes["Slide number"].text_frame
        assert footer.margin_top == 0 and footer.margin_bottom == 0
        for shape in slide.shapes:
            assert 0 <= shape.left and 0 <= shape.top
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
            if shape.has_text_frame:
                assert shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE
                if shape.name != "Slide number":
                    assert all(
                        run.font.size.pt >= 20
                        for p in shape.text_frame.paragraphs
                        for run in p.runs
                    )
    body = next(s for s in prs.slides[1].shapes if s.name == "Slide content")
    assert [p.text for p in body.text_frame.paragraphs] == data.slides[1].bullets
    assert [p._p.pPr.get("rtl") for p in body.text_frame.paragraphs] == ["1", "1", "0"]
    for paragraph in body.text_frame.paragraphs:
        assert paragraph._p.pPr.find(qn("a:buChar")) is not None
        for run in paragraph.runs:
            assert run._r.rPr.find(qn("a:cs")).get("typeface") == "DejaVu Sans"
    native = next(shape.table for shape in prs.slides[2].shapes if shape.has_table)
    table = data.slides[2].table
    assert [[cell.text for cell in row.cells] for row in native.rows] == [
        table.columns,
        *table.rows,
    ]
    assert native._tbl.tblPr.get("rtl") == "1"
    assert native.cell(2, 3).text == "" and native.cell(2, 2).text == "0"
    assert all(
        run.font.size.pt >= 20
        for row in native.rows
        for cell in row.cells
        for p in cell.text_frame.paragraphs
        for run in p.runs
    )
    native.cell(1, 2).text = "250.75"
    body.text_frame.paragraphs[0].runs[0].text = "Updated editable text"
    saved = BytesIO()
    prs.save(saved)
    reopened = Presentation(BytesIO(saved.getvalue()))
    assert (
        next(s.table for s in reopened.slides[2].shapes if s.has_table).cell(1, 2).text == "250.75"
    )
    assert "Updated editable text" in next(
        s.text for s in reopened.slides[1].shapes if s.name == "Slide content"
    )
    with ZipFile(BytesIO(content)) as package:
        assert not any(
            "vba" in name.lower() or "embeddings" in name.lower() or "media/" in name
            for name in package.namelist()
        )
        assert not any(
            b'TargetMode="External"' in package.read(name)
            for name in package.namelist()
            if name.endswith(".rels")
        )


def test_single_slide_and_dense_first_slide_are_content_without_extra_cover_or_data_loss():
    for count in (1, 2, 30):
        first = {
            "title": "بيانات المتابعة",
            "bullets": [f"الطلب {i:05d}: محتوى مؤكد للمراجعة" for i in range(6)],
        }
        data = Deck.model_validate({"slides": [first] * count})
        prs = Presentation(BytesIO(pptx.build(data)))
        assert len(prs.slides) == count
        assert plan_slide(data.slides[0], cover_candidate=count > 1).kind == "content"
        assert (
            next(s.text_frame for s in prs.slides[0].shapes if s.name == "Slide content")
            .paragraphs[5]
            .text
            == first["bullets"][5]
        )


@pytest.mark.parametrize(
    "invalid",
    [
        {"title": "Title", "bullets": ["W" * 180] * 6},
        {"title": "W" * 150, "bullets": ["Short"]},
        {"title": "Title", "bullets": ["First\n" + "\n" * 50 + "Text"]},
        {"title": "Title", "bullets": ["Bad\x00text"]},
        {"title": "Title", "bullets": ["   "]},
        {"title": "Title"},
        {"title": "Title", "bullets": ["Text"], "table": {"columns": ["A"], "rows": [["B"]]}},
        {"title": "Title", "table": {"columns": ["A", "A"], "rows": [["1", "2"]]}},
        {"title": "Title", "table": {"columns": ["A", "B"], "rows": [["1"]]}},
        {"title": "Title", "table": {"columns": ["A"], "rows": [["x" * 81]]}},
        {"title": "Title", "table": {"columns": list("ABCDE"), "rows": [["x"] * 5]}},
        {"title": "Title", "table": {"columns": ["A"], "rows": [["x"]] * 7}},
        {"title": "Title", "table": {"columns": list("ABCD"), "rows": [["W" * 80] * 4] * 6}},
    ],
)
def test_pptx_rejects_dense_or_invalid_plans_instead_of_shrinking_or_omitting(invalid):
    with pytest.raises(ValidationError):
        Deck.model_validate({"slides": [invalid]})


def test_pptx_metric_height_fits_bounded_layout_and_unbroken_tokens():
    data = Deck.model_validate(sample_plan())
    for slide in data.slides:
        layout = plan_slide(slide)
        assert layout.body_size >= 20 and layout.title_size >= 32
        if slide.table is None:
            assert text_height(slide.bullets, layout.body_size, BODY_WIDTH - 24) <= BODY_HEIGHT
        else:
            assert sum(layout.row_heights) <= BODY_HEIGHT
    data = Deck.model_validate({"slides": [{"title": "English", "bullets": ["ID " + "0" * 150]}]})
    prs = Presentation(BytesIO(pptx.build(data)))
    assert (
        next(s.text for s in prs.slides[0].shapes if s.name == "Slide content")
        == data.slides[0].bullets[0]
    )


def test_english_table_remains_ltr_and_does_not_execute_xml_or_links():
    data = Deck.model_validate(
        {
            "slides": [
                {
                    "title": "Records",
                    "table": {
                        "columns": ["ID", "Note"],
                        "rows": [
                            ["00123", "<xml>& =SUM(A1)"],
                            ["00124", "https://example.invalid"],
                        ],
                    },
                }
            ]
        }
    )
    native = next(
        s.table for s in Presentation(BytesIO(pptx.build(data))).slides[0].shapes if s.has_table
    )
    assert native._tbl.tblPr.get("rtl") == "0"
    assert native.cell(1, 1).text == "<xml>& =SUM(A1)"
