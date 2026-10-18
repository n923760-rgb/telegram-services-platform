"""Native PPTX quality gates preserve repair, capture/release and cached delivery."""

import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from pptx import Presentation
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import CostHold, CostUsage, Job, Ledger, Order, Service
from app.orders.engine import confirm_structure, submit
from app.services.base import ServiceError
from app.services.registry import registry
from app.wallet.ledger import balance
from app.workers.runner import deliver_failures, execute_job
from tests.test_foundation import job_for
from tests.test_pptx_quality import sample_plan, sample_source
from tests.test_real_services import setup


@pytest.mark.parametrize("repair", [False, True])
async def test_pptx_native_table_delivery_captures_once_after_bounded_repair(
    tmp_path, monkeypatch, repair
):
    responses = [json.dumps({"deck": sample_plan()})]
    if repair:
        responses.insert(
            0, json.dumps({"deck": {"slides": [{"title": "Dense", "bullets": ["W" * 180] * 6}]}})
        )
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", responses, tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": sample_source()}, price, "native-pptx")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == len(responses) and len(delivery.files) == 1
    prs = Presentation(BytesIO(delivery.files[0][1]))
    native = next(s.table for s in prs.slides[2].shapes if s.has_table)
    assert native.cell(1, 0).text == "00123" and native.cell(1, 2).text == "125.50"
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_pptx_density_invalid_after_one_repair_releases_and_notifies_once(
    tmp_path, monkeypatch
):
    invalid = json.dumps({"deck": {"slides": [{"title": "Dense", "bullets": ["W" * 180] * 6}]}})
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [invalid, invalid], tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": "Content"}, price, "dense-pptx")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 2 and not delivery.files
    assert delivery.errors == [(1, oid, "provider_invalid")]


async def test_pptx_builder_failure_does_not_capture_credit(tmp_path, monkeypatch):
    from app.builders import pptx

    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [json.dumps({"deck": sample_plan()})], tmp_path, monkeypatch
    )

    def failed(*args):
        raise ValueError("synthetic render failure")

    monkeypatch.setattr(pptx, "build", failed)
    oid = await submit(1, "text_to_pptx", {"text": sample_source()}, price, "render-pptx")
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 1 and not delivery.files and len(delivery.errors) == 1


async def test_pptx_delivery_retry_reuses_native_file_before_single_capture(tmp_path, monkeypatch):
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [json.dumps({"deck": sample_plan()})], tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_pptx", {"text": sample_source(), "slide_count": "3"}, price, "retry-pptx"
    )
    jid = await job_for(oid)
    original_send = delivery.send

    async def unavailable(*args):
        raise ConnectionError("synthetic delivery unavailable")

    monkeypatch.setattr(delivery, "send", unavailable)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "delivering"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, price)
    monkeypatch.setattr(delivery, "send", original_send)
    async with sessions.begin() as db:
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC) - timedelta(seconds=1)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == 1 and len(delivery.files) == 1
    assert any(s.has_table for s in Presentation(BytesIO(delivery.files[0][1])).slides[2].shapes)


@pytest.mark.parametrize("count", [1, 3])
@pytest.mark.parametrize("repair", [False, True])
async def test_explicit_count_delivers_exact_native_slides_and_captures_once(
    tmp_path, monkeypatch, count, repair
):
    from tests.test_pptx_count_unit import plan

    responses = [json.dumps(plan(count))]
    if repair:
        responses.insert(0, json.dumps(plan(2)))
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", responses, tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_pptx", {"text": "Source content", "slide_count": str(count)}, price, "count"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004") * len(responses)
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
        assert (
            await db.scalar(
                select(func.count(Ledger.id)).where(
                    Ledger.kind == "capture", Ledger.order_id == oid
                )
            )
            == 1
        )
    assert provider.calls == len(responses) and len(delivery.files) == 1
    assert len(Presentation(BytesIO(delivery.files[0][1])).slides) == count


async def test_wrong_count_after_existing_single_repair_releases_credit_without_file(
    tmp_path, monkeypatch
):
    from tests.test_pptx_count_unit import plan

    bad = json.dumps(plan(2))
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [bad, bad, "unused"], tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_pptx", {"text": "Source content", "slide_count": "3"}, price, "bad-count"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        assert (await db.get(Job, jid)).cost_sar == Decimal(".008")
        assert await db.scalar(select(func.count(CostUsage.id))) == 2
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 2 and len(provider.responses) == 1
    assert not delivery.files and not list(tmp_path.glob("[0-9]*/*/*"))


@pytest.mark.parametrize("count", ["0", "31", "3.0", "٣١"])
async def test_bad_count_has_no_order_cost_hold_or_reservation(tmp_path, monkeypatch, count):
    _, provider, _, _, price = await setup("text_to_pptx", [], tmp_path, monkeypatch)
    with pytest.raises(ServiceError, match="input_invalid"):
        await submit(
            1,
            "text_to_pptx",
            {"text": "Source content", "slide_count": count},
            price,
            "invalid-count",
        )
    async with sessions() as db:
        assert await db.scalar(select(func.count(Order.id))) == 0
        assert await db.scalar(select(func.count(CostHold.id))) == 0
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 0


async def test_explicit_count_approval_reuses_validated_plan_without_another_call(
    tmp_path, monkeypatch
):
    from tests.test_pptx_count_unit import plan

    response = {**plan(3), "ambiguous": True, "question": "اعتماد الهيكل؟"}
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [json.dumps(response)], tmp_path, monkeypatch
    )
    oid = await submit(
        1, "text_to_pptx", {"text": "Source content", "slide_count": "٣"}, price, "count-approval"
    )
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
    await confirm_structure(oid, 1, True)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 1 and len(Presentation(BytesIO(delivery.files[0][1])).slides) == 3


async def test_prepared_v4_file_replays_after_upgrade_with_admin_prices_preserved(
    tmp_path, monkeypatch
):
    from app.services.text_to_pptx.service import TextToPptx
    from tests.test_pptx_count_unit import plan

    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [json.dumps(plan(1))], tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": "Source content"}, price, "v4-cache")
    jid = await job_for(oid)
    send = delivery.send

    async def offline(*args):
        raise ConnectionError("synthetic delivery failure")

    monkeypatch.setattr(delivery, "send", offline)
    await execute_job(ctx, str(jid))
    async with sessions.begin() as db:
        row = await db.get(Service, "text_to_pptx")
        old_schema = {
            **row.input_schema,
            "fields": [f for f in row.input_schema["fields"] if f["name"] != "slide_count"],
        }
        row.version, row.price_halala, row.price_stars, row.input_schema = "4", 777, 25, old_schema
        order = await db.get(Order, oid)
        order.service_version, order.input_schema_snapshot = "4", old_schema
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC)
    async with sessions.begin() as db:
        await registry.sync(db)
        row = await db.get(Service, "text_to_pptx")
        assert row.version == "5" and row.enabled
        assert (row.price_halala, row.price_stars) == (777, 25)

    async def forbidden(*args):
        raise AssertionError("historical file must not be regenerated")

    monkeypatch.setattr(TextToPptx, "run", forbidden)
    monkeypatch.setattr(delivery, "send", send)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price, 0)
    assert provider.calls == 1 and len(Presentation(BytesIO(delivery.files[0][1])).slides) == 1


@pytest.mark.parametrize("old_version", ["1", "3", "4"])
async def test_pptx_version_upgrade_keeps_admin_settings_and_releases_old_pending_job(
    tmp_path, monkeypatch, old_version
):
    _, provider, delivery, ctx, price = await setup("text_to_pptx", [], tmp_path, monkeypatch)
    oid = await submit(1, "text_to_pptx", {"text": "Content"}, price, "old-pptx")
    async with sessions.begin() as db:
        row = await db.get(Service, "text_to_pptx")
        row.version, row.price_halala, row.enabled = old_version, 777, True
        (await db.get(Order, oid)).service_version = old_version
    async with sessions.begin() as db:
        await registry.sync(db)
    async with sessions() as db:
        row = await db.get(Service, "text_to_pptx")
        assert row.version == "5" and row.price_halala == 777 and row.enabled
    await execute_job(ctx, str(await job_for(oid)))
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000, 0)
    assert provider.calls == 0 and not delivery.files


@pytest.mark.parametrize("text", ["\x00", "bad\x1f", "\ud800", "\uffff"])
async def test_pptx_invalid_source_never_creates_order_hold_or_reservation(
    tmp_path, monkeypatch, text
):
    _, provider, _, _, price = await setup("text_to_pptx", [], tmp_path, monkeypatch)
    with pytest.raises(ServiceError, match="input_invalid"):
        await submit(1, "text_to_pptx", {"text": text}, price, "invalid-source")
    async with sessions() as db:
        assert await db.scalar(select(func.count(Order.id))) == 0
        assert await db.scalar(select(func.count(CostHold.id))) == 0
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 0


@pytest.mark.parametrize("repair", [False, True])
async def test_pptx_invented_amount_repair_or_release_records_cost_once(
    tmp_path, monkeypatch, repair
):
    bad = json.dumps({"deck": {"slides": [{"title": "Invoice", "bullets": ["Amount 999.50"]}]}})
    good = json.dumps({"deck": {"slides": [{"title": "Invoice", "bullets": ["Amount 125.50"]}]}})
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [bad, good if repair else bad], tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": "Amount 125.50"}, price, "source-tokens")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    async with sessions() as db:
        result = await db.get(Order, oid)
        assert result.status == ("completed" if repair else "failed")
        if not repair:
            assert result.error_key == "provider_invalid"
        assert (await db.get(Job, jid)).cost_sar == Decimal(".008")
        assert await db.scalar(select(func.count(CostUsage.id))) == 2
        assert (
            await db.scalar(
                select(func.count(Ledger.id)).where(
                    Ledger.kind == ("capture" if repair else "release"), Ledger.order_id == oid
                )
            )
            == 1
        )
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if repair else 1000, 0)
    assert provider.calls == 2 and len(delivery.files) == (1 if repair else 0)
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_pptx_source_bound_approval_reuses_plan_without_provider(tmp_path, monkeypatch):
    response = {
        "ambiguous": True,
        "question": "اعتماد هذا الهيكل؟",
        "deck": {"slides": [{"title": "Invoice", "bullets": ["Amount 125.50"]}]},
    }
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx", [json.dumps(response)], tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": "Amount 125.50"}, price, "source-approval")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
    await confirm_structure(oid, 1, True)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 1 and len(delivery.files) == 1
