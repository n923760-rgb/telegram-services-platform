"""Native, editable table geometry for uneven source-column content."""

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError

from app.builders import pptx
from app.builders.pptx_layout import BODY_HEIGHT, BODY_WIDTH, plan_slide, wrapped_lines
from app.builders.schema import Deck
from tests.pptx_table_fixtures import table_slide


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("language", ["ar", "en"])
def test_uneven_table_fits_without_shrinking_or_losing_source(language):
    source = table_slide(language)
    deck = Deck.model_validate({"slides": [source]})
    layout = plan_slide(deck.slides[0])
    assert sum(layout.row_heights) <= BODY_HEIGHT
    assert layout.column_widths[1] > layout.column_widths[0]
    assert sum(layout.column_widths) == pytest.approx(BODY_WIDTH)
    assert min(layout.column_widths) >= 90
    for index, row in enumerate([source["table"]["columns"], *source["table"]["rows"]]):
        size = 21 if index == 0 else 20
        required = (
            max(
                wrapped_lines(cell, size, width, bold=index == 0)
                for cell, width in zip(row, layout.column_widths, strict=True)
            )
            * size
            * 1.2
            + 16
        )
        assert layout.row_heights[index] >= required

    native = Presentation(BytesIO(pptx.build(deck)))
    shape = next(shape for shape in native.slides[0].shapes if shape.has_table)
    table = shape.table
    assert shape.width == Inches(11.6)
    assert sum(column.width for column in table.columns) == shape.width
    assert [column.width.pt for column in table.columns] == pytest.approx(
        layout.column_widths, abs=0.001
    )
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        source["table"]["columns"],
        *source["table"]["rows"],
    ]
    assert all(
        run.font.size.pt >= 20
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.text_frame.paragraphs
        for run in paragraph.runs
    )
    table.cell(1, 2).text = "250.75"
    saved = BytesIO()
    native.save(saved)
    reopened = Presentation(BytesIO(saved.getvalue()))
    edited = next(shape.table for shape in reopened.slides[0].shapes if shape.has_table)
    assert edited.cell(1, 2).text == "250.75"
    assert edited.cell(1, 0).text == "00000"


@pytest.mark.parametrize("columns", [1, 2, 3, 4])
def test_balanced_short_tables_keep_equal_widths(columns):
    source = {
        "title": "Records",
        "table": {
            "columns": [f"Column {index}" for index in range(columns)],
            "rows": [["Value"] * columns] * 2,
        },
    }
    slide = Deck.model_validate({"slides": [source]}).slides[0]
    layout = plan_slide(slide)
    assert layout.column_widths == pytest.approx([BODY_WIDTH / columns] * columns)


@pytest.mark.parametrize("columns,long_column", [(2, 0), (2, 1), (4, 0), (4, 3)])
def test_variable_columns_keep_order_multiline_headers_and_literal_cells(columns, long_column):
    headers = [f"Column\n{index}" for index in range(columns)]
    row = ["00017", "", "125.50", "2026-10-21"][:columns]
    row[long_column] = table_slide("en")["table"]["rows"][0][1]
    source = {"title": "Records", "table": {"columns": headers, "rows": [row] * 3}}
    plan = Deck.model_validate({"slides": [source]})
    layout = plan_slide(plan.slides[0])
    assert len(layout.column_widths) == columns
    assert min(layout.column_widths) >= 90
    assert sum(layout.row_heights) <= BODY_HEIGHT
    reopened = Presentation(BytesIO(pptx.build(plan)))
    table = next(shape.table for shape in reopened.slides[0].shapes if shape.has_table)
    # python-pptx exposes native soft line breaks as vertical tabs.
    assert [
        [cell.text.replace("\v", "\n") for cell in native_row.cells] for native_row in table.rows
    ] == [headers, row, row, row]


def test_dense_table_still_rejects_and_text_slides_keep_their_layout():
    with pytest.raises(ValidationError, match="too dense"):
        Deck.model_validate(
            {
                "slides": [
                    {
                        "title": "Records",
                        "table": {
                            "columns": list("ABCD"),
                            "rows": [["W" * 80] * 4] * 6,
                        },
                    }
                ]
            }
        )
    slide = Deck.model_validate({"slides": [{"title": "Title", "bullets": ["Content"]}]}).slides[0]
    assert plan_slide(slide).column_widths == ()
