"""Literal restoration precedes native validation and never repairs guesses."""

import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pptx import Presentation
from pydantic import ValidationError
from pydantic_core import PydanticCustomError

from app.providers.storage import LocalStorage, OwnedStorage
from app.services.document_fidelity import CONTACT_TOKEN, source_document_schema
from app.services.protected_values import ProtectedValues
from app.services.text_to_office.schema import WordPlan
from app.services.text_to_office.service import TextToOffice
from app.services.text_to_pdf.schema import PdfPlan
from app.services.text_to_pptx.schema import source_plan_schema
from app.services.text_to_pptx.service import TextToPptx


@pytest.fixture(autouse=True)
def database():
    yield


SOURCE = "Order 00123\nAmount 125.50 SAR\nThe payment has not been approved."


@pytest.mark.parametrize(
    "source",
    [
        SOURCE,
        "٠٠١٢٣ -12.50 12.5% ٠٠١٢٣",
        "+966500000001 a12@example.invalid https://demo.invalid/00123?a=12.50",
        "مراسلة:a12@example.invalid بتاريخ 2026-10-20 الساعة 12:30",
        "Invoice 00123\nInvoice 00123\nContact a@example.invalid a@example.invalid",
        "No numbers or contact information.",
        "Literal [[value-123456abcdef:0]] remains source content.",
    ],
)
def test_exact_round_trip_and_no_mutation_for_contacts_overlapping_numbers(source):
    protected = ProtectedValues(source, CONTACT_TOKEN)
    data = {"title": "Note", "sections": [{"paragraphs": [protected.masked]}]}
    snapshot = json.dumps(data)
    assert protected.resolve(data)["sections"][0]["paragraphs"] == [source]
    assert json.dumps(data) == snapshot
    if protected.values:
        request = json.loads(protected.request())
        assert request["source"] == protected.masked
        assert request["protected_values"] == protected.values
        assert "reference metadata" in protected.instruction()
    else:
        assert protected.request() == source and protected.instruction() == ""


@pytest.mark.parametrize("suffix", ["99999]]", "bad]]", "0", "0]", "]]"])
def test_unknown_or_incomplete_reference_is_rejected_without_source_values(suffix):
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    with pytest.raises(PydanticCustomError) as rejected:
        protected.resolve({"text": protected.prefix + suffix})
    assert rejected.value.type == "protected_value_reference"
    assert "00123" not in str(rejected.value)


def test_missing_information_does_not_require_a_fabricated_document():
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    schema = source_document_schema(WordPlan, SOURCE, protected=protected)
    assert schema.model_validate({"missing_information": True}).document is None


def test_deep_invalid_reference_tree_fails_without_recursion_failure():
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    value = protected.masked
    for _ in range(40):
        value = [value]
    with pytest.raises(PydanticCustomError):
        protected.resolve(value)


@pytest.mark.parametrize("kind", [WordPlan, PdfPlan])
def test_resolved_documents_preserve_original_fidelity_and_are_safe_to_revalidate(kind):
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    schema = source_document_schema(kind, SOURCE, protected=protected)
    encoded = {"document": {"title": "Order", "sections": [{"paragraphs": [protected.masked]}]}}
    result = schema.model_validate(encoded)
    assert result.document.sections[0].paragraphs == [SOURCE]
    assert schema.model_validate(result.model_dump(mode="json")) == result
    assert "00123" not in json.dumps(schema.model_json_schema())
    assert protected.prefix not in json.dumps(schema.model_json_schema())
    duplicate = {
        "document": {"title": "Order", "sections": [{"paragraphs": [protected.masked] * 2}]}
    }
    with pytest.raises(ValidationError, match="preserve source"):
        schema.model_validate(duplicate)
    literal_guess = {
        "document": {
            "title": "Order",
            "sections": [{"paragraphs": [SOURCE.replace("00123", "123")]}],
        }
    }
    with pytest.raises(ValidationError, match="preserve source"):
        schema.model_validate(literal_guess)


def test_pptx_references_restore_before_count_density_and_source_checks():
    protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
    schema = source_plan_schema(SOURCE, 1, protected=protected)
    data = {"deck": {"slides": [{"title": "Order", "bullets": protected.masked.splitlines()}]}}
    result = schema.model_validate(data)
    assert result.deck.slides[0].bullets == SOURCE.splitlines()
    assert schema.model_validate(result.model_dump(mode="json")) == result
    assert "00123" not in json.dumps(schema.model_json_schema())
    with pytest.raises(ValidationError, match="requested slide count"):
        schema.model_validate({"deck": {"slides": data["deck"]["slides"] * 2}})
    with pytest.raises(ValidationError, match="source tokens"):
        schema.model_validate(
            {"deck": {"slides": [{"title": "Order 00123", "bullets": SOURCE.splitlines()}]}}
        )


@pytest.mark.parametrize("kind", ["word", "pptx"])
def test_native_table_cells_restore_before_rectangular_and_source_validation(kind):
    source = "Order 00123\nAmount 125.50"
    protected = ProtectedValues(source, CONTACT_TOKEN)
    ids = list(protected.values)
    table = {"columns": ["Label", "Value"], "rows": [["Order", ids[0]], ["Amount", ids[1]]]}
    if kind == "word":
        schema = source_document_schema(WordPlan, source, protected=protected)
        result = schema.model_validate(
            {"document": {"title": "Order", "sections": [{"tables": [table]}]}}
        )
        rows = result.document.sections[0].tables[0].rows
    else:
        schema = source_plan_schema(source, 1, protected=protected)
        result = schema.model_validate({"deck": {"slides": [{"title": "Order", "table": table}]}})
        rows = result.deck.slides[0].table.rows
    assert rows == [["Order", "00123"], ["Amount", "125.50"]]


@pytest.mark.parametrize("kind", ["word", "pptx"])
async def test_real_editable_output_contains_literal_values_and_no_internal_references(
    tmp_path, kind
):
    calls = []

    async def extract(schema, source, prompt):
        calls.append((source, prompt))
        request = json.loads(source)
        masked = request["source"]
        assert request["protected_values"] and "reference markers" in prompt
        if kind == "word":
            payload = {
                "document": {"title": "Order", "sections": [{"paragraphs": masked.splitlines()}]}
            }
        else:
            payload = {"deck": {"slides": [{"title": "Order", "bullets": masked.splitlines()}]}}
        return schema.model_validate_json(json.dumps(payload))

    service = TextToOffice() if kind == "word" else TextToPptx()
    store = OwnedStorage(LocalStorage(tmp_path), 1)
    service.runtime = SimpleNamespace(storage=store, ai=SimpleNamespace(extract=extract))
    inputs = {"text": SOURCE}
    inputs.update({"target": "word", "mode": "smart"} if kind == "word" else {"slide_count": "1"})
    service.input_schema.validate_inputs(inputs)
    result = await service.run(inputs)
    content = BytesIO(store.read(result.artifacts[0].key))
    if kind == "word":
        text = "\n".join(paragraph.text for paragraph in Document(content).paragraphs)
    else:
        text = "\n".join(
            shape.text
            for slide in Presentation(content).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    assert text.count("00123") == text.count("125.50") == 1
    assert "not been approved" in text and "[[value-" not in text
    assert "[[value-" not in json.dumps(result.continuation)
    assert len(calls) == 1
