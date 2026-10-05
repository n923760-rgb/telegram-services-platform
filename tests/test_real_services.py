from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as Word
from openpyxl import load_workbook
from pptx import Presentation

from app.core.db import sessions
from app.core.models import Job, Order, Service
from app.core.settings import config
from app.orders.engine import confirm_structure, submit
from app.providers.ai.gateway import AI
from app.providers.storage import LocalStorage, OwnedStorage
from app.services.base import ServiceError
from app.wallet.ledger import balance
from app.workers.runner import deliver_confirmations, deliver_failures, execute_job
from tests.test_file_layers import Provider
from tests.test_foundation import fund, job_for


class Delivery:
    def __init__(self, store):
        self.store = store
        self.files = []
        self.results = []
        self.errors = []
        self.confirmations = []

    async def send(self, user_id, order_id, result):
        self.results.append(result)
        self.files.extend(
            (file.filename, self.store.read(file.key, user_id)) for file in result.artifacts
        )

    async def error(self, user_id, order_id, key):
        self.errors.append((user_id, order_id, key))

    async def confirmation(self, user_id, order_id, preview):
        self.confirmations.append((order_id, preview))


async def setup(slug, responses, tmp_path, monkeypatch):
    monkeypatch.setattr(config(), "ai_enabled", True)
    await fund()
    async with sessions.begin() as db:
        service = await db.get(Service, slug)
        service.enabled = True
        price = service.price_halala
    store = LocalStorage(tmp_path)
    provider = Provider(responses)
    delivery = Delivery(store)
    ctx = {
        "delivery": delivery,
        "runtime_factory": lambda jid, uid: SimpleNamespace(
            ai=AI(provider, jid, uid), storage=OwnedStorage(store, uid)
        ),
    }
    return store, provider, delivery, ctx, price


async def test_real_ocr_success_translation_and_capture(tmp_path, monkeypatch):
    store, provider, delivery, ctx, price = await setup(
        "image_to_text",
        ['{"readable":true,"confidence":.99,"text":"مرحبا 123"}', '{"text":"Hello 123"}'],
        tmp_path,
        monkeypatch,
    )
    # JSON requires a leading zero, so use the valid provider response below.
    provider.responses[0] = '{"readable":true,"confidence":0.99,"text":"مرحبا 123"}'
    file = store.save(1, "input.jpg", b"test image", "image/jpeg")
    oid = await submit(
        1,
        "image_to_text",
        {"images": [file.key], "language": "en", "style": "formal"},
        price,
        "ocr",
    )
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).available == 1000 - price
        assert (await db.get(Job, await job_for(oid))).cost_sar == Decimal(".008")
    assert delivery.results[0].text == "Hello 123"
    assert provider.calls == 2


async def test_real_ocr_unreadable_releases_and_notifies(tmp_path, monkeypatch):
    store, provider, delivery, ctx, price = await setup(
        "image_to_text", ['{"readable":false,"confidence":0,"text":""}'], tmp_path, monkeypatch
    )
    file = store.save(1, "input.jpg", b"test", "image/jpeg")
    oid = await submit(
        1, "image_to_text", {"images": [file.key], "language": "none"}, price, "bad-image"
    )
    await execute_job(ctx, str(await job_for(oid)))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        assert (await balance(db, 1)).available == 1000
        assert (await balance(db, 1)).reserved == 0
    assert delivery.errors[0][2] == "ocr_unclear" and len(delivery.errors) == 1
    assert not delivery.files


async def test_real_word_and_excel_outputs(tmp_path, monkeypatch):
    _, provider, delivery, ctx, price = await setup(
        "text_to_office",
        [
            '{"document":{"title":"تقرير","sections":[{"paragraphs":["بيانات مؤكدة"]}]}}',
            '{"table":{"title":"مبيعات","columns":["المنتج","العدد"],"rows":[["تفاح",12]]}}',
        ],
        tmp_path,
        monkeypatch,
    )
    for target in ("word", "excel"):
        oid = await submit(
            1, "text_to_office", {"text": "المعلومات", "target": target}, price, target
        )
        await execute_job(ctx, str(await job_for(oid)))
        async with sessions() as db:
            assert (await db.get(Order, oid)).status == "completed"
    assert any(p.text == "بيانات مؤكدة" for p in Word(BytesIO(delivery.files[0][1])).paragraphs)
    assert load_workbook(BytesIO(delivery.files[1][1])).active["B2"].value == 12
    assert provider.calls == 2


async def test_ambiguity_confirm_resumes_without_second_ai_charge(tmp_path, monkeypatch):
    _, provider, delivery, ctx, price = await setup(
        "text_to_office",
        [
            '{"ambiguous":true,"question":"اعتماد هذا الهيكل؟","document":{"title":"تقرير","sections":[{"heading":"ملخص","paragraphs":["النص"]}]}}'
        ],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(1, "text_to_office", {"text": "نص", "target": "word"}, price, "ambiguous")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await deliver_confirmations(ctx)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "waiting_confirmation"
        assert (await balance(db, 1)).reserved == price
    with pytest.raises(ServiceError, match="not_allowed"):
        await confirm_structure(oid, 2, True)
    await confirm_structure(oid, 1, True)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert provider.calls == 1 and len(delivery.files) == 1 and len(delivery.confirmations) == 1


async def test_ambiguity_cancel_releases(tmp_path, monkeypatch):
    _, _, delivery, ctx, price = await setup(
        "text_to_office",
        [
            '{"ambiguous":true,"question":"هل تعتمد؟","table":{"title":"جدول","columns":["عنصر"],"rows":[["نص"]]}}'
        ],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(
        1, "text_to_office", {"text": "نص", "target": "excel"}, price, "cancel-ambiguous"
    )
    await execute_job(ctx, str(await job_for(oid)))
    await confirm_structure(oid, 1, False)
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "cancelled"
        assert (await balance(db, 1)).available == 1000


async def test_real_service_invalid_json_releases_credit(tmp_path, monkeypatch):
    _, provider, _, ctx, price = await setup("text_to_office", ["{}", "{}"], tmp_path, monkeypatch)
    oid = await submit(1, "text_to_office", {"text": "نص", "target": "word"}, price, "invalid-json")
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "failed"
        assert (await balance(db, 1)).available == 1000
        assert (await db.get(Job, await job_for(oid))).cost_sar == Decimal(".008")
    assert provider.calls == 2


async def test_long_ocr_delivers_txt_docx_and_deletes_all_files(tmp_path, monkeypatch):
    import json

    text = "نص طويل " * 6000
    store, _, delivery, ctx, price = await setup(
        "image_to_text",
        [json.dumps({"readable": True, "confidence": 0.99, "text": text})],
        tmp_path,
        monkeypatch,
    )
    file = store.save(1, "input.jpg", b"test image", "image/jpeg")
    oid = await submit(
        1, "image_to_text", {"images": [file.key], "language": "none"}, price, "long-ocr"
    )
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).reserved == 0
    assert delivery.files[0][0] == "result.txt"
    assert delivery.files[0][1].decode() == text.strip()
    doc = Word(BytesIO(delivery.files[1][1]))
    assert " ".join(p.text for p in doc.paragraphs[1:]) == text.strip()
    assert not list(tmp_path.glob("[0-9]*/*/*"))


async def test_missing_information_releases_reservation(tmp_path, monkeypatch):
    _, _, _, ctx, price = await setup(
        "text_to_office",
        [
            '{"missing_information":true,"document":{"title":"تقرير","sections":[{"paragraphs":["بيانات غير كافية"]}]}}'
        ],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(
        1, "text_to_office", {"text": "غير مكتمل", "target": "word"}, price, "missing-info"
    )
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.error_key == "needs_information"
        assert (await balance(db, 1)).reserved == 0
        assert (await balance(db, 1)).available == 1000


async def test_ambiguity_expiry_releases_reservation(tmp_path, monkeypatch):
    from datetime import UTC, datetime, timedelta

    from app.files.retention import cleanup

    _, _, _, ctx, price = await setup(
        "text_to_office",
        [
            '{"ambiguous":true,"question":"تأكيد؟","document":{"title":"تقرير","sections":[{"paragraphs":["نص"]}]}}'
        ],
        tmp_path,
        monkeypatch,
    )
    monkeypatch.setattr(config(), "storage_root", tmp_path)
    oid = await submit(1, "text_to_office", {"text": "محتوى", "target": "word"}, price, "expired")
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions.begin() as db:
        (await db.get(Order, oid)).updated_at = datetime.now(UTC) - timedelta(minutes=31)
    await cleanup()
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == "cancelled" and order.error_key == "confirmation_expired"
        assert (await balance(db, 1)).available == 1000
        assert (await balance(db, 1)).reserved == 0


async def test_early_plugin_cancellation_eventually_deletes_uploaded_files(tmp_path, monkeypatch):
    from app.files.retention import cleanup_terminal_files
    from app.services.registry import registry

    store, _, _, ctx, price = await setup("image_to_text", [], tmp_path, monkeypatch)
    monkeypatch.setattr(config(), "storage_root", tmp_path)
    file = store.save(1, "input.jpg", b"fixture", "image/jpeg")
    oid = await submit(
        1, "image_to_text", {"images": [file.key], "language": "none"}, price, "cancel-before-run"
    )
    monkeypatch.setattr(registry.types["image_to_text"], "version", "2")
    await execute_job(ctx, str(await job_for(oid)))
    await cleanup_terminal_files()
    await cleanup_terminal_files()
    assert not list(tmp_path.glob("[0-9]*/*/*"))
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.files_deleted and order.status == "cancelled"
        assert (await balance(db, 1)).available == 1000


async def test_database_abort_after_generation_reuses_saved_result(tmp_path, monkeypatch):
    from datetime import UTC, datetime

    from sqlalchemy.exc import DBAPIError

    from app.workers import runner

    _, provider, delivery, ctx, price = await setup(
        "text_to_office",
        ['{"document":{"title":"تقرير","sections":[{"paragraphs":["بيانات"]}]}}'],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(
        1, "text_to_office", {"text": "بيانات", "target": "word"}, price, "db-abort-output"
    )
    jid = await job_for(oid)
    original = runner.lock_order
    calls = 0

    class Abort(Exception):
        sqlstate = "40P01"

    async def lock(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise DBAPIError("safe", None, Abort("deadlock"))
        return await original(*args, **kwargs)

    monkeypatch.setattr(runner, "lock_order", lock)
    await execute_job(ctx, str(jid))
    async with sessions.begin() as db:
        order = await db.get(Order, oid)
        assert order.status == "queued" and order.result
        (await db.get(Job, jid)).next_run_at = datetime.now(UTC)
    await execute_job(ctx, str(jid))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
    assert provider.calls == 1 and len(delivery.files) == 1


async def test_real_pptx_output(tmp_path, monkeypatch):
    _, provider, delivery, ctx, price = await setup(
        "text_to_pptx",
        ['{"deck":{"slides":[{"title":"عنوان","bullets":["نقطة أساسية"]}]}}'],
        tmp_path,
        monkeypatch,
    )
    oid = await submit(1, "text_to_pptx", {"text": "المعلومات"}, price, "pptx")
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        assert (await db.get(Order, oid)).status == "completed"
        assert (await balance(db, 1)).available == 1000 - price
    assert provider.calls == 1
    presentation = Presentation(BytesIO(delivery.files[0][1]))
    assert any(
        shape.has_text_frame and "نقطة أساسية" in shape.text
        for shape in presentation.slides[0].shapes
    )


async def test_pptx_missing_information_releases_reservation(tmp_path, monkeypatch):
    _, _, _, ctx, price = await setup(
        "text_to_pptx", ['{"missing_information":true}'], tmp_path, monkeypatch
    )
    oid = await submit(1, "text_to_pptx", {"text": "غير مكتمل"}, price, "pptx-missing")
    await execute_job(ctx, str(await job_for(oid)))
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.error_key == "needs_information"
        assert (await balance(db, 1)).reserved == 0
        assert (await balance(db, 1)).available == 1000
