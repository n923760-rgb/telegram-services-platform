"""Context-aware source repair preserves PostgreSQL settlement and cached delivery."""

import json
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from pptx import Presentation
from sqlalchemy import func, select

from app.core.db import sessions
from app.core.models import CostHold, CostUsage, Job, Ledger, Order
from app.orders.engine import submit
from app.providers.ai.gateway import AI
from app.providers.storage import OwnedStorage
from app.services.document_fidelity import CONTACT_TOKEN
from app.services.protected_values import ProtectedValues
from app.wallet.ledger import balance
from app.workers.runner import deliver_failures, execute_job
from tests.test_foundation import job_for
from tests.test_real_services import setup
from tests.test_schema_repair_context_unit import SOURCE, plans


@pytest.mark.parametrize(
    "slug,kind,references",
    [
        ("text_to_office", "word", False),
        ("text_to_pdf", "pdf", False),
        ("text_to_pptx", "deck", False),
        ("text_to_office", "word", True),
        ("text_to_pptx", "deck", True),
    ],
)
@pytest.mark.parametrize("recover", [True, False])
async def test_source_repair_context_captures_or_releases_once_and_records_all_costs(
    tmp_path, monkeypatch, slug, kind, references, recover
):
    _, valid, invalid, code = plans(kind)
    valid_text, invalid_text = json.dumps(valid), json.dumps(invalid)
    if references:
        protected = ProtectedValues(SOURCE, CONTACT_TOKEN)
        for marker, literal in protected.values.items():
            valid_text = valid_text.replace(literal, marker)
            invalid_text = invalid_text.replace(literal, marker)
    responses = (
        [invalid_text, valid_text, '{"accepted":true,"issues":[]}']
        if recover
        else [invalid_text] * 2
    )
    store, provider, delivery, ctx, price = await setup(slug, responses, tmp_path, monkeypatch)
    original = provider.extract_from_text
    generation_source = []

    async def checked(schema, source, prompt):
        if provider.calls == 0:
            generation_source.append(source)
        elif provider.calls == 1:
            repair = json.loads(source)
            assert repair["source"] == generation_source[0]
            assert repair["invalid_candidate"] == invalid_text
            assert repair["validation_categories"] == [code]
            assert "untrusted data" in prompt
        elif provider.calls == 2:
            review = json.loads(source)
            assert review["source"] == generation_source[0]
            assert "invalid_candidate" not in review
            assert "[[value-" not in json.dumps(review["candidate"])
        return await original(schema, source, prompt)

    monkeypatch.setattr(provider, "extract_from_text", checked)
    ctx["runtime_factory"] = lambda jid, uid: SimpleNamespace(
        ai=AI(provider, jid, uid, quality_review=True), storage=OwnedStorage(store, uid)
    )
    inputs = {"text": SOURCE}
    if kind == "deck":
        inputs["slide_count"] = "1"
    else:
        inputs["mode"] = "smart"
        if kind == "word":
            inputs["target"] = "word"
    oid = await submit(1, slug, inputs, price, "context-repair")
    jid = await job_for(oid)
    await execute_job(ctx, str(jid))
    await execute_job(ctx, str(jid))
    await deliver_failures(ctx)
    await deliver_failures(ctx)
    calls = 3 if recover else 2
    async with sessions() as db:
        order = await db.get(Order, oid)
        assert order.status == ("completed" if recover else "failed")
        assert order.error_key == (None if recover else "provider_invalid")
        funds = await balance(db, 1)
        assert (funds.available, funds.reserved) == (1000 - price if recover else 1000, 0)
        assert (await db.get(Job, jid)).cost_sar == Decimal(".004") * calls
        assert await db.scalar(select(func.count(CostUsage.id))) == calls
        assert list(await db.scalars(select(CostHold.status))) == ["settled"] * calls
        for ledger_kind, count in (
            ("reserve", 1),
            ("capture", int(recover)),
            ("release", int(not recover)),
        ):
            assert (
                await db.scalar(
                    select(func.count(Ledger.id)).where(
                        Ledger.order_id == oid, Ledger.kind == ledger_kind
                    )
                )
                == count
            )
    assert provider.calls == calls
    assert len(delivery.files) == int(recover)
    assert len(delivery.errors) == int(not recover)
    if references and recover:
        content = BytesIO(delivery.files[0][1])
        if kind == "word":
            text = "\n".join(p.text for p in Document(content).paragraphs)
        else:
            text = "\n".join(
                shape.text
                for slide in Presentation(content).slides
                for shape in slide.shapes
                if shape.has_text_frame
            )
        assert text.count("00123") == text.count("125.50") == 1
        assert "not been approved" in text and "[[value-" not in text
    assert not list(tmp_path.glob("[0-9]*/*/*"))
