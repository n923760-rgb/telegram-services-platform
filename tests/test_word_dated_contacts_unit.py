"""A numeric email local part must remain a complete LTR run in Arabic Word text."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

from app.builders import excel, word
from app.builders.schema import Table
from app.builders.word_schema import WordDocument

CONTACTS = ["2026-10-24@e.invalid", "24-10-2026@e.invalid", "2026-10-24.news@e.invalid"]


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("contact", CONTACTS)
def test_contact_first_matcher_consumes_the_whole_email_before_its_date_prefix(contact):
    text = "راجع " + contact + " قبل 2026/10/24"
    assert word.INLINE_LTR.findall(text) == [contact, "2026/10/24"]
    assert word.INLINE_LTR.sub("", text) == "راجع  قبل "


@pytest.mark.parametrize("contact", CONTACTS)
@pytest.mark.parametrize("professional", [False, True])
@pytest.mark.parametrize("template", [False, True])
def test_word_paragraphs_and_table_cells_keep_arabic_direction_and_complete_contacts(
    contact, professional, template
):
    text = "راجع " + contact
    plan = WordDocument(
        title="المتابعة",
        sections=[
            {"paragraphs": [text]},
            {"tables": [{"columns": ["المرجع"], "rows": [[text]]}]},
        ],
    )
    snapshot = plan.model_dump_json()
    doc = Document(
        BytesIO(word.build(plan, format_dates=professional, professional_template=template))
    )
    body = next(p for p in doc.paragraphs if "راجع" in p.text)
    cell = doc.tables[0].cell(1, 0).paragraphs[0]
    for paragraph in (body, cell):
        assert paragraph.text.replace("\u200e", "") == text
        if not professional:
            assert paragraph.text == text
        assert contact in paragraph.text
        assert paragraph._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "1"
        run = next(r for r in paragraph.runs if contact in r.text)
        assert run.text.replace("\u200e", "") == contact
        assert run._r.rPr.find(qn("w:rtl")).get(qn("w:val")) == "0"
    assert plan.model_dump_json() == snapshot
    saved = BytesIO()
    doc.save(saved)
    reopened = Document(BytesIO(saved.getvalue()))
    assert reopened.tables[0].cell(1, 0).text == doc.tables[0].cell(1, 0).text


@pytest.mark.parametrize("contact", CONTACTS)
def test_excel_keeps_literal_cell_and_uses_arabic_prose_direction(contact):
    text = "راجع " + contact
    sheet = load_workbook(
        BytesIO(excel.build(Table(title="المتابعة", columns=["المرجع"], rows=[[text]])))
    ).active
    assert sheet["A4"].value == text and sheet["A4"].data_type == "s"
    assert sheet["A4"].alignment.horizontal == "right"
    assert sheet["A4"].alignment.readingOrder == 2


@pytest.mark.parametrize("contact", CONTACTS)
def test_english_and_standalone_contacts_still_use_ltr_direction(contact):
    for text in (contact, "Contact " + contact):
        assert not excel._reading_rtl(text)
        paragraph = Document(
            BytesIO(word.build(WordDocument(title="Contacts", sections=[{"paragraphs": [text]}])))
        ).paragraphs[1]
        assert paragraph.text == text
        assert paragraph._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "0"


@pytest.mark.parametrize("template", [False, True])
def test_named_design_titles_headings_and_table_captions_keep_the_complete_contact(template):
    contact = CONTACTS[0]
    text = "راجع " + contact
    plan = WordDocument(
        title=text,
        design={"theme": "modern"},
        sections=[
            {
                "heading": text,
                "tables": [{"caption": text, "columns": ["المرجع"], "rows": [["00123"]]}],
            }
        ],
    )
    doc = Document(BytesIO(word.build(plan, format_dates=True, professional_template=template)))
    paragraphs = [p for p in doc.paragraphs if contact in p.text]
    assert len(paragraphs) == 3
    for paragraph in paragraphs:
        assert paragraph.text.replace("\u200e", "") == text
        assert paragraph._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "1"
        assert any(r.text == "\u200e" + contact + "\u200e" for r in paragraph.runs)
