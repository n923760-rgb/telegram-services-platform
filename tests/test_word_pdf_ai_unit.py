from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from pydantic import ValidationError

from app.builders.word import build
from app.services.base import ServiceError
from app.services.text_to_word_pdf.schema import Inputs, refined_document, revision_schema
from app.services.text_to_word_pdf.service import TextToWordPdf
from tests.word_pdf_ai_fixtures import inputs, revision


@pytest.fixture(autouse=True)
def database():
    yield


def test_mode_choice_discloses_ai_and_safe_skip_default():
    field = TextToWordPdf.input_schema.fields[1]
    assert field.name == "mode" and field.choices == ["direct", "smart"]
    assert not field.required and field.skip_key == "word_pdf_direct"
    assert not TextToWordPdf.needs_ai({"text": "Literal"})
    TextToWordPdf.input_schema.validate_inputs(inputs())
    assert TextToWordPdf.needs_ai(inputs())


def test_word_isolates_inline_contact_and_date_without_changing_source():
    values = Inputs.parse(inputs())
    plan = revision_schema(values).model_validate(revision())
    doc = Document(
        BytesIO(build(refined_document(plan), professional_template=True, format_dates=True))
    )
    text = "\n".join(p.text.replace("\u200e", "") for p in doc.paragraphs)
    assert all(line in text for line in values.lines())
    contact = next(p for p in doc.paragraphs if "sample@example.invalid" in p.text)
    token = next(
        run for run in contact.runs if run.text.replace("\u200e", "") == "sample@example.invalid"
    )
    assert 'w:rtl w:val="0"' in token._r.xml


@pytest.mark.parametrize("language", ["ar", "en"])
@pytest.mark.parametrize("theme", ["modern", "classic", "editorial"])
def test_source_owned_design_and_complete_ordered_blocks(language, theme):
    source = inputs(language)
    values = Inputs.parse(source)
    plan = revision_schema(values).model_validate(revision(source, theme))
    document = refined_document(plan)
    assert document.design.theme == theme
    assert [b.text for s in document.sections for b in s.blocks] == values.lines()
    assert not document.sections[0].tables


@pytest.mark.parametrize(
    "change",
    [
        "number",
        "contact",
        "missing",
        "duplicate",
        "order",
        "heading",
        "title",
        "newline",
        "control",
        "script",
    ],
)
def test_independent_source_guard_rejects_unsafe_output(change):
    source = inputs() | {"title": "طلب 00017"}
    data = deepcopy(revision(source))
    rows = data["sections"][0]["segments"]
    if change == "number":
        rows[1]["text"] = rows[1]["text"].replace("125.50", "999.50")
    elif change == "contact":
        rows[2]["text"] = rows[2]["text"].replace("sample@example.invalid", "other@example.invalid")
    elif change == "missing":
        rows.pop()
    elif change == "duplicate":
        rows.append(rows[0])
    elif change == "order":
        rows.reverse()
    elif change == "heading":
        data["sections"][0]["heading"] = "الطلب 00123"
    elif change == "title":
        data["title"] = "طلب جديد"
    elif change == "newline":
        rows[0]["text"] += "\nنص"
    elif change == "control":
        rows[0]["text"] += "\x00"
    elif change == "script":
        data["design"]["theme"] = "<script>"
    with pytest.raises(ValidationError):
        revision_schema(Inputs.parse(source)).model_validate(data)


@pytest.mark.parametrize("text", ["x" * 9001, "x" * 4001, "\n".join(["x"] * 81), "\x00", "\ud800"])
def test_smart_limits_reject_before_admission(text):
    with pytest.raises(ServiceError, match="input_invalid"):
        TextToWordPdf.needs_ai({"text": text, "mode": "smart"})


async def test_injected_response_revalidated_before_render_or_write():
    data = revision()
    data["sections"][0]["segments"][1]["text"] = "999.50"
    from app.services.text_to_word_pdf.schema import Revision

    renderer = SimpleNamespace(word_pdf=AsyncMock())
    service = TextToWordPdf()
    service.runtime = SimpleNamespace(
        ai=SimpleNamespace(extract=AsyncMock(return_value=Revision.model_validate(data))),
        renderer=renderer,
    )
    with pytest.raises(ServiceError, match="provider_invalid"):
        await service.run(inputs())
    renderer.word_pdf.assert_not_called()
