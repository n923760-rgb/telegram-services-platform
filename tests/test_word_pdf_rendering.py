from io import BytesIO

import pdfplumber

from tests.word_pdf_fixtures import document, renderer


async def test_native_arabic_date_has_correct_visual_glyph_order():
    result = await renderer().word_pdf(document())
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        # Reading-order extraction can look correct even when rendered glyphs are reversed.
        lines = {}
        for char in pdf.pages[0].chars:
            lines.setdefault(round(char["top"]), []).append(char)
        visual_lines = [
            "".join(c["text"] for c in sorted(chars, key=lambda c: c["x0"]))
            for chars in lines.values()
        ]
    assert sum("2026-10-09" in line for line in visual_lines) == 2
    assert not any("09-10-2026" in line for line in visual_lines)


async def test_native_export_keeps_word_text_and_embedded_fonts_on_a4_page():
    from docx import Document
    from pypdf import PdfReader

    from tests.word_pdf_fixtures import TEXT

    result = await renderer().word_pdf(document())
    doc = Document(BytesIO(result.docx))
    assert [p.text.replace("\u200e", "") for p in doc.paragraphs] == [
        "تقرير الطلبات",
        *TEXT.splitlines(),
    ]
    reader = PdfReader(BytesIO(result.pdf), strict=True)
    assert len(reader.pages) == 1
    page = reader.pages[0]
    assert abs(float(page.mediabox.width) - 595.28) < 1
    assert abs(float(page.mediabox.height) - 841.89) < 1
    text = page.extract_text()
    assert text.count("00123") == text.count("125.50") == text.count("2026-10-09") == 2
    fonts = list(page["/Resources"]["/Font"].values())
    assert any("DejaVuSans" in str(font.get_object()["/BaseFont"]) for font in fonts)
    for font in fonts:
        node = font.get_object()
        if node["/Subtype"] == "/Type0":
            node = node["/DescendantFonts"][0].get_object()
        descriptor = node.get("/FontDescriptor")
        assert descriptor and any(
            key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3")
        )


async def test_native_export_preserves_literal_markup_blank_lines_and_no_fallback_title():
    from docx import Document

    from app.builders.word_schema import WordDocument

    values = ["  Invoice 00123  ", "", "{{ 7 * 7 }} <tag> & {% if True %}", "سطر عربي"]
    result = await renderer().word_pdf(
        WordDocument(title="Hidden", sections=[{"paragraphs": values}]), include_title=False
    )
    assert [p.text for p in Document(BytesIO(result.docx)).paragraphs] == values
    from pypdf import PdfReader

    text = PdfReader(BytesIO(result.pdf)).pages[0].extract_text()
    assert "Hidden" not in text and "{{ 7 * 7 }} <tag> & {% if True %}" in text


async def test_native_multpage_table_preserves_identifiers_dates_and_headers():
    from pypdf import PdfReader

    from tests.word_pdf_fixtures import table_document

    result = await renderer().word_pdf(table_document())
    pages = PdfReader(BytesIO(result.pdf)).pages
    assert 2 <= len(pages) <= 5
    text = "\n".join(page.extract_text() for page in pages)
    for identifier in range(100, 155):
        assert text.count(f"{identifier:05d}") == 1
    assert text.count("2026-10-09") == text.count("125.50") == 55
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        for page in pdf.pages:
            assert all(
                0 <= char["x0"] < char["x1"] <= page.width
                and 0 <= char["top"] < char["bottom"] <= page.height
                for char in page.chars
            )
            assert "الرقم" in page.extract_text()[::-1] or "الرقم" in page.extract_text()
