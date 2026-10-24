from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.builders import word
from app.builders.word_schema import WordDocument, WordTable
from tests.word_pdf_fixtures import wide_table_document


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("template", [False, True])
@pytest.mark.parametrize("theme", ["modern", "classic", "editorial"])
@pytest.mark.parametrize("arabic", [False, True])
def test_table_uses_final_design_margins_and_preserves_every_cell(template, theme, arabic):
    source = wide_table_document(arabic=arabic, design=theme)
    before = source.model_dump()
    doc = Document(BytesIO(word.build(source, professional_template=template)))
    section = doc.sections[0]
    available = section.page_width - section.left_margin - section.right_margin
    table = doc.tables[0]
    assert abs(sum(column.width for column in table.columns) - available) <= 3810
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        source.sections[0].tables[0].columns,
        *source.sections[0].tables[0].rows,
    ]
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows)
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(
        cell.paragraphs[0].runs[0].font.size == Pt(10) for row in table.rows for cell in row.cells
    )
    assert source.model_dump() == before
    if not arabic:
        for index, token in [(0, "Reference"), (2, "Responsible"), (4, "Amount")]:
            font = word._table_font(True)
            assert table.columns[index].width >= Pt(font.getlength(token) / 4 + 10)


@pytest.mark.parametrize("date", [False, True])
def test_impossible_minima_and_long_identifiers_stay_bounded(date):
    values = ["2026-10-25" if date else "W" * 240] * 6
    source = WordTable(columns=list("ABCDEF"), rows=[values])
    available = Cm(16)
    widths = word._table_widths(source, available)
    assert all(width > 0 for width in widths)
    assert 0 <= available - sum(widths) < 6


def test_missing_local_font_keeps_editable_docx_generation_available(monkeypatch):
    def missing(*args, **kwargs):
        raise OSError("font unavailable")

    monkeypatch.setattr(word.ImageFont, "truetype", missing)
    word._table_font.cache_clear()
    data = WordDocument(
        title="Report",
        sections=[
            {
                "tables": [
                    {"columns": ["Reference", "Description"], "rows": [["00123", "Source text"]]}
                ]
            }
        ],
    )
    try:
        doc = Document(BytesIO(word.build(data)))
    finally:
        word._table_font.cache_clear()
    assert [cell.text for cell in doc.tables[0].rows[1].cells] == ["00123", "Source text"]


def test_plain_table_without_design_uses_existing_page_margins():
    data = wide_table_document().model_copy(update={"design": None})
    doc = Document(BytesIO(word.build(data, professional_template=True)))
    section = doc.sections[0]
    assert section.left_margin == section.right_margin
    assert abs(section.left_margin - Cm(2.5)) <= 635
    available = section.page_width - section.left_margin - section.right_margin
    assert abs(sum(column.width for column in doc.tables[0].columns) - available) <= 3810
