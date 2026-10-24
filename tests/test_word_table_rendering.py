from io import BytesIO

import pdfplumber
import pytest
from docx import Document

from tests.word_pdf_fixtures import renderer, wide_table_document


@pytest.fixture(autouse=True)
def database():
    yield


@pytest.mark.parametrize("arabic", [False, True])
async def test_wide_multpage_table_keeps_words_records_and_repeated_headers(arabic):
    source = wide_table_document(arabic=arabic)
    result = await renderer().word_pdf(source)
    doc = Document(BytesIO(result.docx))
    assert len(doc.tables[0].rows) == 36
    with pdfplumber.open(BytesIO(result.pdf)) as pdf:
        assert 2 <= len(pdf.pages) <= 8
        text = "\n".join(page.extract_text() for page in pdf.pages)
        for index in range(100, 135):
            assert text.count(f"{index:05d}") == 1
        assert text.count("2026-10-25") == text.count("125.50") == 35
        for page in pdf.pages:
            assert all(
                0 <= char["x0"] < char["x1"] <= page.width
                and 0 <= char["top"] < char["bottom"] <= page.height
                for char in page.chars
            )
            if not arabic:
                words = {word["text"] for word in page.extract_words()}
                assert {"Reference", "Responsible", "department", "Amount", "information"} <= words
    # DOCX retains native editable values independently of PDF glyph extraction.
    for row, expected in zip(
        doc.tables[0].rows[1:], source.sections[0].tables[0].rows, strict=True
    ):
        assert [cell.text.replace("\u200e", "") for cell in row.cells] == expected
