"""Report title typography must preserve editable content and script parity."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.builders.word import build
from app.builders.word_schema import WordDocument


@pytest.mark.parametrize("template", [False, True])
@pytest.mark.parametrize("theme", ["modern", "classic", "editorial"])
@pytest.mark.parametrize(
    "title",
    ["تقرير متابعة مشروع متجر إلكتروني", "Project progress report for the online store"],
)
def test_long_report_title_keeps_content_and_matching_arabic_latin_sizes(template, theme, title):
    plan = WordDocument(
        title=title,
        design={"theme": theme},
        sections=[
            {"heading": "Status", "paragraphs": ["00123 remains pending until 2026-11-15."]},
            {"tables": [{"columns": ["ID", "Amount"], "rows": [["00017", "125.50"]]}]},
            {"title": title, "page_break_before": True, "paragraphs": ["Keep this second title."]},
        ],
    )
    snapshot = plan.model_dump()
    output = Document(BytesIO(build(plan, professional_template=template)))
    assert plan.model_dump() == snapshot
    assert [p.text for p in output.paragraphs if p.text] == [
        title,
        "Status",
        "00123 remains pending until 2026-11-15.",
        title,
        "Keep this second title.",
    ]
    assert [[c.text for c in r.cells] for r in output.tables[0].rows] == [
        ["ID", "Amount"],
        ["00017", "125.50"],
    ]
    titles = [p for p in output.paragraphs if p.style.name == "Title"]
    for paragraph in titles:
        for run in paragraph.runs:
            assert run.font.size.pt == 20
            assert run._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "40"
    assert titles[-1].paragraph_format.page_break_before


@pytest.mark.parametrize("template", [False, True])
def test_unthemed_literal_export_does_not_gain_report_design(template):
    plan = WordDocument(
        title="Keep this literal customer title unchanged", sections=[{"paragraphs": ["Body"]}]
    )
    output = Document(BytesIO(build(plan, include_title=False, professional_template=template)))
    assert [p.text for p in output.paragraphs] == ["Body"]
